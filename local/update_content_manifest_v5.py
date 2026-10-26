from __future__ import annotations

from pathlib import Path

from docx import Document


ROOT = Path(r"F:\作业\aibola\reports\nc_submission_revision_20260824\NC_Submission_Updated_20260824")
SOURCE = ROOT / "Content_Manifest_Updated_v4_20260824.docx"
TARGET = ROOT / "Content_Manifest_Updated_v5_20260824.docx"

REPLACEMENTS = {
    "Updated 2026-08-24 | Original Figs. 1 and 3 retained; Fig. 6 and Supplementary Figs. 1–4 code-generated in English":
        "Updated 2026-08-24 | Original Figs. 1 and 3 retained; English Supplementary Figs. 1–4 embedded in both main text and SI",
    "DiseaseX_Nexus_NC_Manuscript_v6_20260824.docx/.pdf":
        "DiseaseX_Nexus_NC_Manuscript_v7_20260824.docx/.pdf",
    "Supplementary/SI_Figures/ and SUPPLEMENTARY_ENGLISH_FIGURE_EMBEDDING_V4.json":
        "Supplementary/SI_Figures/, MAIN_ENGLISH_FIGURE_EMBEDDING_V7.json and SUPPLEMENTARY_ENGLISH_FIGURE_EMBEDDING_V4.json",
    "English Supplementary Figs. 1–4, retained original Main Figs. 1 and 3, and byte-identical DOCX embedding attestation":
        "English Supplementary Figs. 1–4, retained original Main Figs. 1 and 3, and byte-identical embedding attestations for both main text and SI",
}


def replace_text(paragraph) -> int:
    original = paragraph.text
    updated = original
    for old, new in REPLACEMENTS.items():
        updated = updated.replace(old, new)
    if updated == original:
        return 0
    paragraph.text = updated
    return 1


document = Document(SOURCE)
changes = sum(replace_text(paragraph) for paragraph in document.paragraphs)
for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            changes += sum(replace_text(paragraph) for paragraph in cell.paragraphs)
document.save(TARGET)
print(f"changes={changes}")
