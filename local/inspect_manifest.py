from pathlib import Path
from docx import Document

p=Path(r'F:\作业\aibola\reports\nc_submission_revision_20260824\NC_Submission_Updated_20260824\Content_Manifest_Updated_20260824.docx')
d=Document(p)
print('PARAGRAPHS')
for i,x in enumerate(d.paragraphs): print(i, x.style.name, repr(x.text))
print('TABLES',len(d.tables))
for ti,t in enumerate(d.tables):
 print('TABLE',ti,len(t.rows),len(t.columns))
 for row in t.rows: print([c.text for c in row.cells])
