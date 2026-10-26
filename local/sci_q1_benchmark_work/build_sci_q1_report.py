from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"F:\作业\aibola")
REPORT_DIR = ROOT / "reports" / "sci_q1_ewnet_benchmark_optimization_20260822"
EVIDENCE_DIR = REPORT_DIR / "evidence"
OUTPUT_DIR = REPORT_DIR / "figures"
DOCX_PATH = REPORT_DIR / "SCI一区论文EWNet对比试验优化结果报告_20260822.docx"

NAVY = "17365D"
BLUE = "2E74B5"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "6B7280"
GREEN = "2F855A"
GOLD = "9A6700"
RED = "9B1C1C"
WHITE = "FFFFFF"
BLACK = "111827"


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_fixed(table, widths_dxa: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa, strict=True):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, size=11, bold=False, color=BLACK, italic=False) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.extend([r_fonts, color, underline])
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_para(doc: Document, text: str = "", *, bold=False, color=BLACK, size=11, align=None, before=0, after=6, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, color=color, italic=italic)
    return p


def add_bullet(doc: Document, text: str, *, level=0) -> None:
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.167
    p.paragraph_format.left_indent = Inches(0.5 if level == 0 else 0.75)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    r = p.add_run(text)
    set_run_font(r)


def add_callout(doc: Document, title: str, text: str, fill: str, accent: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_fixed(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=11.5, bold=True, color=accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.10
    r2 = p2.add_run(text)
    set_run_font(r2, size=10.5)
    add_para(doc, after=2)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_dxa: list[int], font_size=9.5) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_fixed(table, widths_dxa)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=font_size, bold=True, color=NAVY)
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
    set_table_fixed(table, widths_dxa)
    add_para(doc, after=2)


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, before, after, color in (
        ("Heading 1", 16, 16, 8, BLUE),
        ("Heading 2", 13, 12, 6, BLUE),
        ("Heading 3", 12, 8, 4, NAVY),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("疫情自动获取系统｜SCI 一区论文对比试验")
    set_run_font(r, size=8.5, color=MID_GRAY)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "D7DBE2")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("内部研究报告  |  第 ")
    set_run_font(fr, size=8.5, color=MID_GRAY)
    add_field(fp, "PAGE")
    fr2 = fp.add_run(" 页")
    set_run_font(fr2, size=8.5, color=MID_GRAY)


def add_picture(doc: Document, path: Path, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(path), width=Inches(6.3))
    cp = add_para(doc, caption, size=9, color=MID_GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    cp.paragraph_format.keep_with_next = False


def generate_figures(metrics: list[dict[str, str]], paired: list[dict[str, str]], selection_score: float) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Arial Unicode MS", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False

    order = [
        "last_value",
        "EWNet_author_code",
        "rolling_mean_4",
        "kg_validation_only_convex_ensemble.v2",
        "seasonal_naive_52",
    ]
    labels = {
        "last_value": "Last-value 基线",
        "EWNet_author_code": "论文 EWNet（作者代码）",
        "rolling_mean_4": "4周滚动均值",
        "kg_validation_only_convex_ensemble.v2": "项目优化组合",
        "seasonal_naive_52": "52周季节基线",
    }
    by_method = {r["method"]: r for r in metrics}
    vals = [float(by_method[m]["wape"]) * 100 for m in order]
    colors = ["#6B7280", "#2563EB", "#9CA3AF", "#D97706", "#CBD5E1"]
    fig, ax = plt.subplots(figsize=(10, 5.3), dpi=180)
    bars = ax.barh([labels[m] for m in order], vals, color=colors)
    ax.axvline(30, color="#9B1C1C", ls="--", lw=1.4, label="目标门槛 30%")
    for b, v in zip(bars, vals, strict=True):
        ax.text(v + 0.7, b.get_y() + b.get_height() / 2, f"{v:.2f}%", va="center", fontsize=10)
    ax.set_xlabel("WAPE（越低越好）")
    ax.set_title("锁定测试集：各方法误差对比（2015-W41 至 W53）", fontweight="bold")
    ax.set_xlim(0, max(vals) + 10)
    ax.grid(axis="x", alpha=0.25)
    ax.legend(loc="lower right", frameon=False)
    fig.tight_layout()
    fig.savefig(OUTPUT_DIR / "01_test_wape_comparison.png", bbox_inches="tight", facecolor="white")
    plt.close(fig)

    weeks = [f"W{int(r['week']):02d}" for r in paired]
    x = np.arange(len(weeks))
    fig, ax = plt.subplots(figsize=(10, 5.5), dpi=180)
    ax.plot(x, [float(r["actual"]) for r in paired], marker="o", lw=2.6, color="#111827", label="实际病例")
    ax.plot(x, [float(r["EWNet_author_code"]) for r in paired], marker="s", lw=1.8, color="#2563EB", label="论文 EWNet")
    ax.plot(x, [float(r["kg_validation_only_convex_ensemble.v2"]) for r in paired], marker="^", lw=1.8, color="#D97706", label="项目优化组合")
    ax.plot(x, [float(r["last_value"]) for r in paired], ls="--", lw=1.5, color="#6B7280", label="Last-value")
    ax.set_xticks(x, weeks)
    ax.set_ylabel("周病例数")
    ax.set_title("锁定测试周实际值与预测轨迹", fontweight="bold")
    ax.grid(alpha=0.22)
    ax.legend(ncol=2, frameon=False)
    fig.tight_layout()
    fig.savefig(OUTPUT_DIR / "02_test_trajectory.png", bbox_inches="tight", facecolor="white")
    plt.close(fig)

    selected_test = float(by_method["kg_validation_only_convex_ensemble.v2"]["wape"]) * 100
    fig, ax = plt.subplots(figsize=(9.5, 4.8), dpi=180)
    bars = ax.bar(["历史滚动折 +\n论文验证折", "锁定测试折"], [selection_score * 100, selected_test], color=["#2F855A", "#D97706"], width=0.55)
    for b, v in zip(bars, [selection_score * 100, selected_test], strict=True):
        ax.text(b.get_x() + b.get_width() / 2, v + 1.0, f"{v:.2f}%", ha="center", fontsize=12, fontweight="bold")
    ax.axhline(30, color="#9B1C1C", ls="--", lw=1.4)
    ax.text(1.42, 30.8, "30% 门槛", ha="right", color="#9B1C1C", fontsize=9)
    ax.set_ylabel("WAPE（越低越好）")
    ax.set_title("项目优化组合的验证—测试泛化差距", fontweight="bold")
    ax.set_ylim(0, 55)
    ax.grid(axis="y", alpha=0.22)
    fig.tight_layout()
    fig.savefig(OUTPUT_DIR / "03_generalization_gap.png", bbox_inches="tight", facecolor="white")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(10, 5.3), dpi=180)
    actual = np.array([float(r["actual"]) for r in paired])
    for key, label, color in (
        ("EWNet_author_code", "论文 EWNet", "#2563EB"),
        ("kg_validation_only_convex_ensemble.v2", "项目优化组合", "#D97706"),
        ("last_value", "Last-value", "#6B7280"),
    ):
        pred = np.array([float(r[key]) for r in paired])
        ax.plot(x + 1, np.abs(actual - pred), marker="o", lw=1.8, label=label, color=color)
    ax.set_xlabel("预测步长（周）")
    ax.set_ylabel("绝对误差（病例数）")
    ax.set_title("误差随预测步长扩大：突发上升未被模型捕获", fontweight="bold")
    ax.grid(alpha=0.22)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(OUTPUT_DIR / "04_horizon_absolute_error.png", bbox_inches="tight", facecolor="white")
    plt.close(fig)


def build_docx() -> None:
    summary = read_json(EVIDENCE_DIR / "outputs" / "final_summary.json")
    prior = read_json(EVIDENCE_DIR / "prior_comparison_summary.json")
    metrics = read_csv(EVIDENCE_DIR / "outputs" / "test_metrics.csv")
    paired = read_csv(EVIDENCE_DIR / "outputs" / "test_paired_predictions.csv")
    lock = read_json(EVIDENCE_DIR / "outputs" / "locked_selection.json")
    generate_figures(metrics, paired, float(lock["selected_score"]))

    by_method = {r["method"]: r for r in metrics}
    selected = by_method["kg_validation_only_convex_ensemble.v2"]
    ewnet = by_method["EWNet_author_code"]
    baseline = by_method["last_value"]
    old_project = prior["project_method"]["metrics"]
    test_actual = np.array([float(r["actual"]) for r in paired])
    early_mean = test_actual[:5].mean()
    late_mean = test_actual[-5:].mean()
    surge_ratio = late_mean / early_mean

    doc = Document()
    configure_doc(doc)

    add_para(doc, "技术研究报告", bold=True, color=BLUE, size=10.5, before=12, after=2)
    add_para(doc, "SCI 一区论文 EWNet 对比试验优化结果", bold=True, color=NAVY, size=23, after=4)
    add_para(doc, "同数据、同切分、测试集一次性开启；如实记录负泛化结果", color=MID_GRAY, size=13, after=14)

    meta = [
        ["项目", "疫情自动获取系统 / epidemic-openclaw"],
        ["基准论文", "Panja et al. (2023), Neural Networks, DOI 10.1016/j.neunet.2023.05.049"],
        ["数据集", "新加坡登革热周病例（838 周）"],
        ["执行日期", "2026-08-22"],
        ["最终判定", "优化协议完成；独立测试未优于论文方法或最强简单基线"],
    ]
    add_table(doc, ["项目", "内容"], meta, [2700, 6660], font_size=10)
    add_callout(
        doc,
        "结论先行",
        "本轮满足“选择 SCI 一区论文并做无泄漏对比”的要求，但优化模型没有在锁定测试集上取得提升。项目组合模型的测试 WAPE 为 45.51%，高于论文 EWNet 的 41.07% 和 Last-value 基线的 40.05%。因此不能把 19.16% 的验证分数表述为最终准确率，也不能据此替换现有预测方案。",
        "FFF4E5",
        GOLD,
    )

    doc.add_heading("1. 任务范围与最终状态", level=1)
    add_para(doc, "本轮只在新的隔离实验目录中执行，不部署、不写数据库、不创建正式 forecast run、不 promotion，也不修改 active pointer。目标是把论文资格、数据切分、模型选择、独立测试和失败原因放在同一可审计链路中。")
    add_table(
        doc,
        ["验收项", "结果", "说明"],
        [
            ["论文等级", "PASS", "Neural Networks 为 SCIE；2024 JCR 在 Computer Science, AI 与 Neurosciences 均为 Q1"],
            ["同数据同切分", "PASS", "838 周；训练 799、验证 26、测试 13"],
            ["测试泄漏防护", "PASS", "模型选择期间未读取 2015-W41 至 W53 的测试标签"],
            ["优化效果", "FAIL", "项目优化组合未优于 EWNet 或最强简单基线"],
            ["非变更边界", "PASS", "数据库/正式预测/发布/指针变更均为 0"],
        ],
        [2350, 1250, 5760],
        font_size=9.2,
    )
    add_para(doc, "最终状态：SCI_Q1_EWNET_OPTIMIZATION_COMPLETED_NEGATIVE_GENERALIZATION_RESULT", bold=True, color=RED, size=10.5)

    doc.add_heading("2. 对比论文资格核验", level=1)
    add_para(doc, "选定论文：Epicasting: An Ensemble Wavelet Neural Network for forecasting epidemics。论文发表于 Elsevier 期刊 Neural Networks，卷 165（2023），页 185–212。该论文提出 EWNet：先用最大重叠离散小波变换分解非平稳疫情序列，再由自回归前馈神经网络分别建模各分量并集成。")
    add_table(
        doc,
        ["证据维度", "核验结果"],
        [
            ["正式出版身份", "ScienceDirect 正式文章；DOI 10.1016/j.neunet.2023.05.049"],
            ["期刊等级", "SCIE；2024 JCR Q1（Computer Science, Artificial Intelligence；Neurosciences）"],
            ["作者实现", "mad-stat/Epicasting；固定 commit 3740776f65c6f4157e67e6a78f71daeade45246b"],
            ["可比性", "作者仓库提供 15 个疫情数据集、EWNet 实现及建议超参数；本项目沿用相同新加坡登革热数据"],
        ],
        [2400, 6960],
        font_size=9.4,
    )
    p = add_para(doc, "来源：", size=9.5, color=MID_GRAY, after=2)
    add_hyperlink(p, "ScienceDirect 论文页", "https://www.sciencedirect.com/science/article/abs/pii/S0893608023002939")
    p.add_run("；")
    add_hyperlink(p, "作者 GitHub 仓库", "https://github.com/mad-stat/Epicasting")
    p.add_run("；")
    add_hyperlink(p, "2024 JCR/SCIE 分区清单", "https://uefiscdi.gov.ro/resource-860007-clasament.ais.jcr.2024.pdf")

    doc.add_heading("3. 数据与无泄漏实验协议", level=1)
    add_table(
        doc,
        ["阶段", "范围", "用途"],
        [
            ["训练", "前 799 周", "拟合候选模型"],
            ["论文验证折", "2015-W15 至 W40（26 周）", "与 2010–2014 同季节滚动折共同选择模型"],
            ["同季节滚动折", "2010–2014 各 W41 至 W53（5×13 周）", "检验跨年季节泛化；权重 70%"],
            ["锁定测试", "2015-W41 至 W53（13 周）", "选择完成并写入锁文件后，仅开启一次"],
        ],
        [2100, 3300, 3960],
        font_size=9.3,
    )
    add_bullet(doc, "候选池共 65 个配置，61 个完成；包含 Ridge、季节缩放、历史类比、树模型与简单基线。")
    add_bullet(doc, "选择目标为：70% × 2010–2014 同季节滚动折 WAPE + 30% × 论文验证折 WAPE。")
    add_bullet(doc, "最终组合权重只由训练/验证信息决定：Ridge(raw, α=100) 31.84%；Ridge(log1p, α=10) 35.44%；12 年季节缩放 32.72%。")
    add_bullet(doc, f"锁定选择 SHA-256：{lock['lock_sha256']}。")
    add_bullet(doc, "WAPE 越低越好；本文中的“预测吻合度”仅定义为 max(0, 1−WAPE)×100%，不是分类准确率。")

    doc.add_heading("4. 独立测试结果", level=1)
    metric_rows = []
    label_map = {
        "last_value": "Last-value 基线",
        "EWNet_author_code": "论文 EWNet（作者代码）",
        "rolling_mean_4": "4周滚动均值",
        "kg_validation_only_convex_ensemble.v2": "项目优化组合",
        "seasonal_naive_52": "52周季节基线",
    }
    for method in ("last_value", "EWNet_author_code", "rolling_mean_4", "kg_validation_only_convex_ensemble.v2", "seasonal_naive_52"):
        row = by_method[method]
        metric_rows.append([
            label_map[method],
            f"{float(row['wape']) * 100:.2f}%",
            f"{float(row['agreement']) * 100:.2f}%",
            f"{float(row['mae']):.2f}",
            f"{float(row['rmse']):.2f}",
        ])
    add_table(doc, ["方法", "WAPE", "吻合度*", "MAE", "RMSE"], metric_rows, [3300, 1350, 1550, 1500, 1660], font_size=9.2)
    add_para(doc, "* 吻合度=max(0,1−WAPE)×100%，仅用于直观表达计数预测误差，不是分类准确率。", size=8.8, color=MID_GRAY, after=8)
    add_picture(doc, OUTPUT_DIR / "01_test_wape_comparison.png", "图 1  锁定测试集 WAPE。项目优化组合未越过 30% 目标，也未优于论文 EWNet 或 Last-value。")
    add_picture(doc, OUTPUT_DIR / "02_test_trajectory.png", "图 2  2015-W41 至 W53 的实际值与预测轨迹。三类预测均未捕获后半段快速上升。")

    doc.add_heading("5. 优化前后对比", level=1)
    optimized_wape = float(selected["wape"])
    old_wape = float(old_project["wape"])
    add_table(
        doc,
        ["项目方法版本", "测试 WAPE", "相对论文 EWNet", "结论"],
        [
            ["原项目 lagged_count_boosting.v1", f"{old_wape * 100:.2f}%", f"高 {(old_wape / float(ewnet['wape']) - 1) * 100:.2f}%", "接近 EWNet，但未超过"],
            ["本轮验证集优化组合", f"{optimized_wape * 100:.2f}%", f"高 {(optimized_wape / float(ewnet['wape']) - 1) * 100:.2f}%", "优化后反而退化"],
        ],
        [3000, 1600, 2100, 2660],
        font_size=9.3,
    )
    add_picture(doc, OUTPUT_DIR / "03_generalization_gap.png", "图 3  项目优化组合在选择样本上 WAPE 为 19.16%，但锁定测试升至 45.51%，存在明显泛化差距。")
    add_callout(
        doc,
        "为什么不能继续调到测试集变好？",
        "测试集已按协议开启一次。此后若根据测试结果继续改权重、特征或参数，测试集就会变成新的验证集，所得数字不能再作为独立对比结果。正确做法是保留本次负结果，并在新的未见年份或新数据集上预注册下一轮方案。",
        "FDECEC",
        RED,
    )

    doc.add_heading("6. 泛化失败的可解释原因", level=1)
    add_bullet(doc, f"测试后半段发生强烈水平跃迁：前 5 周平均 {early_mean:.1f} 例，后 5 周平均 {late_mean:.1f} 例，约为前者的 {surge_ratio:.2f} 倍。")
    add_bullet(doc, "所有方法都在 W40 之后进行 13 周多步预测，期间没有吸收新观测；因此无法在线修正对突发上升的低估。")
    add_bullet(doc, "项目候选只使用单变量历史病例数。没有天气、媒介密度、干预或报告延迟等外生信号时，提前识别结构突变的能力有限。")
    add_bullet(doc, "组合权重显著受 2015-W15 至 W40 的平稳验证段影响，而测试段的分布与其不同，导致“验证上最优”不等于“新阶段最优”。")
    add_bullet(doc, "EWNet 同样低估突发上升，但其小波分解与非线性分量建模使测试误差略小于项目组合；简单 Last-value 在该 13 周窗口反而最稳。")
    add_picture(doc, OUTPUT_DIR / "04_horizon_absolute_error.png", "图 4  误差随预测步长扩大。W49 以后实际病例快速上升，项目组合误差增长最快。")

    doc.add_heading("7. 结果应如何表述", level=1)
    add_table(
        doc,
        ["可以表述", "不可表述"],
        [
            ["已完成 SCI JCR Q1 论文的同数据、同切分、同测试窗对比", "项目优化后准确率达到 70%"],
            ["模型选择过程未使用锁定测试标签", "19.16% 是独立测试 WAPE"],
            ["独立测试中 EWNet 与 Last-value 优于本轮项目组合", "项目已超越论文方法"],
            ["负结果揭示突发上升与长步长预测的泛化风险", "本次单一 13 周窗口代表所有国家/病种"],
        ],
        [4680, 4680],
        font_size=9.3,
    )
    add_para(doc, "推荐论文叙述：本项目建立了面向高等级论文基准的可审计复现实验链路。验证集驱动的组合模型未能在独立突发期泛化，说明疫情预测系统必须同时报告基线、分布漂移、预测步长和负结果，不能只展示最优验证分数。")

    doc.add_heading("8. 下一步优化方案（不污染本次测试）", level=1)
    add_table(
        doc,
        ["优先级", "方案", "验收方式"],
        [
            ["P0", "把长达 13 周的一次性预测改为每周滚动更新，同时保留 1–4 周固定 horizon", "在 3 个未见季节上预注册 rolling-origin，逐周更新且无未来泄漏"],
            ["P0", "建立分布漂移/突发门禁；检测到水平跃迁时切换 EWNet、Last-value 与稳健组合", "门禁只使用 origin 之前数据；报告切换次数和误报"],
            ["P1", "引入外生变量版本（天气、媒介、报告延迟），与纯病例数模型分轨比较", "同一 target week、同一可用时点、消融实验"],
            ["P1", "采用 nested rolling-origin：外层留一年测试，内层调参，禁止复用已开启窗口", "至少 3 个外层年份；报告均值、分布和失败年份"],
            ["P2", "模型平均由静态权重升级为受约束动态权重，但保留基线 guard", "新 fresh 窗口 WAPE 改善且覆盖率不下降"],
        ],
        [1050, 5000, 3310],
        font_size=9.0,
    )

    doc.add_heading("9. 审计与可复现性", level=1)
    audit_rows = [
        ["数据工作簿 SHA-256", summary["data"]["workbook_sha256"]],
        ["派生 CSV SHA-256", lock["data_sha256"]],
        ["选择锁 SHA-256", lock["lock_sha256"]],
        ["作者仓库 commit", summary["paper"]["author_repository_commit"]],
        ["测试", "VERIFY_PASS；6/6 单元测试通过"],
        ["测试集消费", "一次；test_used_for_selection=false"],
        ["数据库/发布副作用", "database_writes=0；forecast_runs_created=0；promotion=0；active_pointer_changes=0"],
    ]
    add_table(doc, ["证据", "值"], audit_rows, [2600, 6760], font_size=8.8)
    add_para(doc, "本报告与图表由本轮真实 CSV/JSON 工件生成。完整逐周预测、模型搜索表、锁文件与 MANIFEST.sha256 位于随报告同目录的 evidence/ 下。", size=9.5, color=MID_GRAY)

    doc.add_page_break()
    doc.add_heading("附录 A：13 周逐行对账", level=1)
    appendix_rows = []
    for r in paired:
        appendix_rows.append([
            f"2015-W{int(r['week']):02d}",
            f"{float(r['actual']):.0f}",
            f"{float(r['EWNet_author_code']):.1f}",
            f"{float(r['kg_validation_only_convex_ensemble.v2']):.1f}",
            f"{float(r['last_value']):.1f}",
        ])
    add_table(doc, ["目标周", "实际", "论文 EWNet", "项目优化组合", "Last-value"], appendix_rows, [1800, 1400, 2100, 2300, 1760], font_size=8.8)
    add_para(doc, "说明：以上 13 行均来自 outputs/test_paired_predictions.csv；没有补零、插值或后验修正。", size=9, color=MID_GRAY)

    doc.add_heading("附录 B：参考来源", level=1)
    refs = [
        ("论文正式页", "https://www.sciencedirect.com/science/article/abs/pii/S0893608023002939"),
        ("作者代码与数据", "https://github.com/mad-stat/Epicasting"),
        ("2024 JCR/SCIE 分区清单", "https://uefiscdi.gov.ro/resource-860007-clasament.ais.jcr.2024.pdf"),
    ]
    for idx, (label, url) in enumerate(refs, 1):
        p = add_para(doc, f"{idx}. {label}：", size=10, after=4)
        add_hyperlink(p, url, url)

    core_props = doc.core_properties
    core_props.title = "SCI 一区论文 EWNet 对比试验优化结果报告"
    core_props.subject = "疫情自动获取系统：同数据、同切分、无泄漏对比试验"
    core_props.author = "Codex"
    core_props.keywords = "EWNet; SCI Q1; dengue; forecasting; WAPE; reproducibility"
    core_props.comments = "Generated from verified experiment artifacts; negative generalization result retained."

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)
    manifest_lines = []
    for path in sorted((REPORT_DIR / "evidence").rglob("*")):
        if path.is_file():
            manifest_lines.append(f"{sha256_file(path)}  {path.relative_to(REPORT_DIR).as_posix()}")
    for path in sorted(OUTPUT_DIR.glob("*.png")):
        manifest_lines.append(f"{sha256_file(path)}  {path.relative_to(REPORT_DIR).as_posix()}")
    manifest_lines.append(f"{sha256_file(DOCX_PATH)}  {DOCX_PATH.name}")
    (REPORT_DIR / "MANIFEST.sha256").write_text("\n".join(manifest_lines) + "\n", encoding="utf-8")
    print(DOCX_PATH)


if __name__ == "__main__":
    build_docx()
