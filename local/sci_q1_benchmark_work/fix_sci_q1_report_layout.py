from pathlib import Path

from docx import Document


DOCX = Path(r"F:\作业\aibola\reports\sci_q1_ewnet_benchmark_optimization_20260822\SCI一区论文EWNet对比试验优化结果报告_20260822.docx")


doc = Document(DOCX)

# Keep the compact task-status table (third table) together instead of
# orphaning its final row on the next page.
status_table = doc.tables[2]
all_status_paragraphs = [p for row in status_table.rows for cell in row.cells for p in cell.paragraphs]
for paragraph in all_status_paragraphs[:-1]:
    paragraph.paragraph_format.keep_with_next = True

# The audit heading must travel with its table; forcing it to a fresh page is
# clearer than leaving a lone table header at the foot of the prior page.
for paragraph in doc.paragraphs:
    if paragraph.text.strip() == "9. 审计与可复现性":
        paragraph.paragraph_format.page_break_before = True
        paragraph.paragraph_format.keep_with_next = True
        break

doc.save(DOCX)
