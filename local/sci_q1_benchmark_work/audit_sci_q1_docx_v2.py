from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn


path = Path(r"F:\作业\aibola\reports\sci_q1_ewnet_benchmark_optimization_20260822\SCI一区论文EWNet对比试验优化结果报告_20260822.docx")
doc = Document(path)
section = doc.sections[0]
assert section.page_width == 7772400
assert section.page_height == 10058400
assert section.top_margin == 914400
assert section.bottom_margin == 914400
assert section.left_margin == 914400
assert section.right_margin == 914400
assert section.header_distance == 449580
assert section.footer_distance == 449580

for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
    style = doc.styles[style_name]
    assert style._element.rPr.rFonts.get(qn("w:eastAsia")) == "Microsoft YaHei"

for table_index, table in enumerate(doc.tables):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    assert tbl_w is not None and tbl_w.get(qn("w:w")) == "9360", table_index
    assert tbl_ind is not None and tbl_ind.get(qn("w:w")) == "120", table_index
    assert tbl_layout is not None and tbl_layout.get(qn("w:type")) == "fixed", table_index
    grid_sum = sum(int(col.get(qn("w:w"))) for col in table._tbl.tblGrid)
    assert grid_sum == 9360, (table_index, grid_sum)

with ZipFile(path) as zf:
    xml = zf.read("word/document.xml").decode("utf-8")
    assert "{{" not in xml and "}}" not in xml
    assert "PLACEHOLDER" not in xml

print(f"DOCX_AUDIT_PASS tables={len(doc.tables)} sections={len(doc.sections)}")
