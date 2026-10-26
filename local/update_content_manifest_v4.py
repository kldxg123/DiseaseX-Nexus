from __future__ import annotations

from pathlib import Path

from docx import Document


ROOT = Path(r"F:\作业\aibola\reports\nc_submission_revision_20260824\NC_Submission_Updated_20260824")
SOURCE = ROOT / "Content_Manifest_Updated_v3_20260824.docx"
TARGET = ROOT / "Content_Manifest_Updated_v4_20260824.docx"

REPLACEMENTS = {
    "Updated 2026-08-24 | Page-21 prose normalization, Python-generated figures and authenticated English interface evidence":
        "Updated 2026-08-24 | Original Figs. 1 and 3 retained; Fig. 6 and Supplementary Figs. 1–4 code-generated in English",
    "Main Figs. 1, 3 and 6 are now Python-generated; Main Fig. 2 remains unchanged by instruction.":
        "Main Figs. 1 and 3 retain the original submission artwork. Main Fig. 2 remains unchanged. Main Fig. 6 and Supplementary Figs. 1–4 are code-generated; all Supplementary Fig. 1–4 titles, legends and axes are English.",
    "DiseaseX_Nexus_NC_Manuscript_v5_20260824.docx/.pdf":
        "DiseaseX_Nexus_NC_Manuscript_v6_20260824.docx/.pdf",
    "Supplementary/Supplementary_Information_Updated_v2_20260824.docx/.pdf":
        "Supplementary/Supplementary_Information_Updated_v4_20260824.docx/.pdf",
    "Supplementary/Python_Generated_Figures_20260824/":
        "Supplementary/SI_Figures/ and SUPPLEMENTARY_ENGLISH_FIGURE_EMBEDDING_V4.json",
    "Python source, authenticated Brazil dengue export and deterministic Main Figs. 1 and 3":
        "English Supplementary Figs. 1–4, retained original Main Figs. 1 and 3, and byte-identical DOCX embedding attestation",
}


def replace_text(paragraph) -> int:
    original = paragraph.text
    updated = original
    for old, new in REPLACEMENTS.items():
        updated = updated.replace(old, new)
    if updated == original:
        return 0
    paragraph.text = updated
    return 1


document = Document(SOURCE)
changes = sum(replace_text(paragraph) for paragraph in document.paragraphs)
for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            changes += sum(replace_text(paragraph) for paragraph in cell.paragraphs)
document.save(TARGET)
print(f"changes={changes}")
