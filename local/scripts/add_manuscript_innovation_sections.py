from __future__ import annotations

import argparse
import hashlib
import shutil
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from lxml import etree


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}


INTRO_TEXT = (
    "To our knowledge, DiseaseX Nexus is the first implemented framework to organize "
    "epidemic intelligence as a proactive disease-control decision-support loop rather "
    "than a static monitoring pipeline. Its novelty is fourfold. First, the framework "
    "couples dynamic sensing with a feedback-driven active-learning process, allowing "
    "new observations, matured outcomes and recorded failures to revise model "
    "recommendations as surveillance conditions change. This design enables rapid "
    "adaptation across pathogens without assuming that one forecasting model is "
    "universally transferable. Second, it introduces an architecture that links an LLM "
    "with deep-learning models: the LLM retrieves evidence, maps data characteristics "
    "and coordinates tools, while inspectable numerical models perform training and "
    "prediction. This division of labour establishes a new clinician-model working "
    "relationship at the analytic workstation: clinicians retain authority and review, "
    "whereas models continuously monitor signals, propose analyses and expose their "
    "evidence and uncertainty. Third, when a previously unseen Disease X emerges, the "
    "system can locate it within a disease and pathogen tree, traverse the knowledge "
    "graph to identify analogous transmission, surveillance and feature profiles, "
    "retrieve relevant forecasting experience and assemble candidate methods for "
    "controlled evaluation. Fourth, governance identities and execution contracts make "
    "each transition from detection to recommendation and prediction auditable. "
    "Together, these contributions operationalize active learning, dynamic sensing and "
    "intelligent prediction for unknown pathogens while preserving human oversight."
)


DISCUSSION_TEXT = (
    "The central innovation should therefore be interpreted at the architectural level "
    "rather than as a claim of universal predictive superiority. Static surveillance "
    "systems mainly summarize a fixed set of indicators after data arrive; DiseaseX "
    "Nexus instead links dynamic sensing, feedback-driven active learning and governed "
    "model execution in a continuously revisable loop. The LLM and deep-learning "
    "components have complementary responsibilities: the LLM translates an evolving "
    "data profile into evidence-backed candidate tasks, whereas deterministic or deep "
    "numerical models generate forecasts that can be inspected and challenged. At the "
    "analytic workstation, this creates a clinician-model relationship based on shared "
    "control rather than automation: the system maintains continuous vigilance and "
    "proposes evidence-linked actions, while clinicians authorize execution, evaluate "
    "plausibility and determine public-health use. For an unseen Disease X, placement in "
    "the pathogen and disease tree can support retrieval of analogous features, "
    "transferable forecasting experience and candidate methods from the knowledge graph. "
    "Such analogical transfer is a starting hypothesis, not proof of biological "
    "equivalence; every proposed feature and method must still pass temporal-availability, "
    "local-validity and prospective-evaluation gates. The contribution is thus a "
    "governed mechanism for converting partial prior knowledge into testable predictions "
    "more rapidly, while keeping uncertainty and human accountability explicit."
)


INTRO_ANCHOR = "Here, we evaluate the claims that the current implementation can support."
DISCUSSION_ANCHOR = "Data heterogeneity remains the principal obstacle."


def paragraph_text(paragraph: etree._Element) -> str:
    return "".join(paragraph.xpath(".//w:t/text()", namespaces=NS))


def make_paragraph_like(anchor: etree._Element, text: str) -> etree._Element:
    paragraph = etree.Element(f"{{{W_NS}}}p")
    paragraph_properties = anchor.find(f"{{{W_NS}}}pPr")
    if paragraph_properties is not None:
        paragraph.append(deepcopy(paragraph_properties))

    run = etree.SubElement(paragraph, f"{{{W_NS}}}r")
    anchor_run = anchor.find(f"{{{W_NS}}}r")
    if anchor_run is not None:
        run_properties = anchor_run.find(f"{{{W_NS}}}rPr")
        if run_properties is not None:
            run.append(deepcopy(run_properties))

    text_node = etree.SubElement(run, f"{{{W_NS}}}t")
    text_node.text = text
    return paragraph


def insert_after_unique_anchor(
    body: etree._Element,
    anchor_prefix: str,
    text: str,
) -> None:
    paragraphs = body.xpath("./w:p", namespaces=NS)
    if any(paragraph_text(paragraph) == text for paragraph in paragraphs):
        raise ValueError("The requested paragraph already exists in the source document.")

    matches = [
        paragraph
        for paragraph in paragraphs
        if paragraph_text(paragraph).startswith(anchor_prefix)
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Expected one anchor beginning with {anchor_prefix!r}; found {len(matches)}."
        )
    matches[0].addnext(make_paragraph_like(matches[0], text))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def media_hashes(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        return {
            info.filename: hashlib.sha256(archive.read(info.filename)).hexdigest()
            for info in archive.infolist()
            if info.filename.startswith("word/media/")
        }


def edit_docx(source: Path, output: Path) -> None:
    if source.resolve() == output.resolve():
        raise ValueError("The output must not overwrite the source document.")
    if not source.is_file():
        raise FileNotFoundError(source)

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(output.suffix + ".tmp")
    if temporary.exists():
        temporary.unlink()

    with zipfile.ZipFile(source, "r") as source_archive:
        document_xml = source_archive.read("word/document.xml")
        root = etree.fromstring(document_xml)
        body = root.find(f"{{{W_NS}}}body")
        if body is None:
            raise ValueError("The Word document body was not found.")

        insert_after_unique_anchor(body, INTRO_ANCHOR, INTRO_TEXT)
        insert_after_unique_anchor(body, DISCUSSION_ANCHOR, DISCUSSION_TEXT)

        revised_xml = etree.tostring(
            root,
            encoding="UTF-8",
            xml_declaration=True,
            standalone=True,
        )

        with zipfile.ZipFile(temporary, "w") as output_archive:
            for info in source_archive.infolist():
                data = (
                    revised_xml
                    if info.filename == "word/document.xml"
                    else source_archive.read(info.filename)
                )
                output_archive.writestr(info, data)

    shutil.move(temporary, output)

    with zipfile.ZipFile(output) as archive:
        bad_member = archive.testzip()
        if bad_member is not None:
            raise ValueError(f"Corrupt ZIP member in output: {bad_member}")

    document = Document(output)
    all_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    if all_text.count(INTRO_TEXT) != 1 or all_text.count(DISCUSSION_TEXT) != 1:
        raise ValueError("The revised paragraphs were not found exactly once after save.")
    if media_hashes(source) != media_hashes(output):
        raise ValueError("Embedded media changed during the minimal OOXML edit.")

    print(f"source_sha256={sha256(source)}")
    print(f"output_sha256={sha256(output)}")
    print(f"output={output}")
    print(f"paragraph_count={len(document.paragraphs)}")
    print(f"media_count={len(media_hashes(output))}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    edit_docx(args.source, args.output)


if __name__ == "__main__":
    main()
