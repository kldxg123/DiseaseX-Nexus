from __future__ import annotations

import argparse
import hashlib
import shutil
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from lxml import etree


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}


INTRO_ANCHOR = (
    "To our knowledge, DiseaseX Nexus is the first implemented framework to organize"
)
DISCUSSION_ANCHOR = (
    "The central innovation should therefore be interpreted at the architectural level"
)


INTRO_PLACEHOLDER = (
    "[PLACEHOLDER FOR INTRODUCTION SYNTHESIS: Add a concise concluding paragraph that "
    "elevates these innovations into a broader research contribution and summarizes how "
    "active learning, dynamic sensing and knowledge-guided adaptation can transform "
    "multi-pathogen epidemic intelligence, including the response to a previously unseen "
    "Disease X.]"
)


DISCUSSION_PLACEHOLDER = (
    "[PLACEHOLDER FOR DISCUSSION SYNTHESIS: Add a concise concluding paragraph that "
    "synthesizes the scientific and public-health significance of the framework, clarifies "
    "the clinician-model working relationship and states the long-term vision for "
    "trustworthy intelligent prediction of unknown pathogens without extending claims "
    "beyond the evidence reported in this study.]"
)


def paragraph_text(paragraph: etree._Element) -> str:
    return "".join(paragraph.xpath(".//w:t/text()", namespaces=NS))


def highlighted_paragraph_like(
    anchor: etree._Element,
    text: str,
) -> etree._Element:
    paragraph = etree.Element(f"{{{W_NS}}}p")
    paragraph_properties = anchor.find(f"{{{W_NS}}}pPr")
    if paragraph_properties is not None:
        paragraph.append(deepcopy(paragraph_properties))

    run = etree.SubElement(paragraph, f"{{{W_NS}}}r")
    anchor_run = anchor.find(f"{{{W_NS}}}r")
    if anchor_run is not None:
        anchor_run_properties = anchor_run.find(f"{{{W_NS}}}rPr")
        if anchor_run_properties is not None:
            run_properties = deepcopy(anchor_run_properties)
            for old_highlight in run_properties.findall(f"{{{W_NS}}}highlight"):
                run_properties.remove(old_highlight)
            run.append(run_properties)
        else:
            run_properties = etree.SubElement(run, f"{{{W_NS}}}rPr")
    else:
        run_properties = etree.SubElement(run, f"{{{W_NS}}}rPr")

    if run.find(f"{{{W_NS}}}rPr") is None:
        run_properties = etree.SubElement(run, f"{{{W_NS}}}rPr")
    else:
        run_properties = run.find(f"{{{W_NS}}}rPr")
    highlight = etree.SubElement(run_properties, f"{{{W_NS}}}highlight")
    highlight.set(f"{{{W_NS}}}val", "yellow")

    text_node = etree.SubElement(run, f"{{{W_NS}}}t")
    text_node.text = text
    return paragraph


def insert_after_unique_anchor(
    body: etree._Element,
    anchor_prefix: str,
    placeholder: str,
) -> None:
    paragraphs = body.xpath("./w:p", namespaces=NS)
    if any(paragraph_text(paragraph) == placeholder for paragraph in paragraphs):
        raise ValueError("The requested placeholder already exists in the source.")
    matches = [
        paragraph
        for paragraph in paragraphs
        if paragraph_text(paragraph).startswith(anchor_prefix)
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Expected one anchor beginning with {anchor_prefix!r}; found {len(matches)}."
        )
    matches[0].addnext(highlighted_paragraph_like(matches[0], placeholder))


def media_hashes(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        return {
            info.filename: hashlib.sha256(archive.read(info.filename)).hexdigest()
            for info in archive.infolist()
            if info.filename.startswith("word/media/")
        }


def validate_highlighted_placeholder(
    body: etree._Element,
    placeholder: str,
) -> None:
    matches = [
        paragraph
        for paragraph in body.xpath("./w:p", namespaces=NS)
        if paragraph_text(paragraph) == placeholder
    ]
    if len(matches) != 1:
        raise ValueError(f"Expected one placeholder paragraph; found {len(matches)}.")
    highlighted_runs = matches[0].xpath(
        ".//w:r[w:rPr/w:highlight[@w:val='yellow']]",
        namespaces=NS,
    )
    if len(highlighted_runs) != 1:
        raise ValueError("The placeholder is not contained in one yellow-highlighted run.")


def edit_docx(source: Path, output: Path) -> None:
    if source.resolve() == output.resolve():
        raise ValueError("The output must not overwrite the source document.")
    if not source.is_file():
        raise FileNotFoundError(source)

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(output.suffix + ".tmp")
    if temporary.exists():
        temporary.unlink()

    with zipfile.ZipFile(source, "r") as source_archive:
        root = etree.fromstring(source_archive.read("word/document.xml"))
        body = root.find(f"{{{W_NS}}}body")
        if body is None:
            raise ValueError("The Word document body was not found.")

        insert_after_unique_anchor(body, INTRO_ANCHOR, INTRO_PLACEHOLDER)
        insert_after_unique_anchor(
            body,
            DISCUSSION_ANCHOR,
            DISCUSSION_PLACEHOLDER,
        )
        validate_highlighted_placeholder(body, INTRO_PLACEHOLDER)
        validate_highlighted_placeholder(body, DISCUSSION_PLACEHOLDER)

        revised_xml = etree.tostring(
            root,
            encoding="UTF-8",
            xml_declaration=True,
            standalone=True,
        )
        with zipfile.ZipFile(temporary, "w") as output_archive:
            for info in source_archive.infolist():
                data = (
                    revised_xml
                    if info.filename == "word/document.xml"
                    else source_archive.read(info.filename)
                )
                output_archive.writestr(info, data)

    shutil.move(temporary, output)

    with zipfile.ZipFile(output) as archive:
        bad_member = archive.testzip()
        if bad_member is not None:
            raise ValueError(f"Corrupt ZIP member in output: {bad_member}")
        root = etree.fromstring(archive.read("word/document.xml"))
        body = root.find(f"{{{W_NS}}}body")
        if body is None:
            raise ValueError("The saved Word document body was not found.")
        validate_highlighted_placeholder(body, INTRO_PLACEHOLDER)
        validate_highlighted_placeholder(body, DISCUSSION_PLACEHOLDER)

    document = Document(output)
    all_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    if all_text.count(INTRO_PLACEHOLDER) != 1:
        raise ValueError("The Introduction placeholder was not saved exactly once.")
    if all_text.count(DISCUSSION_PLACEHOLDER) != 1:
        raise ValueError("The Discussion placeholder was not saved exactly once.")
    if media_hashes(source) != media_hashes(output):
        raise ValueError("Embedded media changed during the minimal OOXML edit.")

    print(f"output={output}")
    print(f"paragraph_count={len(document.paragraphs)}")
    print(f"media_count={len(media_hashes(output))}")
    print("yellow_placeholders=2")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    edit_docx(args.source, args.output)


if __name__ == "__main__":
    main()
