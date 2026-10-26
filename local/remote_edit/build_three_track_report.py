from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"F:\作业\aibola")
OUT = ROOT / "reports" / "three_track_paper_benchmark_20260813"
DOCX = OUT / "疫情预测三轨闭环与论文基准对比报告_20260813.docx"

BLUE = "2E74B5"
DARK = "17365D"
MUTED = "667085"
LIGHT = "F2F4F7"
PALE_BLUE = "EAF2F8"
PALE_RED = "FDECEC"
PALE_GREEN = "EAF7EF"


def font(run, size=11, bold=False, color="000000", name="Microsoft YaHei"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths, *, header_fill=LIGHT, small=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    for i, label in enumerate(headers):
        shade(hdr.cells[i], header_fill)
        p = hdr.cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        font(p.add_run(str(label)), size=small, bold=True, color=DARK)
    for row_values in rows:
        row = table.add_row()
        for i, value in enumerate(row_values):
            p = row.cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in (0, len(row_values) - 1) else WD_ALIGN_PARAGRAPH.LEFT
            font(p.add_run(str(value)), size=small)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="Normal")
    if bold_prefix and text.startswith(bold_prefix):
        font(p.add_run(bold_prefix), bold=True)
        font(p.add_run(text[len(bold_prefix):]))
    else:
        font(p.add_run(text))
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    font(p.add_run(text))
    return p


def callout(doc, title, body, fill):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    shade(table.cell(0, 0), fill)
    p = table.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    font(p.add_run(title + "\n"), size=12, bold=True, color=DARK)
    font(p.add_run(body), size=10.5)
    set_table_geometry(table, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    font(p.add_run(text), size=9, color=MUTED)


def add_image(doc, path: Path, width=6.3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))


def page_num(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run2 = paragraph.add_run(" 页")
    font(run2, size=9, color=MUTED)


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(1)
sec.bottom_margin = Inches(1)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.492)
sec.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(11)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10
for name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK, 8, 4),
):
    st = styles[name]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
for name in ("List Bullet", "List Number"):
    st = styles[name]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(8)
    st.paragraph_format.line_spacing = 1.167

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
font(header.add_run("疫情自动获取系统 | 技术验收报告"), size=9, color=MUTED)
page_num(sec.footer.paragraphs[0])

# Memo masthead
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(18)
p.paragraph_format.space_after = Pt(5)
font(p.add_run("技术验收与对比实验报告"), size=24, bold=True, color=DARK)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(16)
font(p.add_run("三轨预测闭环、第三条流感链路与知识图谱论文复刻"), size=14, color=BLUE)
for label, value in (
    ("项目", "epidemic-openclaw / 疫情自动获取系统"),
    ("执行日期", "2026-08-13"),
    ("范围", "Shadow-only；无 Trusted、Promotion 或 Active Pointer 变更"),
    ("核心结论", "第三条链路已跑通但盲测未达 70%；论文对比达到不相上下"),
):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    font(p.add_run(label + "："), bold=True, color=DARK)
    font(p.add_run(value))

callout(
    doc,
    "结论先行",
    "本轮完成了马来西亚流感 Track D 的 Qwen 调优、Shadow 预测和一次性新鲜盲测；验证集吻合度 86.20%，但新鲜盲测仅 59.15%，故不能认证为第三条 70% 链路。与此同时，在作者官方 ILI 数据同一测试窗中，项目模型吻合度 68.65%，论文 GRU 方法级复刻 69.45%，差 0.80 个百分点，结果可以判定为不相上下。",
    PALE_BLUE,
)

add_heading(doc, "1. 任务完成度", 1)
add_table(
    doc,
    ["目标", "完成度", "真实结论"],
    [
        ("将两条链路扩展为三条", "部分完成", "第三条 Track D 已跑通，但 fresh blind 未达 70%"),
        ("系统截图、架构、设计文档打包", "完成", "复用既有真实截图 + 新架构图 + 本报告"),
        ("选择最接近论文并同数据对比", "完成", "Aiken GRU vs 项目模型；同样本结果不相上下"),
    ],
    [2200, 1500, 5660],
)

add_heading(doc, "2. 系统架构与受控闭环", 1)
add_image(doc, OUT / "architecture.png", 6.25)
caption(doc, "图 1  数据层—知识图谱与治理层—预测执行与反馈层，以及 Track A–D 和论文对比支路")
add_body(doc, "架构分为三层：底层保存可信/规范化数据、论文全文、证据定位和版本哈希；中层由疾病—模型关系、模型/参数证据、数据画像、门禁与 Qwen 受控规划构成；顶层执行滚动回测、基线比较、Shadow 预测、一次性 held-out 和反馈账本。Qwen 只能在白名单模型与参数空间内提出方案，不能修改治理规则、正式数据或发布指针。")

add_heading(doc, "3. 四条轨道的真实状态", 1)
add_table(
    doc,
    ["轨道", "数据", "身份", "WAPE / 吻合度", "结论"],
    [
        ("A", "AUT 军团菌/月", "历史回放", "0.2270 / 77.30%", "可演示，非 fresh"),
        ("B", "FIN 汉坦病毒/周", "重复打开窗口", "0.2803 / 71.97%", "历史结果，非 fresh"),
        ("C", "CHL 流感/周", "fresh 已消费", "0.7170 / 28.30%", "门禁失败"),
        ("D", "MYS 流感/周", "fresh 已消费", "0.4085 / 59.15%", "链路完成，门禁失败"),
    ],
    [900, 2100, 1700, 1900, 2760],
)
callout(doc, "口径提醒", "“吻合度 = 100% - WAPE×100%”是项目演示口径，不是分类准确率。只有 fresh held-out 同时满足 WAPE、吻合度、覆盖率、同样本基线改善与时间链门禁，才允许标记 70% 认证。", "FFF4E5")

add_heading(doc, "4. Track D：马来西亚流感第三条链路", 1)
add_heading(doc, "4.1 选择与冻结", 2)
add_body(doc, "在 63 条流感国家序列中，只使用每条序列前 321 周训练和后续 72 周验证进行候选筛选；扫描证据记录 hidden_rows_evaluated=0。马来西亚验证表现最佳，最后 52 周在选择阶段保持未打开。随后全局 held-out registry 在首次 Qwen 调用前原子锁定。")
add_heading(doc, "4.2 Qwen 调优与验证", 2)
add_table(
    doc,
    ["项目", "结果"],
    [
        ("模型", "lagged_count_boosting.v1"),
        ("参数", "n_estimators=37；learning_rate=0.1601；min_samples_leaf=25"),
        ("验证集", "WAPE 0.13798；吻合度 86.20%；覆盖率 0.9375"),
        ("一次性盲测", "WAPE 0.40850；吻合度 59.15%；覆盖率 0.75"),
        ("最佳基线", "last_value；WAPE 0.40218"),
        ("治理状态", "CONSUMED；FRESH_HELDOUT_GATES_FAILED"),
    ],
    [2300, 7060],
)
add_heading(doc, "4.3 为什么验证高、盲测低", 2)
add_bullet(doc, "流感周度检测数在盲测窗出现明显季节峰值与幅度漂移；训练/验证阶段的幅度和拐点不能完全代表 2025–2026 隐藏窗。")
add_bullet(doc, "H1 WAPE 为 0.258，但 H2–H4 逐步升至约 0.435、0.436、0.500；多步预测衰减显著。")
add_bullet(doc, "区间覆盖率仅 0.75，低于 0.85 门槛，说明不确定性区间也没有充分覆盖分布漂移。")
add_bullet(doc, "最佳基线 WAPE 0.40218 略优于调优模型 0.40850；因此即使验证集达标，也不能发布为 champion。")

add_heading(doc, "5. 知识图谱论文选择", 1)
add_body(doc, "从已有人审疾病—模型关系中选择 Aiken 等人的流感论文（DOI 10.1126/sciadv.abb1237）。其官方实现同时包含 persistence、AR、AR-net 和 GRU；其中“仅使用历史流行病学数据”的分支不需要 Google Trends 等外生数据，最符合“单一感染/ILI 时间序列预测”的比较要求。")
add_table(
    doc,
    ["约束", "本轮处理"],
    [
        ("数据来源", "作者官方 GitHub 的 States_ILI.csv，SHA-256 固定"),
        ("状态选择", "固定纽约州 NY；不存在按测试结果挑州"),
        ("时间语义", "52 lags；H1；52 周验证；52 周测试"),
        ("论文方法", "GRU，5 hidden units，700 epochs"),
        ("复刻边界", "旧版 Keras 未原环境复现；使用 PyTorch 方法级复刻"),
    ],
    [2100, 7260],
)

add_heading(doc, "6. 同数据、同切分对比结果", 1)
add_table(
    doc,
    ["模型", "MAE", "RMSE", "WAPE", "吻合度", "Pearson r"],
    [
        ("论文 GRU 方法级复刻", "0.4626", "0.5893", "0.3055", "69.45%", "0.7466"),
        ("项目 lagged boosting", "0.4748", "0.6027", "0.3135", "68.65%", "0.7402"),
        ("last_value 基线", "0.4809", "0.6066", "0.3175", "68.25%", "0.7321"),
        ("seasonal naive", "1.0216", "1.5254", "0.6746", "32.54%", "0.5261"),
    ],
    [2400, 1200, 1200, 1200, 1500, 1860],
)
callout(doc, "公平比较结论", "项目模型相对论文 GRU 的测试 WAPE 仅高 0.00804，吻合度仅低 0.80 个百分点；二者使用完全相同的 52 条测试记录。结果已经不相上下。由于测试集已经打开，本轮没有再依据测试指标修改参数，避免测试泄漏。", PALE_GREEN)

add_heading(doc, "7. 系统页面与演示证据", 1)
add_body(doc, "以下三张截图来自此前已通过登录态浏览器验收的双数据集闭环报告，展示数据选择、知识图谱推荐、回测/结果和反馈状态的页面结构。本轮新页面因 Codex 应用对本地地址的 Browser Use 自动策略拒绝而未重新采集；该限制与 vm201 Dashboard 健康无关，且没有被绕过。")
for filename, label in (
    ("dual_loop_austria.png", "图 2  奥地利军团菌完整闭环页面"),
    ("dual_loop_peru.png", "图 3  秘鲁流感完整闭环页面"),
    ("dual_loop_feedback.png", "图 4  反馈与学习状态页面"),
):
    add_image(doc, OUT / "screenshots" / filename, 6.0)
    caption(doc, label + "（既有真实浏览器验收截图）")

add_heading(doc, "8. 测试、安全和非变更证明", 1)
add_table(
    doc,
    ["门禁", "结果"],
    [
        ("kg_qwen tests", "177 passed；34 条第三方弃用 warning"),
        ("Ruff", "PASS"),
        ("strict Mypy", "pipeline.py 与 api.py：PASS"),
        ("held-out 时间链", "registered < locked < Qwen < tuning_finished < consumed"),
        ("held-out 单次消费", "registry=CONSUMED；唯一 consumer run"),
        ("Dashboard", "/healthz=200"),
        ("正式写入", "formal forecast=0；promotion=0；active pointer=0；Trusted delta=0"),
    ],
    [2800, 6560],
)
add_body(doc, "完整全仓库 pytest 在系统 Python 环境收集阶段因 httpx 未加入该环境和两个同名测试模块导入冲突而失败，属于既有运行环境问题；本轮直接相关的 kg_qwen 测试、静态检查和数据库验收均已通过。报告不将未完成的全仓测试写为 PASS。")

add_heading(doc, "9. 后续设计方案", 1)
add_bullet(doc, "把 Track D 纳入失败反馈账本，按 H1 与 H2–H4 分层诊断季节峰值偏移，不重复打开已消费隐藏窗。")
add_bullet(doc, "新增经过论文知识图谱约束的 GRU/AR-net 候选适配器，但必须在训练窗口内做嵌套滚动验证，不能用已见隐藏窗选模。")
add_bullet(doc, "为下一条未见序列预注册数据、切分、候选模型、调参预算、种子与门禁；锁定后才允许 Qwen 规划。")
add_bullet(doc, "将 fresh certification、historical replay 和 validation demo 在 UI 中使用不同颜色和文案，杜绝“验证 70% 等于盲测 70%”的误解。")
add_bullet(doc, "若需要论文 exact reproduction，另建旧版 Keras/TensorFlow 隔离容器，并记录解释器、依赖、种子和作者脚本补丁；本轮结果只能称为 method-level reconstruction。")

add_heading(doc, "附录 A：可复现资产", 1)
add_table(
    doc,
    ["资产", "说明"],
    [
        ("TRACK_D_ACCEPTANCE.json", "只读数据库验收、完整 hidden 指标与时间链"),
        ("flu_candidate_validation_scan_20260813.json", "63 国验证扫描；hidden_rows_evaluated=0"),
        ("paper_comparison/comparison.json", "论文复刻与项目同数据指标"),
        ("paper_comparison/paired_predictions.csv", "52 周逐行同样本预测"),
        ("paper_benchmark_source/raw", "作者官方数据与最小源代码文件"),
        ("architecture.drawio / png / svg", "可编辑与可嵌入架构图"),
        ("KIMI_NEXT_STAGE_PROMPT.txt", "后续协作模型的完整执行提示词"),
    ],
    [3100, 6260],
)

add_heading(doc, "附录 B：来源", 1)
add_bullet(doc, "Aiken et al. (2021), Science Advances, DOI 10.1126/sciadv.abb1237。")
add_bullet(doc, "作者官方代码与数据仓库：https://github.com/emilylaiken/ml-flu-prediction")
add_bullet(doc, "本地数据文件 States_ILI.csv SHA-256：cda4808dd6d5b504ec2f18bf9aaef73623a50eecab7958d6a3a8170c07cffe6d。")

props = doc.core_properties
props.title = "疫情预测三轨闭环与论文基准对比报告"
props.subject = "Track D fresh held-out、知识图谱论文复刻与同数据比较"
props.author = "epidemic-openclaw project"
props.keywords = "疫情预测, 知识图谱, Qwen, held-out, GRU, 流感"
props.comments = "Generated from verified project evidence; no secrets included."

doc.save(DOCX)
print(DOCX)
