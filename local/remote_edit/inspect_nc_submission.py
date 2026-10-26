from __future__ import annotations

import csv
import hashlib
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from PIL import Image


ROOT = Path(r"F:\作业\aibola\reports\nc_submission_revision_20260823")
UNPACKED = ROOT / "source_unpacked"
OUT = ROOT / "inspection_inventory.json"
TEXT_OUT = ROOT / "inspection_fulltext.txt"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def docx_info(path: Path) -> dict[str, object]:
    doc = Document(path)
    paragraphs = []
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text:
            paragraphs.append(
                {
                    "index": index,
                    "style": paragraph.style.name if paragraph.style else None,
                    "text": text,
                }
            )
    tables = []
    for table_index, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append({"index": table_index, "rows": rows})

    media = []
    with zipfile.ZipFile(path) as archive:
        for member in sorted(n for n in archive.namelist() if n.startswith("word/media/")):
            data = archive.read(member)
            item: dict[str, object] = {
                "member": member,
                "bytes": len(data),
                "sha256": hashlib.sha256(data).hexdigest(),
            }
            try:
                from io import BytesIO

                with Image.open(BytesIO(data)) as image:
                    item["width_px"], item["height_px"] = image.size
                    item["format"] = image.format
            except Exception as exc:  # image inventory must remain best effort
                item["image_error"] = type(exc).__name__
            media.append(item)

    all_text = "\n".join(p["text"] for p in paragraphs)
    all_text += "\n" + "\n".join(" | ".join(r) for t in tables for r in t["rows"])
    keywords = [
        "ar(2)",
        "2407.00332",
        "rsos.240248",
        "s41467-024-52504-1",
        "awbe",
        "adaptive weight",
        "singapore",
        "hong kong",
        "2012",
        "2019",
        "2023",
        "wape",
        "88.66",
        "87.20",
        "77.69",
    ]
    lower = all_text.lower()
    hits = {keyword: len(re.findall(re.escape(keyword), lower)) for keyword in keywords}
    return {
        "path": str(path),
        "bytes": path.stat().st_size,
        "sha256": sha256(path),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "inline_shape_count": len(doc.inline_shapes),
        "paragraphs": paragraphs,
        "tables": tables,
        "media": media,
        "keyword_hits": hits,
    }


def csv_info(path: Path) -> dict[str, object]:
    rows: list[list[str]] = []
    encoding = "utf-8-sig"
    try:
        with path.open("r", encoding=encoding, newline="") as handle:
            rows = list(csv.reader(handle))
    except UnicodeDecodeError:
        encoding = "gb18030"
        with path.open("r", encoding=encoding, newline="") as handle:
            rows = list(csv.reader(handle))
    text = "\n".join(",".join(row) for row in rows)
    return {
        "path": str(path),
        "bytes": path.stat().st_size,
        "sha256": sha256(path),
        "encoding": encoding,
        "row_count": len(rows),
        "column_count_max": max((len(row) for row in rows), default=0),
        "header": rows[0] if rows else [],
        "first_rows": rows[1:4],
        "last_rows": rows[-3:] if len(rows) > 1 else [],
        "keyword_hits": {
            keyword: text.lower().count(keyword)
            for keyword in [
                "ar(2)",
                "2407.00332",
                "rsos.240248",
                "awbe",
                "singapore",
                "hong kong",
                "wape",
            ]
        },
    }


def main() -> None:
    docx_files = sorted(UNPACKED.rglob("*.docx"))
    csv_files = sorted(UNPACKED.rglob("*.csv"))
    image_files = sorted(
        path
        for path in UNPACKED.rglob("*")
        if path.is_file() and path.suffix.lower() in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}
    )
    inventory = {
        "docx": [docx_info(path) for path in docx_files],
        "csv": [csv_info(path) for path in csv_files],
        "images": [],
    }
    for path in image_files:
        with Image.open(path) as image:
            inventory["images"].append(
                {
                    "path": str(path),
                    "bytes": path.stat().st_size,
                    "sha256": sha256(path),
                    "width_px": image.width,
                    "height_px": image.height,
                    "format": image.format,
                }
            )
    OUT.write_text(json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8")

    sections = []
    for item in inventory["docx"]:
        sections.append(f"===== {item['path']} =====")
        for paragraph in item["paragraphs"]:
            sections.append(
                f"P{paragraph['index']} [{paragraph['style']}] {paragraph['text']}"
            )
        for table in item["tables"]:
            sections.append(f"--- TABLE {table['index']} ---")
            for row in table["rows"]:
                sections.append(" | ".join(row))
    TEXT_OUT.write_text("\n".join(sections), encoding="utf-8")
    print(
        json.dumps(
            {
                "docx_count": len(docx_files),
                "csv_count": len(csv_files),
                "image_count": len(image_files),
                "inventory": str(OUT),
                "fulltext": str(TEXT_OUT),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
