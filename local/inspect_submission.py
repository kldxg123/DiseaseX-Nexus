from pathlib import Path
from docx import Document

ROOT = Path(r"F:\作业\aibola\reports\nc_submission_revision_20260824\NC_Submission_Updated_20260824")

def inspect_doc(path: Path, needles: list[str]) -> None:
    print(f"DOC={path}")
    doc = Document(path)
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if any(n.lower() in text.lower() for n in needles):
            print(f"P[{i}] style={p.style.name!r}: {text}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)}")

inspect_doc(
    ROOT / "DiseaseX_Nexus_NC_Manuscript_v4_20260824.docx",
    [
        "Hong Kong influenza author-output benchmark",
        "The benchmark used the official repository",
        "DiseaseX searched 450",
        "Results_point.zip",
        "b0bd4cc",
    ],
)

inspect_doc(
    ROOT / "Supplementary" / "Supplementary_Information_Updated_20260824.docx",
    [
        "b0bd4cc",
        "27c3003",
        "Results_point.zip",
        "selection-input",
        "selection-lock",
        "Hong Kong",
        "Supplementary Note 4",
    ],
)

for note in sorted((ROOT / "Supplementary").glob("SI_Note_*.docx")):
    inspect_doc(note, ["b0bd4cc", "27c3003", "Results_point.zip", "Hong Kong", "selection-input"])
