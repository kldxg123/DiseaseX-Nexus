from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import sys
import zipfile
from pathlib import Path, PurePosixPath

from docx import Document


ROOT = Path(r"F:\作业\aibola\reports\nc_manuscript_comparison_alignment_20260823")
DOCX = ROOT / "source" / "DiseaseX_Nexus_NC_Manuscript_v2_20260822.docx"
ZIP = ROOT / "source" / "NC投稿完善_20260822.zip"
EXTRACT = ROOT / "source" / "zip_extracted"
AUDIT = ROOT / "audit"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def safe_member(name: str) -> bool:
    p = PurePosixPath(name.replace("\\", "/"))
    return not p.is_absolute() and ".." not in p.parts and not re.match(r"^[A-Za-z]:", name)


def inspect_zip() -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    with zipfile.ZipFile(ZIP) as zf:
        for info in zf.infolist():
            mode = (info.external_attr >> 16) & 0o170000
            is_symlink = mode == 0o120000
            rows.append(
                {
                    "name": info.filename,
                    "size": info.file_size,
                    "compressed_size": info.compress_size,
                    "safe_path": safe_member(info.filename),
                    "is_symlink": is_symlink,
                }
            )
        bad = [r for r in rows if not r["safe_path"] or r["is_symlink"]]
        if bad:
            raise RuntimeError(f"Unsafe ZIP members: {bad[:10]}")
        if EXTRACT.exists():
            shutil.rmtree(EXTRACT)
        EXTRACT.mkdir(parents=True)
        zf.extractall(EXTRACT)
    return rows


def inspect_docx(path: Path, label: str) -> dict[str, object]:
    doc = Document(path)
    paragraphs = []
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:
            paragraphs.append({"index": i, "style": p.style.name if p.style else None, "text": text})
    tables = []
    for ti, table in enumerate(doc.tables):
        data = []
        for row in table.rows:
            data.append([cell.text.strip() for cell in row.cells])
        tables.append({"index": ti, "rows": data})
    media = []
    with zipfile.ZipFile(path) as zf:
        for name in sorted(n for n in zf.namelist() if n.startswith("word/media/")):
            raw = zf.read(name)
            media.append({"name": name, "size": len(raw), "sha256": hashlib.sha256(raw).hexdigest()})
    return {
        "label": label,
        "path": str(path),
        "sha256": sha256(path),
        "paragraphs": paragraphs,
        "tables": tables,
        "media": media,
        "inline_shapes": len(doc.inline_shapes),
        "section_count": len(doc.sections),
    }


def candidate_docs() -> list[Path]:
    return sorted(p for p in EXTRACT.rglob("*.docx") if not p.name.startswith("~$"))


def main() -> None:
    AUDIT.mkdir(parents=True, exist_ok=True)
    zip_rows = inspect_zip()
    docs = [inspect_docx(DOCX, "primary_manuscript")]
    for p in candidate_docs():
        try:
            docs.append(inspect_docx(p, f"zip:{p.relative_to(EXTRACT)}"))
        except Exception as exc:
            docs.append({"label": f"zip:{p.relative_to(EXTRACT)}", "error": repr(exc), "path": str(p)})

    files = []
    for p in sorted(EXTRACT.rglob("*")):
        if p.is_file():
            files.append(
                {
                    "path": str(p.relative_to(EXTRACT)),
                    "size": p.stat().st_size,
                    "sha256": sha256(p),
                    "suffix": p.suffix.lower(),
                }
            )

    payload = {
        "source_docx_sha256": sha256(DOCX),
        "source_zip_sha256": sha256(ZIP),
        "zip_member_count": len(zip_rows),
        "zip_members": zip_rows,
        "files": files,
        "documents": docs,
    }
    (AUDIT / "source_inventory.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    terms = re.compile(
        r"EWNet|Epicasting|WAPE|AR\(2\)|ARIMA|agreement|accuracy|准确率|吻合度|"
        r"benchmark|对比|Singapore|新加坡|dengue|登革热|2015|2019|2024|SCI|Q1|"
        r"lagged_count|last[-_ ]?value|seasonal|RMSE|MAE",
        re.IGNORECASE,
    )
    lines = ["# Source inventory", "", f"ZIP members: {len(zip_rows)}", f"Extracted files: {len(files)}", ""]
    for d in docs:
        lines.append(f"## {d.get('label')}")
        if "error" in d:
            lines.append(str(d["error"]))
            continue
        lines.append(f"Paragraphs: {len(d['paragraphs'])}; tables: {len(d['tables'])}; media: {len(d['media'])}")
        lines.append("")
        lines.append("### Matched paragraphs")
        for p in d["paragraphs"]:
            if terms.search(p["text"]):
                lines.append(f"- P{p['index']} [{p['style']}]: {p['text']}")
        lines.append("")
        lines.append("### Tables")
        for table in d["tables"]:
            flat = " | ".join(" || ".join(row) for row in table["rows"])
            if terms.search(flat):
                lines.append(f"- Table {table['index']}: {flat}")
        lines.append("")
    (AUDIT / "source_inventory.md").write_text("\n".join(lines), encoding="utf-8")
    print(json.dumps({"zip_members": len(zip_rows), "files": len(files), "docx_count": len(docs)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
