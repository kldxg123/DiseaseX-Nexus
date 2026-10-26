from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path

from docx import Document
from PIL import Image


ROOT = Path(r"F:\作业\aibola\reports\nc_manuscript_comparison_alignment_20260823")
DOCX = ROOT / "source" / "DiseaseX_Nexus_NC_Manuscript_v2_20260822.docx"
OUT = ROOT / "audit" / "original_media"

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
}


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    metadata = []
    with zipfile.ZipFile(DOCX) as zf:
        for name in sorted(n for n in zf.namelist() if n.startswith("word/media/")):
            raw = zf.read(name)
            dest = OUT / Path(name).name
            dest.write_bytes(raw)
            try:
                with Image.open(dest) as im:
                    width, height, fmt = im.width, im.height, im.format
            except Exception:
                width = height = None
                fmt = None
            metadata.append(
                {
                    "name": name,
                    "path": str(dest),
                    "size": len(raw),
                    "sha256": hashlib.sha256(raw).hexdigest(),
                    "width": width,
                    "height": height,
                    "format": fmt,
                }
            )

    doc = Document(DOCX)
    body_order = []
    p_index = -1
    t_index = -1
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            p_index += 1
            p = doc.paragraphs[p_index]
            blips = child.xpath(".//a:blip")
            images = []
            for blip in blips:
                rid = blip.get(f"{{{NS['r']}}}embed")
                if rid and rid in doc.part.rels:
                    images.append(doc.part.rels[rid].target_ref)
            body_order.append({"kind": "paragraph", "index": p_index, "style": p.style.name, "text": p.text, "images": images})
        elif child.tag.endswith("}tbl"):
            t_index += 1
            body_order.append({"kind": "table", "index": t_index, "rows": [[c.text for c in r.cells] for r in doc.tables[t_index].rows]})

    (OUT / "media_inventory.json").write_text(
        json.dumps({"media": metadata, "body_order": body_order}, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps({"media": metadata, "body_items": len(body_order)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
