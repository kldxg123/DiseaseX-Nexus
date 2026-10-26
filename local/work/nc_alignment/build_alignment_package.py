from __future__ import annotations

import csv
import hashlib
import json
import math
import shutil
import zipfile
from copy import deepcopy
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"F:\作业\aibola\reports\nc_manuscript_comparison_alignment_20260823")
SOURCE = ROOT / "source" / "DiseaseX_Nexus_NC_Manuscript_v2_20260822.docx"
MEDIA = ROOT / "audit" / "original_media"
EWNET = Path(r"F:\作业\aibola\reports\sci_q1_ewnet_benchmark_optimization_20260822")
OUTPUTS = EWNET / "evidence" / "outputs"
EVIDENCE = EWNET / "evidence"
OUT = ROOT / "final"
FIG = OUT / "Supplementary" / "Figures"
TAB = OUT / "Supplementary" / "Tables"
EVI = OUT / "Evidence"
AUDIT = OUT / "Audit"

MAIN_DOC = OUT / "DiseaseX_Nexus_NC_Manuscript_v3_ComparisonAligned_20260823.docx"
SI_DOC = OUT / "DiseaseX_Nexus_Supplementary_ComparisonExperiment_20260823.docx"
FINAL_ZIP = ROOT / "DiseaseX_Nexus_NC_ComparisonAligned_Package_20260823.zip"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_paragraph(paragraph, text: str, *, bold_prefix: str | None = None) -> None:
    clear_paragraph(paragraph)
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        run.bold = True
        paragraph.add_run(text[len(bold_prefix) :])
    else:
        paragraph.add_run(text)


def set_picture(paragraph, image_path: Path, width_inches: float = 6.5) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(width_inches))


def set_cell_text(cell, value: object, font_size: float = 8.5) -> None:
    cell.text = "" if value is None else str(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        for r in p.runs:
            r.font.size = Pt(font_size)


def replace_table(table, rows: list[list[object]], font_size: float = 8.5) -> None:
    while len(table.rows) < len(rows):
        table.add_row()
    while len(table.rows) > len(rows):
        table._tbl.remove(table.rows[-1]._tr)
    for ri, row in enumerate(rows):
        tr_pr = table.rows[ri]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if ri == 0:
            tbl_header = OxmlElement("w:tblHeader")
            tbl_header.set(qn("w:val"), "true")
            tr_pr.append(tbl_header)
        for ci, value in enumerate(row):
            if ci >= len(table.rows[ri].cells):
                continue
            set_cell_text(table.rows[ri].cells[ci], value, font_size)
            if ri == 0:
                for run in table.rows[ri].cells[ci].paragraphs[0].runs:
                    run.bold = True


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_table(doc: Document, rows: list[list[object]], widths: list[float] | None = None, font_size: float = 8.0):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        tr_pr = table.rows[ri]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if ri == 0:
            tbl_header = OxmlElement("w:tblHeader")
            tbl_header.set(qn("w:val"), "true")
            tr_pr.append(tbl_header)
        for ci, value in enumerate(row):
            cell = table.cell(ri, ci)
            set_cell_text(cell, value, font_size)
            if ri == 0:
                shade_cell(cell, "D9EAF7")
                for r in cell.paragraphs[0].runs:
                    r.bold = True
            if widths:
                cell.width = Inches(widths[ci])
    return table


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run(text)


def copy_image(src: Path, name: str) -> Path:
    dst = FIG / name
    shutil.copy2(src, dst)
    return dst


def font(size: int, bold: bool = False):
    candidates = [Path(r"C:\Windows\Fonts\arialbd.ttf") if bold else Path(r"C:\Windows\Fonts\arial.ttf")]
    for p in candidates:
        if p.exists():
            return ImageFont.truetype(str(p), size)
    return ImageFont.load_default()


def compose_main_figure() -> Path:
    top = Image.open(EWNET / "figures" / "02_test_trajectory.png").convert("RGB")
    bottom = Image.open(EWNET / "figures" / "01_test_wape_comparison.png").convert("RGB")
    width = max(top.width, bottom.width)
    def scale(im: Image.Image, w: int) -> Image.Image:
        if im.width == w:
            return im
        h = round(im.height * w / im.width)
        return im.resize((w, h), Image.Resampling.LANCZOS)
    top, bottom = scale(top, width), scale(bottom, width)
    gap = 50
    canvas = Image.new("RGB", (width, top.height + bottom.height + gap), "white")
    canvas.paste(top, (0, 0))
    canvas.paste(bottom, (0, top.height + gap))
    draw = ImageDraw.Draw(canvas)
    draw.text((18, 10), "a", fill="#111827", font=font(44, True))
    draw.text((18, top.height + gap + 10), "b", fill="#111827", font=font(44, True))
    out = FIG / "Fig4_EWNet_locked_test_system_outputs.png"
    canvas.save(out, optimize=True)
    return out


def load_metrics() -> pd.DataFrame:
    metrics = pd.read_csv(OUTPUTS / "test_metrics.csv")
    prior = json.loads((EVIDENCE / "prior_comparison_summary.json").read_text(encoding="utf-8"))
    p = prior["project_method"]["metrics"]
    extra = pd.DataFrame(
        [
            {
                "method": "lagged_count_boosting.v1 (pre-optimization)",
                "wape": p["wape"],
                "agreement": p["agreement"],
                "mae": p["mae"],
                "rmse": p["rmse"],
                "smape": p["smape"],
                "prediction_sha256": "see prior_project_predictions.csv",
                "n": 13,
            }
        ]
    )
    return pd.concat([metrics, extra], ignore_index=True)


DISPLAY = {
    "last_value": "Last-value baseline",
    "EWNet_author_code": "EWNet (author code)",
    "rolling_mean_4": "Rolling mean (4 weeks)",
    "kg_validation_only_convex_ensemble.v2": "DiseaseX optimized ensemble",
    "seasonal_naive_52": "Seasonal naive (52 weeks)",
    "lagged_count_boosting.v1 (pre-optimization)": "DiseaseX lagged-count boosting",
}


def make_full_wape_figure(metrics: pd.DataFrame) -> Path:
    frame = metrics.copy()
    frame["label"] = frame["method"].map(DISPLAY)
    order = frame.sort_values("wape", ascending=False)
    colors = ["#7C3AED" if "optimized" in x.lower() else "#2563EB" if "EWNet" in x else "#64748B" for x in order["label"]]
    fig, ax = plt.subplots(figsize=(11, 6.4))
    bars = ax.barh(order["label"], order["wape"] * 100, color=colors)
    ax.set_xlabel("Locked-test WAPE (%) — lower is better")
    ax.set_title("Singapore dengue: all methods on the same 13-week locked test")
    ax.grid(axis="x", alpha=0.25)
    for b, val in zip(bars, order["wape"] * 100):
        ax.text(val + 0.45, b.get_y() + b.get_height() / 2, f"{val:.2f}%", va="center", fontsize=9)
    ax.set_xlim(0, max(order["wape"] * 100) + 8)
    fig.text(0.01, 0.01, "Source: immutable system outputs test_metrics.csv and prior_comparison_summary.json; n=13 common weeks.", fontsize=8)
    fig.tight_layout(rect=(0, 0.04, 1, 1))
    out = FIG / "S6_all_method_locked_test_WAPE.png"
    fig.savefig(out, dpi=220, bbox_inches="tight")
    plt.close(fig)
    return out


def make_tracks_figure() -> Path:
    labels = ["A\nAUT legionellosis\nhistorical replay", "B\nFIN hantavirus\nrepeated window", "C\nCHL influenza\nfresh blind", "E\nSGP dengue\nEWNet benchmark"]
    wape = np.array([22.7, 28.0, 71.7, 45.51])
    agreement = 100 - wape
    decisions = ["not fresh", "not fresh", "gates failed", "retrospective; no promotion"]
    x = np.arange(len(labels))
    fig, ax = plt.subplots(figsize=(11, 6.2))
    bars = ax.bar(x, agreement, color=["#64748B", "#64748B", "#DC2626", "#7C3AED"])
    ax.axhline(70, color="#B91C1C", linestyle="--", linewidth=1.5, label="70% display threshold")
    ax.set_ylim(0, 100)
    ax.set_ylabel("Agreement = max(0, 1 − WAPE) × 100 (%)")
    ax.set_xticks(x, labels)
    ax.set_title("Evaluation identity governs interpretation, not the display score alone")
    ax.grid(axis="y", alpha=0.2)
    for b, a, d in zip(bars, agreement, decisions):
        ax.text(b.get_x() + b.get_width() / 2, a + 2, f"{a:.1f}%\n{d}", ha="center", va="bottom", fontsize=9)
    ax.legend(loc="upper right")
    fig.text(0.01, 0.01, "Agreement is not classification accuracy. Track E uses the corrected SCI Q1 EWNet locked-test result.", fontsize=8)
    fig.tight_layout(rect=(0, 0.04, 1, 1))
    out = FIG / "Fig5_governance_tracks_corrected.png"
    fig.savefig(out, dpi=220, bbox_inches="tight")
    plt.close(fig)
    return out


def make_qwen_figure() -> Path:
    rounds = [1, 2, 3]
    wape = [65.01, 60.83, 58.36]
    fig, ax = plt.subplots(figsize=(9.5, 5.4))
    ax.plot(rounds, wape, marker="o", linewidth=2.5, color="#2563EB")
    ax.axhline(30, color="#DC2626", linestyle="--", label="Prespecified target (30%)")
    for x, y in zip(rounds, wape):
        ax.text(x, y + 1.7, f"{y:.2f}%", ha="center")
    ax.set_xticks(rounds)
    ax.set_xlabel("Accepted Qwen proposal round")
    ax.set_ylabel("Validation WAPE (%)")
    ax.set_ylim(20, 72)
    ax.set_title("Controlled Qwen proposal loop improved validation error but missed its gate")
    ax.grid(alpha=0.25)
    ax.legend()
    fig.tight_layout()
    out = FIG / "S8_qwen_validation_rounds.png"
    fig.savefig(out, dpi=220, bbox_inches="tight")
    plt.close(fig)
    return out


def make_system_snapshot(metrics: pd.DataFrame) -> Path:
    summary = json.loads((OUTPUTS / "final_summary.json").read_text(encoding="utf-8"))
    size = (1800, 1120)
    im = Image.new("RGB", size, "#F4F7FB")
    d = ImageDraw.Draw(im)
    title = font(48, True)
    h2 = font(30, True)
    body = font(25)
    small = font(21)
    d.rounded_rectangle((45, 40, 1755, 1080), radius=28, fill="white", outline="#CBD5E1", width=3)
    d.text((90, 75), "Archived system result snapshot — SCI Q1 EWNet benchmark", fill="#0F172A", font=title)
    d.text((90, 150), "Immutable output artifacts; not a live dashboard capture", fill="#475569", font=small)
    cards = [
        (90, 220, 535, 420, "Dataset", "Singapore dengue\n838 weekly observations\n799 / 26 / 13 split"),
        (585, 220, 1030, 420, "Comparator", "EWNet author code\nNeural Networks (2023)\nDOI 10.1016/j.neunet.2023.05.049"),
        (1080, 220, 1710, 420, "Governance", "Retrospective test, opened once\nNo test labels used for selection\nDB writes / promotion / pointer changes = 0"),
    ]
    for x1, y1, x2, y2, head, text in cards:
        d.rounded_rectangle((x1, y1, x2, y2), radius=18, fill="#EEF4FF", outline="#93C5FD", width=2)
        d.text((x1 + 24, y1 + 20), head, fill="#1D4ED8", font=h2)
        d.multiline_text((x1 + 24, y1 + 70), text, fill="#1E293B", font=body, spacing=8)
    d.text((90, 470), "Locked-test metrics (13 common weeks)", fill="#0F172A", font=h2)
    columns = [("Method", 90), ("WAPE", 980), ("Agreement", 1180), ("MAE", 1420), ("RMSE", 1580)]
    y = 530
    d.rounded_rectangle((80, y - 10, 1720, y + 55), radius=10, fill="#DDE7F4")
    for name, x in columns:
        d.text((x, y), name, fill="#0F172A", font=small)
    rows = metrics.copy()
    rows["label"] = rows["method"].map(DISPLAY)
    rows = rows.sort_values("wape")
    y += 75
    for _, row in rows.iterrows():
        d.text((90, y), str(row["label"]), fill="#1E293B", font=small)
        d.text((980, y), f"{row['wape']*100:.2f}%", fill="#1E293B", font=small)
        d.text((1180, y), f"{row['agreement']*100:.2f}%", fill="#1E293B", font=small)
        d.text((1420, y), f"{row['mae']:.2f}", fill="#1E293B", font=small)
        d.text((1580, y), f"{row['rmse']:.2f}", fill="#1E293B", font=small)
        y += 62
    d.text((90, 1015), "Outcome: optimization did not generalize; Last-value < EWNet < prior DiseaseX < optimized ensemble in WAPE.", fill="#991B1B", font=small)
    out = FIG / "S9_archived_system_result_snapshot.png"
    im.save(out, optimize=True)
    return out


def metric_rows(metrics: pd.DataFrame) -> list[list[object]]:
    frame = metrics.copy()
    frame["label"] = frame["method"].map(DISPLAY)
    frame = frame.sort_values("wape")
    rows: list[list[object]] = [["Method", "WAPE (%) ↓", "Agreement (%)", "MAE", "RMSE", "n"]]
    for _, r in frame.iterrows():
        rows.append([r["label"], f"{r['wape']*100:.2f}", f"{r['agreement']*100:.2f}", f"{r['mae']:.2f}", f"{r['rmse']:.2f}", int(r["n"])])
    return rows


def write_csv(path: Path, rows: list[list[object]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        csv.writer(fh).writerows(rows)


def build_figures_and_tables() -> dict[str, Path]:
    for d in (OUT, FIG, TAB, EVI, AUDIT):
        d.mkdir(parents=True, exist_ok=True)
    metrics = load_metrics()
    figures = {
        "fig1": copy_image(MEDIA / "image1.jpeg", "S1_main_Fig1_system_architecture.jpeg"),
        "fig2": copy_image(MEDIA / "image2.jpeg", "S2_main_Fig2_knowledge_graph.jpeg"),
        "fig3": copy_image(MEDIA / "image3.jpeg", "S3_main_Fig3_Brazil_profile.jpeg"),
        "fig4": compose_main_figure(),
        "fig5": make_tracks_figure(),
        "wape": make_full_wape_figure(metrics),
        "gap": copy_image(EWNET / "figures" / "03_generalization_gap.png", "S7_generalization_gap.png"),
        "horizon": copy_image(EWNET / "figures" / "04_horizon_absolute_error.png", "S8_horizon_absolute_error.png"),
        "qwen": make_qwen_figure(),
        "snapshot": make_system_snapshot(metrics),
    }

    tables = {
        "asset": [
            ["Asset", "Count", "Interpretation"],
            ["Unique full-text versions", "1,574", "Versioned source documents for method and parameter evidence"],
            ["EvidenceSpan objects", "4,502", "Locatable excerpts with page/offset and dual hashes"],
            ["Assertions", "3,791", "Disease–model–condition and parameter-evidence relations"],
            ["Data profiles", "26", "Country–disease–indicator–frequency identities"],
            ["Execution contracts", "17", "Machine-consumable packages; all deferred in the audited snapshot"],
        ],
        "metrics": metric_rows(metrics),
        "tracks": [
            ["Track", "Series", "Evaluation identity", "WAPE", "Agreement", "Coverage", "Decision"],
            ["A", "AUT legionellosis, month", "Historical replay", "22.7%", "77.3%", "100%", "Not fresh"],
            ["B", "FIN hantavirus, week", "Repeated evaluation", "28.0%", "72.0%", "97.9%", "Reclassified non-fresh"],
            ["C", "CHL influenza, week", "Fresh blind test", "71.7%", "28.3%", "89.6%", "Gates failed; not certified"],
            ["E", "SGP dengue, week", "SCI Q1 EWNet retrospective test", "45.51%", "54.49%", "Not evaluated", "Optimized model did not beat EWNet/baseline; no promotion"],
        ],
        "brazil": [
            ["Candidate", "Evidence-fit", "Parameter confidence", "Gate result"],
            ["hurdle_count_ss.v1", "87.25", "75.00", "Eligible; rank 1"],
            ["negbin_dglm.v1", "85.44", "56.25", "Below parameter gate"],
            ["renewal_corri.v1", "85.25", "75.00", "Eligible; rank 3"],
            ["bayes_local_level.v1", "80.00", "80.00", "Eligible; rank 4"],
            ["bayes_local_trend.v1", "80.00", "80.00", "Eligible; rank 5"],
            ["gamma_poisson_dglm.v1", "71.50", "50.00", "Below parameter gate"],
        ],
        "feedback": [
            ["Series", "KG model WAPE", "Baseline WAPE", "Registered outcome"],
            ["AUT legionellosis, 240 monthly observations", "0.462 (negbin_dglm.v1)", "0.307 (seasonal naive)", "EMPIRICAL_BASELINE_OVERRIDE"],
            ["PER influenza, 53 weekly observations", "0.336 (hurdle_count_ss.v1)", "0.253 (rolling mean)", "EMPIRICAL_BASELINE_OVERRIDE"],
        ],
        "screen": [
            ["Candidate", "Median WAPE, first 3 windows", "Fourth-window WAPE", "Agreement", "Decision"],
            ["THA influenza", "38.83%", "34.56%", "65.44%", "Not promoted"],
            ["AUS influenza", "27.38%", "38.50%", "61.50%", "Not promoted"],
            ["MEX influenza", "41.52%", "41.93%", "58.07%", "Not promoted"],
        ],
        "qwen": [
            ["Item", "Value", "Interpretation"],
            ["Train / validation / hidden", "321 / 72 / 52", "Temporal split; hidden consumed once"],
            ["Accepted rounds / trials", "3 / 76", "Qwen proposes; deterministic code executes"],
            ["Validation WAPE", "0.6501 → 0.6083 → 0.5836", "Improved but missed 0.30 target"],
            ["Hidden WAPE", "0.7170", "Failed certification"],
            ["Rolling-mean hidden WAPE", "0.8248", "Champion reduced error but remained above gate"],
        ],
        "qa": [
            ["Artifact/scope", "Result", "Interpretation"],
            ["EWNet author-code reproduction", "PASS", "RMSE/MAE/sMAPE closely match paper Table 5"],
            ["Common locked test keys", "13/13", "All comparison methods evaluated on same weeks"],
            ["Future-label sentinel", "PASS", "Predictions unchanged after post-origin perturbation"],
            ["Deterministic repeat", "PASS", "Prediction SHA-256 repeated exactly"],
            ["Database/promotion/active pointer", "0 / 0 / 0", "Offline benchmark; no operational publication"],
            ["Optimization conclusion", "NEGATIVE GENERALIZATION", "Validation-selected ensemble did not beat EWNet or Last-value"],
        ],
        "planned": [
            ["Question", "Minimum design", "Required output", "Current status"],
            ["Multi-pathogen accuracy", "Registered held-out windows; fixed and domain baselines", "Horizon-specific error with paired uncertainty", "DATA/EXPERIMENT REQUIRED"],
            ["Prospective early warning", "Real-time origins; fixed peak definition", "Lead time, missed peaks, false alarms", "DATA/EXPERIMENT REQUIRED"],
            ["Uncertainty calibration", "Probabilistic forecasts on common targets", "CRPS, coverage, interval width", "DATA/EXPERIMENT REQUIRED"],
            ["Self-learning benefit", "Frozen no-feedback control vs feedback system", "Paired change in error and decision quality", "EXPERIMENT REQUIRED"],
            ["Component contribution", "Ablate KG, LLM, contract, feedback", "Performance, safety, auditability", "EXPERIMENT REQUIRED"],
        ],
        "provenance": [
            ["Artifact", "Identity / SHA-256", "Role"],
            ["Paper", "10.1016/j.neunet.2023.05.049", "SCI Q1 EWNet comparator"],
            ["Author repository", "commit 3740776f65c6f4157e67e6a78f71daeade45246b", "EWNet implementation"],
            ["Source workbook", "19b55bfcede7539fd15087f21d0b07e9d2e78dd9b19c6584057fecc721cd325a", "838-week source series"],
            ["Canonical CSV", "86f057db4b7e0f56ec5d73b86f90738d2450b92fbd4c1a3109e0cb06e4f025a9", "Normalized benchmark input"],
            ["Selection lock", "41905d950f2071a7d080240ca2066fac18c14a5b9c4f195c2da6e7ce22ada312", "Frozen before test"],
            ["Selected prediction", "ce7bfbb50206ba20735630bd4408ff48a94c3424a792ea1c9faa2cd92d489ecb", "Exactly repeated locked-test prediction"],
        ],
    }
    for name, rows in tables.items():
        write_csv(TAB / f"SI_Table_{name}.csv", rows)

    for src in [
        OUTPUTS / "final_summary.json",
        OUTPUTS / "test_metrics.csv",
        OUTPUTS / "test_paired_predictions.csv",
        OUTPUTS / "locked_selection.json",
        OUTPUTS / "model_search.csv",
        OUTPUTS / "tuning_predictions.csv",
        EVIDENCE / "prior_comparison_summary.json",
        EVIDENCE / "prior_project_predictions.csv",
    ]:
        shutil.copy2(src, EVI / src.name)
    return {**figures, "metrics_df": metrics, "tables": tables}


def build_main(assets: dict[str, object]) -> None:
    doc = Document(SOURCE)
    repl = {
        2: "Infectious-disease forecasting systems must adapt to changing pathogens, reporting regimes and data availability, yet model selection and updating are often fixed, weakly documented or difficult to audit. We present DiseaseX Nexus, a local-first framework that uses large-language-model agents to retrieve evidence and propose bounded model configurations, a versioned knowledge graph to connect diseases, data profiles, models and parameter provenance, and deterministic code to execute backtests and forecasts only after an explicit authorization contract. The implemented knowledge assets comprised 1,574 full-text versions, 4,502 evidence spans, 3,791 assertions, 26 data profiles and 17 execution contracts. A Brazil dengue case study demonstrates data profiling, evidence-scored model admission and a deliberately deferred execution state. In a separate author-code-aligned Singapore dengue benchmark against the 2023 SCI Q1 EWNet paper, the 13-week locked test yielded WAPE 40.05% for a last-value baseline, 41.07% for EWNet, 42.46% for the initial DiseaseX model and 45.51% for a validation-selected optimized ensemble. The negative generalization result was retained without promotion. Governance-aware evaluation further showed that the only fresh blind test failed prespecified error gates. These results establish an auditable agent-orchestration and feedback framework while delimiting prospective, multi-pathogen performance claims that require further experiments.",
        12: "Here we evaluate what the current implementation can substantiate. We first quantify the knowledge and contract assets, then trace a non-executed Brazil dengue recommendation from live observations to an authorization gate. We separately analyse a 2023 SCI Q1 EWNet Singapore dengue benchmark and four governed evaluation tracks. The results support the feasibility, auditability and error-containing behaviour of the orchestration framework. They do not establish superiority over state-of-the-art systems, general multi-pathogen accuracy, calibrated uncertainty or operational early-warning benefit; those claims remain planned rather than inferred from incomplete evidence.",
        29: "A SCI Q1 EWNet benchmark exposes a negative generalization result",
        30: "We next examined an executed benchmark independent of the Brazil case. The comparator was Epicasting: An Ensemble Wavelet Neural Network for forecasting epidemics (EWNet), published in Neural Networks in 2023 (ref. 26). The authors' repository supplied the single-variable Singapore dengue weekly case series and implementation. The canonical input contained 838 observations and used the paper-aligned temporal split of 799 training, 26 validation and 13 locked-test weeks. The source workbook and normalized CSV were frozen by SHA-256 before selection.",
        31: "Six methods were evaluated on the same 13 target weeks: EWNet rerun from the authors' code, the initial DiseaseX lagged-count boosting model, a validation-only optimized DiseaseX convex ensemble, and last-value, four-week rolling-mean and 52-week seasonal-naive baselines. Candidate selection used five seasonally matched historical folds plus the 26-week validation segment; the 13 test labels were not used for model or weight selection and were opened once. All methods issued a fixed 13-week multi-step forecast from the common origin, and independent recomputation reproduced the archived metrics and prediction hashes.",
        32: "On the locked test, the last-value baseline obtained the lowest WAPE (40.05%; MAE 161.08; RMSE 216.76), followed by EWNet (41.07%; MAE 165.18; RMSE 218.89). The initial DiseaseX lagged-count model obtained WAPE 42.46%. The validation-selected optimized ensemble obtained WAPE 45.51% (MAE 183.04; RMSE 229.50), 10.81% higher relative WAPE than EWNet and 13.64% higher than last value. Thus the attempted optimization improved its historical/validation objective but worsened locked-test generalization. Agreement, max(0, 1 − WAPE) × 100, was 54.49% for the optimized ensemble and is a display transformation of aggregate error, not classification accuracy (Fig. 4 and Table 2).",
        34: "Fig. 4 Singapore dengue SCI Q1 EWNet benchmark from locked system outputs. a, Actual counts and fixed-origin 13-week predictions from EWNet, the DiseaseX validation-selected ensemble and baselines. b, Common-sample test WAPE; lower is better. The project optimization did not beat EWNet or the last-value baseline. This was a retrospective test opened once, not a fresh prospective evaluation.",
        35: "Table 2 Singapore dengue SCI Q1 EWNet benchmark on 13 paired locked-test weeks",
        36: "The governance identity of this experiment was RETROSPECTIVE_TEST_EXECUTED_ONCE with no database write, forecast publication, promotion or active-pointer change. It demonstrates author-code reproduction, a frozen selection boundary and retention of an unfavourable result. It does not show that DiseaseX outperforms EWNet or a strong naive baseline, and it cannot substitute for prospective evaluation. The comparator's single-series design made it suitable for a same-data method comparison, but the short 13-week terminal window limits generalization claims.",
        38: "Four result identities illustrate why metric provenance matters (Fig. 5 and Table 3). Track A, monthly Austrian legionellosis, was a historical replay with WAPE 22.7%, agreement 77.3% and interval coverage 100%. Track B, weekly Finnish hantavirus, had WAPE 28.0%, agreement 72.0% and coverage 97.9%, but prior test-window use required reclassification as RETROSPECTIVE_REPLAY_NOT_FRESH. Track E is the corrected Singapore EWNet retrospective benchmark: the DiseaseX optimized ensemble had WAPE 45.51% and agreement 54.49% and did not beat EWNet or last value. Only A and B exceeded the 70% display threshold, and neither was fresh.",
        41: "Fig. 5 Scores interpreted with their evaluation identities. Agreement is max(0, 1 − WAPE) × 100 and is not classification accuracy. Tracks A and B exceeded 70% but were historical or repeatedly evaluated. Track C was the only fresh blind test and failed the prespecified WAPE and agreement gates. Track E is the corrected SCI Q1 EWNet retrospective benchmark and also remained unpromoted. Scores across tracks are not direct estimates of comparative model efficacy because diseases, frequencies, windows and evaluation identities differ.",
        54: "The SCI Q1 EWNet benchmark provides a calibrated view of predictive capability. The authors' EWNet reproduction closely matched the published test metrics, but a simple last-value baseline was slightly better on this single 13-week window. The initial DiseaseX model was close to EWNet, whereas a later validation-selected ensemble generalized worse. This negative result reinforces the need for simple baselines, external comparators and once-only test use. It also illustrates why an adaptive architecture should be evaluated separately from any one selected forecast: orchestration can be correct even when optimization fails, provided the system records the outcome and does not promote it.",
        56: "This study has five main limitations. First, the primary Brazil case stops at recommendation and does not test forecasts. Second, only one fresh blind test was available and it failed. Third, the direct Singapore comparator is EWNet from a 2023 SCI Q1 article (ref. 26), but the comparison is limited to one 13-week terminal window and the optimized DiseaseX ensemble did not beat EWNet or last value. Fourth, the knowledge assets are a verified snapshot; incremental production construction was incomplete. Fifth, the current evaluation does not estimate decision utility, fairness, cross-site transportability, probabilistic calibration or prospective early-warning lead time. These limitations preclude claims of clinical or public-health effectiveness and constrain the paper to system validation and transparent benchmark evidence.",
        71: "Singapore dengue EWNet benchmark and locked-test models",
        72: "The benchmark used the Singapore dengue weekly case series distributed with the EWNet authors' repository at commit 3740776f65c6f4157e67e6a78f71daeade45246b. The source workbook contained 838 observations and had SHA-256 19b55bfcede7539fd15087f21d0b07e9d2e78dd9b19c6584057fecc721cd325a; the normalized CSV had SHA-256 86f057db4b7e0f56ec5d73b86f90738d2450b92fbd4c1a3109e0cb06e4f025a9. The chronological split was 799 training, 26 validation and 13 locked-test observations. The test segment covered 2015-W41 to W53 and was not used during selection.",
        73: "EWNet was rerun from the authors' implementation after three documented syntax/runtime repairs that did not alter the MODWT, autoregressive neural-network structure, p = 1 setting or 500 repeats. The rerun closely reproduced the paper's Singapore test RMSE, MAE and sMAPE. The initial DiseaseX model was lagged_count_boosting.v1. A subsequent optimization searched 65 validation-eligible candidates using five historical W41–W53 folds plus the paper validation period; a convex ensemble of two ridge components and one seasonally scaled component was locked before test. Every method then produced a fixed 13-week multi-step forecast from the same origin. A future-label sentinel and deterministic repeat preserved the selected prediction SHA-256 ce7bfbb50206ba20735630bd4408ff48a94c3424a792ea1c9faa2cd92d489ecb.",
        77: "The primary point-forecast metric was weighted absolute percentage error, WAPE = Σ|yₜ − ŷₜ| / Σ|yₜ|. Secondary metrics were mean absolute error, MAE = n⁻¹Σ|yₜ − ŷₜ|, and root mean squared error, RMSE = [n⁻¹Σ(yₜ − ŷₜ)²]¹ᐟ². Agreement was defined for display as max(0, 1 − WAPE) × 100% and was not treated as accuracy. Interval coverage was the proportion of observed targets inside the corresponding forecast interval. Metrics were computed on common target keys. Because the executed comparisons were not powered or preregistered for inferential testing, no P values or confidence intervals for between-model differences were generated; comparisons are descriptive.",
        80: "The final material package passed package-level verification and secret scanning. Test reports with 22, 175 and 178 passing tests refer to different scopes and dates and were not merged into a single total. For the EWNet benchmark, the authors' code closely reproduced the paper-reported Singapore RMSE, MAE and sMAPE; the common 13-week predictions, selection lock, future-label sentinel and deterministic repeat were archived with SHA-256 identities. File hashes, data manifests, model identities and evaluation states are retained in the Supplementary Information. One strict typing check in an earlier packaged environment was not executed because librt.internal was unavailable; this is an environment limitation rather than a passed test.",
        84: "The Singapore benchmark series and EWNet implementation are available from the authors' Epicasting repository and are identified here by repository commit and file hashes. The normalized series, locked selection, paired predictions, metrics and non-sensitive audit objects are included in the authors' reproducibility package. Brazil dengue observations were retrieved under the source key who_dengue_global_validated and are represented in the archived data-profile object. [PLACEHOLDER: deposit the exact processed datasets, source-response manifests and non-sensitive audit objects in a persistent repository; insert accession and DOI. Any source-specific redistribution restrictions must be stated.]",
        121: "26. Panja, M., Chakraborty, T., Kumar, U. & Liu, N. Epicasting: An Ensemble Wavelet Neural Network for forecasting epidemics. Neural Networks 165, 185–212 (2023). https://doi.org/10.1016/j.neunet.2023.05.049",
    }
    for idx, text in repl.items():
        set_paragraph(doc.paragraphs[idx], text)
    set_picture(doc.paragraphs[33], assets["fig4"], 6.4)
    set_picture(doc.paragraphs[40], assets["fig5"], 6.4)
    doc.paragraphs[35].paragraph_format.page_break_before = True
    doc.paragraphs[42].paragraph_format.page_break_before = True
    replace_table(doc.tables[1], assets["tables"]["metrics"], 7.8)
    replace_table(doc.tables[2], assets["tables"]["tracks"], 7.4)
    doc.core_properties.title = "DiseaseX Nexus: Self-improving LLM agents for adaptive infectious disease forecasting"
    doc.core_properties.comments = "Comparison experiment aligned to SCI Q1 EWNet evidence on 23 August 2026."
    doc.save(MAIN_DOC)


def reset_body_keep_design(doc: Document) -> None:
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def build_supplementary(assets: dict[str, object]) -> None:
    doc = Document(SOURCE)
    reset_body_keep_design(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Supplementary Information")
    r.bold = True
    r.font.size = Pt(20)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DiseaseX Nexus: Self-improving LLM agents for adaptive infectious disease forecasting")
    r.italic = True
    r.font.size = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Comparison-aligned version · 23 August 2026")

    doc.add_heading("Supplementary Note 1. Comparison-experiment replacement map", level=1)
    doc.add_paragraph(
        "The submitted draft used a 412-week, AR(2)-based Singapore benchmark with a 311/50/50 split. "
        "That experiment remains a historical artifact but is not the current primary comparator. The corrected manuscript uses the 2023 "
        "Neural Networks EWNet article, the authors' 838-week Singapore dengue series, a 799/26/13 temporal split and a once-opened 13-week locked test. "
        "All numbers below are taken from immutable system output files; no value was transcribed from the old figure."
    )
    mapping_rows = [
        ["Manuscript item", "Old content", "Corrected content", "Evidence"],
        ["Abstract / Results", "AR(2), 50-week test, WAPE 11.34–13.82%", "EWNet locked test; WAPE 40.05–50.81%; negative generalization", "final_summary.json; test_metrics.csv"],
        ["Methods", "412 points; 311/50/50; rolling one-step", "838 points; 799/26/13; fixed-origin 13-week multi-step", "data_identity.json; locked_selection.json"],
        ["Primary paper", "Older/incorrect dengue comparator citation", "Panja et al., Neural Networks 165 (2023), DOI 10.1016/j.neunet.2023.05.049", "Publisher metadata; author repository"],
        ["Figure 4 / Table 2", "AR(2) trajectory and four-method table", "EWNet, prior DiseaseX, optimized DiseaseX and three baselines", "paired predictions; recomputed metrics"],
        ["Track E", "Agreement 87.20%", "Agreement 54.49%; retrospective; not promoted", "corrected locked-test output"],
        ["Discussion", "Project ensemble competitive with AR(2)", "Optimization worsened held-out generalization; simple baseline won", "comparison and generalization-gap outputs"],
    ]
    add_table(doc, mapping_rows, font_size=7.5)

    doc.add_heading("Supplementary Note 2. Benchmark identity and protocol", level=1)
    doc.add_paragraph(
        "The comparator is EWNet, a MODWT-based ensemble of autoregressive neural networks. The authors' implementation was rerun at repository "
        "commit 3740776f65c6f4157e67e6a78f71daeade45246b after three documented runtime-only repairs. The reproduction closely matched the paper's "
        "Singapore RMSE, MAE and sMAPE. The project optimization used only historical folds and the 26-week validation period; test labels were opened once after locking."
    )
    add_table(doc, assets["tables"]["provenance"], font_size=7.3)

    doc.add_heading("Supplementary Figures", level=1)
    figure_items = [
        ("S1", assets["fig1"], "Copy of main Fig. 1. DiseaseX Nexus system architecture."),
        ("S2", assets["fig2"], "Copy of main Fig. 2. Knowledge-graph provenance and execution-contract assets."),
        ("S3", assets["fig3"], "Copy of main Fig. 3. Brazil dengue data profile and evidence-scored admission."),
        ("S4", assets["fig4"], "Copy of main Fig. 4. SCI Q1 EWNet benchmark from locked system outputs."),
        ("S5", assets["fig5"], "Copy of main Fig. 5. Evaluation identity and corrected Track E result."),
        ("S6", assets["wape"], "All six methods on the same 13-week locked test. Lower WAPE is better."),
        ("S7", assets["gap"], "Historical/validation selection versus locked-test WAPE. The optimized ensemble failed to generalize."),
        ("S8", assets["horizon"], "Absolute error by test horizon from the archived evaluation output."),
        ("S9", assets["qwen"], "Controlled Qwen proposal rounds. Validation WAPE improved but remained above the gate."),
        ("S10", assets["snapshot"], "Archived system-result snapshot assembled directly from immutable JSON/CSV outputs; not a live dashboard capture."),
    ]
    for label, path, caption in figure_items:
        doc.add_heading(f"Supplementary Figure {label}", level=2)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(6.45))
        add_caption(doc, caption)

    doc.add_heading("Supplementary Tables", level=1)
    table_items = [
        ("S1", "Versioned knowledge and execution assets", "asset"),
        ("S2", "SCI Q1 EWNet benchmark locked-test metrics", "metrics"),
        ("S3", "Governance-aware evaluation tracks", "tracks"),
        ("S4", "Brazil dengue R3C candidate ranking", "brazil"),
        ("S5", "Empirical feedback overrides", "feedback"),
        ("S6", "Pseudo-blind influenza screening", "screen"),
        ("S7", "Controlled Qwen tuning", "qwen"),
        ("S8", "Verification and non-mutation evidence", "qa"),
        ("S9", "Analyses still required before prospective performance claims", "planned"),
        ("S10", "Benchmark provenance and hashes", "provenance"),
    ]
    for index, (label, title, key) in enumerate(table_items):
        if index:
            doc.add_page_break()
        heading = doc.add_heading(f"Supplementary Table {label}. {title}", level=2)
        heading.paragraph_format.keep_with_next = True
        add_table(doc, assets["tables"][key], font_size=7.5 if key in {"metrics", "tracks", "planned", "provenance"} else 8.0)
        doc.add_paragraph()

    doc.add_heading("Supplementary Note 3. Interpretation limits", level=1)
    doc.add_paragraph(
        "WAPE is an aggregate scale-normalized error. Agreement is defined only as max(0, 1 − WAPE) × 100 and is not classification accuracy. "
        "The EWNet benchmark is retrospective and limited to 13 terminal weeks. The result supports reproducibility and governance claims, but not "
        "state-of-the-art superiority, prospective early-warning effectiveness or clinical/public-health validation."
    )
    doc.core_properties.title = "DiseaseX Nexus Supplementary Information"
    doc.core_properties.comments = "Consolidated figures, tables, benchmark evidence and replacement map."
    doc.save(SI_DOC)


def write_mapping_files() -> None:
    rows = [
        ["序号", "位置", "原内容", "问题", "替换内容", "直接证据"],
        [1, "摘要", "50周AR(2)对比，WAPE 11.34/12.80%", "与当前SCI Q1实验不一致", "13周EWNet锁定测试；Last-value 40.05%，EWNet 41.07%，初始项目模型42.46%，优化集成45.51%", "final_summary.json; test_metrics.csv"],
        [2, "Results 对比实验小节", "412点、311/50/50、rolling one-step", "旧实验协议", "838点、799/26/13、固定起点13周多步预测", "data identity; locked selection"],
        [3, "原 Supplementary Fig.1", "AR(2)四方法图", "旧对比论文与旧切分", "EWNet系统锁定输出轨迹+WAPE复合图", "test_paired_predictions.csv; test_metrics.csv"],
        [4, "Table 2", "AR(2)四方法指标", "缺少EWNet、优化前后与三个基线", "六方法同窗指标表", "test_metrics.csv; prior_comparison_summary.json"],
        [5, "Track E / Fig.5 / Table 3", "WAPE 12.80%，agreement 87.20%", "沿用旧实验", "WAPE 45.51%，agreement 54.49%，retrospective/not promoted", "corrected locked-test output"],
        [6, "Discussion", "项目集成接近AR(2)", "未反映优化失败", "初始项目模型接近EWNet；新优化集成泛化变差；Last-value最优", "generalization gap; locked test"],
        [7, "Methods", "旧数据源、系数与0.05权重网格", "不属于EWNet实验", "作者仓库commit、EWNet运行修复、65候选、验证锁定与future-sentinel", "FINAL_REPORT.md; locked_selection.json"],
        [8, "Reference 26", "错误的2022巴西城市论文", "不是本文实际比较方法", "Panja et al., Neural Networks 165 (2023), DOI 10.1016/j.neunet.2023.05.049", "publisher and article DOI"],
        [9, "Supplementary", "9个分散Note及13个CSV，实验图不全", "对比实验分散且旧数据混入", "合并为1份SI，包含主文全部图表、4张评估图、系统结果快照和10张表", "本次生成包"],
    ]
    write_csv(AUDIT / "对比实验数据图表替换清单_20260823.csv", rows)
    lines = ["# 对比实验数据与图表替换清单", "", "本清单只把附件内容当作待核验材料；最终数值以锁定 JSON/CSV 为准。", ""]
    for r in rows[1:]:
        lines.extend([f"## {r[0]}. {r[1]}", f"- 原内容：{r[2]}", f"- 问题：{r[3]}", f"- 替换：{r[4]}", f"- 证据：{r[5]}", ""])
    (AUDIT / "对比实验数据图表替换清单_20260823.md").write_text("\n".join(lines), encoding="utf-8")


def package() -> None:
    manifest_lines = []
    for p in sorted(x for x in OUT.rglob("*") if x.is_file() and x.name != "MANIFEST.sha256"):
        manifest_lines.append(f"{sha256(p)}  {p.relative_to(OUT).as_posix()}")
    (OUT / "MANIFEST.sha256").write_text("\n".join(manifest_lines) + "\n", encoding="utf-8")
    if FINAL_ZIP.exists():
        FINAL_ZIP.unlink()
    with zipfile.ZipFile(FINAL_ZIP, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as zf:
        for p in sorted(x for x in OUT.rglob("*") if x.is_file()):
            zf.write(p, p.relative_to(OUT).as_posix())


def main() -> None:
    if OUT.exists():
        shutil.rmtree(OUT)
    assets = build_figures_and_tables()
    build_main(assets)
    build_supplementary(assets)
    write_mapping_files()
    package()
    print(
        json.dumps(
            {
                "main_docx": str(MAIN_DOC),
                "supplementary_docx": str(SI_DOC),
                "zip": str(FINAL_ZIP),
                "zip_sha256": sha256(FINAL_ZIP),
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
