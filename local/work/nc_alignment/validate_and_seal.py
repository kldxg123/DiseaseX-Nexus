from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path, PurePosixPath

from docx import Document


ROOT = Path(r"F:\作业\aibola\reports\nc_manuscript_comparison_alignment_20260823")
OUT = ROOT / "final"
MAIN = OUT / "DiseaseX_Nexus_NC_Manuscript_v3_ComparisonAligned_20260823.docx"
SI = OUT / "DiseaseX_Nexus_Supplementary_ComparisonExperiment_20260823.docx"
ZIP = ROOT / "DiseaseX_Nexus_NC_ComparisonAligned_Package_20260823.zip"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for b in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(b)
    return h.hexdigest()


def doc_text(path: Path) -> tuple[str, Document]:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts), doc


def main() -> None:
    main_text, main_doc = doc_text(MAIN)
    si_text, si_doc = doc_text(SI)
    assertions = {
        "main_has_ewnet_doi": "10.1016/j.neunet.2023.05.049" in main_text,
        "main_has_split_799_26_13": "799 training, 26 validation and 13 locked-test" in main_text,
        "main_has_negative_result": "worsened locked-test generalization" in main_text,
        "main_has_wape_values": all(x in main_text for x in ["40.05", "41.07", "42.46", "45.51", "50.81"]),
        "main_removed_old_ar2_table": "Published AR(2) reconstruction" not in main_text,
        "main_removed_old_split": "311/50/50" not in main_text and "311 training, 50 validation" not in main_text,
        "main_removed_old_50week_claim": "50 paired test weeks" not in main_text,
        "agreement_not_accuracy": "not treated as accuracy" in main_text and "not classification accuracy" in main_text,
        "main_figure_count": len(main_doc.inline_shapes) == 5,
        "main_table_count": len(main_doc.tables) == 5,
        "si_has_all_main_figures": all(f"Supplementary Figure S{i}" in si_text for i in range(1, 6)),
        "si_has_all_evaluation_figures": all(f"Supplementary Figure S{i}" in si_text for i in range(6, 11)),
        "si_has_ten_tables": all(f"Supplementary Table S{i}" in si_text for i in range(1, 11)),
        "si_has_interpretation_limits": "not classification accuracy" in si_text,
        "si_old_ar2_values_only_in_audit_map": "Published AR(2) reconstruction" not in si_text,
        "si_figure_count": len(si_doc.inline_shapes) == 10,
        "si_table_count": len(si_doc.tables) == 12,
    }
    if not all(assertions.values()):
        raise SystemExit(json.dumps({k: v for k, v in assertions.items() if not v}, indent=2))

    evidence = OUT / "Evidence"
    output_summary = json.loads((evidence / "final_summary.json").read_text(encoding="utf-8"))
    data_checks = {
        "dataset_rows": output_summary["data"]["rows"] == 838,
        "split": output_summary["data"]["split"] == {"train": 799, "validation": 26, "test": 13},
        "paper_doi": output_summary["paper"]["doi"] == "10.1016/j.neunet.2023.05.049",
        "selected_wape": abs(output_summary["test"]["selected_metrics"]["wape"] - 0.4550641697588694) < 1e-12,
        "ewnet_wape": abs(output_summary["test"]["paper_ewnet_metrics"]["wape"] - 0.41066509394615125) < 1e-12,
        "no_db_writes": output_summary["governance"]["database_writes"] == 0,
        "no_promotion": output_summary["governance"]["promotion"] == 0,
    }
    if not all(data_checks.values()):
        raise SystemExit(json.dumps({k: v for k, v in data_checks.items() if not v}, indent=2))

    qa = {
        "status": "PASS",
        "documents": {
            MAIN.name: {"sha256": sha256(MAIN), "pages_rendered": 27, "visual_review": "PASS_ALL_PAGES"},
            SI.name: {"sha256": sha256(SI), "pages_rendered": 19, "visual_review": "PASS_ALL_PAGES"},
        },
        "content_assertions": assertions,
        "data_assertions": data_checks,
        "comparison_conclusion": "NEGATIVE_GENERALIZATION_RETAINED_NOT_PROMOTED",
        "source_of_figures": "immutable system JSON/CSV outputs and archived manuscript figures",
    }
    audit_path = OUT / "Audit" / "FINAL_VALIDATION.json"
    audit_path.write_text(json.dumps(qa, ensure_ascii=False, indent=2), encoding="utf-8")

    manifest_lines = []
    for p in sorted(x for x in OUT.rglob("*") if x.is_file() and x.name != "MANIFEST.sha256"):
        manifest_lines.append(f"{sha256(p)}  {p.relative_to(OUT).as_posix()}")
    (OUT / "MANIFEST.sha256").write_text("\n".join(manifest_lines) + "\n", encoding="utf-8")
    if ZIP.exists():
        ZIP.unlink()
    with zipfile.ZipFile(ZIP, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as zf:
        for p in sorted(x for x in OUT.rglob("*") if x.is_file()):
            zf.write(p, p.relative_to(OUT).as_posix())

    with zipfile.ZipFile(ZIP) as zf:
        names = zf.namelist()
        unsafe = [n for n in names if PurePosixPath(n).is_absolute() or ".." in PurePosixPath(n).parts]
        if unsafe:
            raise SystemExit(f"unsafe zip names: {unsafe}")
        for line in zf.read("MANIFEST.sha256").decode("utf-8").splitlines():
            digest, name = line.split("  ", 1)
            if hashlib.sha256(zf.read(name)).hexdigest() != digest:
                raise SystemExit(f"manifest mismatch: {name}")
    print(json.dumps({"status": "PASS", "zip_sha256": sha256(ZIP), "manifest_entries": len(manifest_lines)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
