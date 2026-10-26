# ruff: noqa: E501
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from talk_content import SECTIONS

ROOT = Path(r"F:\作业\aibola")
SOURCE = ROOT / "DiseaseX_Nexus_NC_Manuscript_v10_English_Polished.docx"
FIGURE_DIR = ROOT / "reports" / "hk_flu_comparison_eeg_style_20260827"
OUTPUT_MANUSCRIPT = ROOT / "DiseaseX_Nexus_NC_Manuscript_v11_New_Comparison_Figures.docx"
OUTPUT_SCRIPT = ROOT / "DiseaseX_Nexus_逐段中文讲稿_多病种自适应与泛化.docx"

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x0B, 0x25, 0x45)
MUTED = RGBColor(0x55, 0x55, 0x55)
GOLD = RGBColor(0x7A, 0x5A, 0x00)
LIGHT_FILL = "F4F6F9"


def set_run_font(
    run,
    *,
    name: str = "Calibri",
    size: float | None = None,
    color: RGBColor | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def replace_text(paragraph, text: str) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    set_run_font(run)


def replace_picture(paragraph, image_path: Path, alt_text: str) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(4)
    inline_shape = paragraph.add_run().add_picture(str(image_path), width=Inches(6.35))
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("name", image_path.stem)
    doc_pr.set("descr", alt_text)


def build_manuscript() -> None:
    doc = Document(SOURCE)
    paragraphs = doc.paragraphs

    replace_text(
        paragraphs[2],
        "Forecasting systems for infectious diseases must adapt to changes in pathogens, "
        "reporting regimes and data availability. However, model selection and updating are "
        "often fixed, poorly documented or difficult to audit. We present DiseaseX Nexus, a "
        "local-first framework in which LLM agents retrieve evidence and propose bounded model "
        "configurations; a versioned knowledge graph links diseases, data profiles, models and "
        "parameter provenance; and deterministic code executes backtests and forecasts only "
        "after approval through an explicit authorization contract. The implemented snapshot "
        "comprised 1,574 full-text versions, 4,502 evidence spans, 3,791 assertions, 26 data "
        "profiles and 17 execution contracts. A Brazil dengue case study demonstrates data "
        "profiling, evidence-scored model admission and a deliberately deferred execution state. "
        "In a 4,968-row matched-pair Hong Kong influenza benchmark against predictions released "
        "by the authors of a 2024 Nature Communications study, the validation-locked horizon "
        "router achieved a main-period WAPE of 33.94%, compared with 34.18% for paper AWBE. In "
        "the post-COVID period, however, the corresponding WAPE values were 49.35% and 48.88%, "
        "respectively, and external generalization was therefore not confirmed. The independent "
        "DiseaseX ensemble improved substantially over the constant baseline but remained less "
        "accurate than AWBE. Governance-aware evaluation further showed that the only fresh blind "
        "test failed the prespecified error gates. Together, these results establish an auditable "
        "framework for cross-disease agent orchestration and feedback while preventing a small "
        "retrospective improvement from being misrepresented as stable predictive superiority.",
    )
    replace_text(
        paragraphs[32],
        "On the primary 4,968-row test, the DiseaseX independent ensemble reduced WAPE from "
        "75.94% for the paper constant baseline to 50.27%, but remained 16.09 percentage points "
        "above paper AWBE at 34.18%. The validation-locked KG horizon router achieved a WAPE of "
        "33.94%, an MAE of 1.6376 and an RMSE of 2.7255, compared with a WAPE of 34.18%, an MAE "
        "of 1.6491 and an RMSE of 2.7472 for AWBE (Supplementary Figs. 1–3 and Table 2). The "
        "0.24-percentage-point WAPE difference corresponds to a relative reduction of 0.70%. "
        "Horizon-specific analysis showed that the independent DiseaseX ensemble outperformed "
        "AWBE only at horizon 0 and was progressively worse at longer horizons. In the 423-row "
        "post-COVID external test, the router yielded a WAPE of 49.35%, compared with 48.88% for "
        "AWBE (Supplementary Fig. 4). The main-period gain was therefore not confirmed externally. "
        "Agreement is defined as max(0, 1 − WAPE) × 100 and is not a measure of classification "
        "accuracy.",
    )
    replace_text(
        paragraphs[42],
        "Four additional result identities illustrate the importance of metric provenance "
        "(Table 3). Track A, monthly Austrian legionellosis, was a historical replay with a WAPE "
        "of 22.7%, agreement of 77.3% and interval coverage of 100%. Track B, weekly Finnish "
        "hantavirus, had a WAPE of 28.0%, agreement of 72.0% and coverage of 97.9%; however, "
        "previous use of the test window required reclassification as "
        "RETROSPECTIVE_REPLAY_NOT_FRESH. Track E was the Hong Kong influenza author-output "
        "benchmark: the main-period KG router had a WAPE of 33.94% and agreement of 66.06%, "
        "whereas the corresponding post-COVID external values were 49.35% and 50.65%. Tracks A, "
        "B and E therefore did not constitute new blind tests.",
    )
    replace_text(
        paragraphs[58],
        "The Hong Kong influenza benchmark provides a calibrated assessment of predictive "
        "capability. The DiseaseX independent ensemble reduced WAPE from 75.94% for the constant "
        "baseline to 50.27%, but did not approach paper AWBE at 34.18%. The validation-locked "
        "horizon router reduced WAPE by 0.24 percentage points relative to AWBE during the long "
        "main period, but its WAPE was 0.47 percentage points higher in the post-COVID external "
        "period. This reversal demonstrates the need to evaluate horizon-aware orchestration under "
        "temporal distribution shift and shows why a marginal retrospective gain should not be "
        "presented as stable superiority. It also separates the portability of the orchestration "
        "mechanism from the generalization of a particular predictive configuration.",
    )
    replace_text(
        paragraphs[82],
        "For the present horizon-specific DiseaseX–AWBE comparison, exploratory 95% intervals "
        "were estimated from 2,500 paired row bootstrap resamples, and I² summarized heterogeneity "
        "across calendar-year strata. These descriptors quantify uncertainty and temporal "
        "instability in the released prediction panel but do not fully account for serial "
        "dependence. Planned prospective analyses will therefore use paired block-bootstrap "
        "intervals, report horizon-specific MAE, WAPE, RMSE, continuous ranked probability score, "
        "nominal and empirical coverage, and interval width, and adjust for the prespecified family "
        "of primary comparisons. Peak-warning lead time will be defined before label access by "
        "using a fixed peak threshold and tolerance window. These prospective analyses have not "
        "been executed and are not included in the present Results.",
    )

    figures = {
        33: (
            "figure2_main_error_distribution.png",
            "Hong Kong influenza main-period paired absolute-error distributions and row-wise accuracy-rank counts.",
        ),
        35: (
            "figure3_paired_effect_synthesis.png",
            "Paired horizon-specific WAPE differences between the independent DiseaseX ensemble and paper AWBE.",
        ),
        37: (
            "figure4_main_model_horizon_landscape.png",
            "Main-period model-by-horizon WAPE landscape across standalone models and orchestration.",
        ),
        44: (
            "figure5_postcovid_generalization_landscape.png",
            "Post-COVID external generalization landscape across model families and forecast bands.",
        ),
    }
    for index, (filename, alt_text) in figures.items():
        replace_picture(paragraphs[index], FIGURE_DIR / filename, alt_text)

    captions = {
        34: "Supplementary Fig. 1 Hong Kong influenza main-period paired error distributions and row-wise rank counts. a, Absolute-error distributions for the constant baseline, the DiseaseX independent ensemble, paper AWBE and the validation-locked KG horizon router across 4,968 common forecasts; the display is capped at the pooled 99th percentile, while all rows are retained in the statistics. b, Counts of best or tied-best, middle and worst row-wise ranks. The router uses DiseaseX at horizon 0 and paper AWBE at horizons 1–8 and is not an independent DiseaseX model.",
        36: "Supplementary Fig. 2 Paired evidence synthesis for the DiseaseX independent ensemble versus paper AWBE. a, Horizon-specific differences in WAPE with 95% intervals from 2,500 paired row bootstrap resamples; negative values favour DiseaseX. b, Relative bootstrap precision. c, I² heterogeneity across calendar-year strata. d, Pairwise WAPE differences among the focal methods. The intervals and heterogeneity estimates are exploratory descriptors of the released prediction panel.",
        38: "Supplementary Fig. 3 Main-period model–horizon performance landscape across 4,968 paired forecasts. Panels report the top model–horizon combinations, the complete WAPE heat map, model-family composition, the best standalone method by horizon, family-by-horizon performance and forecast-band summaries. WAPE is the primary point-forecast metric. The orchestration family is shown separately because the KG horizon router retains paper AWBE at horizons 1–8.",
        45: "Supplementary Fig. 4 Post-COVID external generalization landscape for 423 matched rows from 5 March 2023 to 21 January 2024. Panels report the top model–horizon combinations, standalone winners, evaluation-row composition, family-by-forecast-band WAPE, DiseaseX absolute-error distributions and WAPE improvement relative to the constant baseline. All configurations were frozen before external evaluation. The KG horizon router did not outperform paper AWBE and is paper-dependent rather than an independent DiseaseX replacement.",
    }
    for index, caption in captions.items():
        replace_text(paragraphs[index], caption)

    doc.core_properties.title = (
        "DiseaseX Nexus: Self-improving LLM agents for adaptive infectious-disease forecasting"
    )
    doc.core_properties.subject = "Manuscript with WAPE-centred comparison figures"
    doc.save(OUTPUT_MANUSCRIPT)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)


def style_script_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, RGBColor(0x1F, 0x4D, 0x78), 10, 5),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(header.add_run("DiseaseX Nexus | 中文逐段讲稿"), size=9, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(footer.add_run("第 "), size=9, color=MUTED)
    add_page_field(footer)
    set_run_font(footer.add_run(" 页"), size=9, color=MUTED)


def add_title_block(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    set_run_font(kicker.add_run("论文汇报讲稿"), size=10, color=GOLD, bold=True)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    title_run = title.add_run("DiseaseX Nexus")
    title_run.add_break()
    title_run.add_text("多病种自适应与泛化能力")
    set_run_font(
        title_run,
        size=26,
        color=DARK_BLUE,
        bold=True,
    )
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(22)
    set_run_font(subtitle.add_run("按论文段落顺序编排的中文讲解稿"), size=14, color=BLUE)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(48)
    set_run_font(
        meta.add_run("建议时长：30–40 分钟 | 主指标：WAPE | 更新日期：2026-08-28"),
        size=10.5,
        color=MUTED,
    )
    lead = doc.add_paragraph()
    lead.paragraph_format.left_indent = Inches(0.2)
    lead.paragraph_format.right_indent = Inches(0.2)
    lead.paragraph_format.space_before = Pt(8)
    lead.paragraph_format.space_after = Pt(8)
    shade_paragraph(lead, LIGHT_FILL)
    set_run_font(lead.add_run("核心主线："), bold=True, color=DARK_BLUE)
    set_run_font(
        lead.add_run(
            "本论文的贡献不是宣称一个模型适用于所有病种，而是构建一套能够识别不同病种、"
            "数据语义、时间尺度和证据条件，并据此选择、约束、评估和更新模型的可信自适应机制。"
            "现有结果支持机制的可迁移性与可审计性，但尚不支持稳定的跨时期预测优势。"
        )
    )
    doc.add_page_break()


def add_segment(doc: Document, title: str, anchor: str, speech: str) -> None:
    doc.add_paragraph(title, style="Heading 3")
    source = doc.add_paragraph()
    source.paragraph_format.space_after = Pt(3)
    set_run_font(source.add_run("对应原文："), size=9.5, color=MUTED, bold=True)
    set_run_font(source.add_run(anchor), size=9.5, color=MUTED, italic=True)
    paragraph = doc.add_paragraph()
    set_run_font(paragraph.add_run("讲述："), bold=True, color=DARK_BLUE)
    set_run_font(paragraph.add_run(speech))


def build_script() -> None:
    doc = Document()
    style_script_document(doc)
    add_title_block(doc)
    for section_index, (section_title, segments) in enumerate(SECTIONS):
        if section_index > 0:
            doc.add_section(WD_SECTION.NEW_PAGE)
        doc.add_paragraph(section_title, style="Heading 1")
        for title, anchor, speech in segments:
            add_segment(doc, title, anchor, speech)
    doc.core_properties.title = "DiseaseX Nexus 中文逐段讲稿：多病种自适应与泛化能力"
    doc.core_properties.subject = "Paper presentation script"
    doc.core_properties.author = ""
    doc.save(OUTPUT_SCRIPT)


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    required = [
        SOURCE,
        FIGURE_DIR / "figure2_main_error_distribution.png",
        FIGURE_DIR / "figure3_paired_effect_synthesis.png",
        FIGURE_DIR / "figure4_main_model_horizon_landscape.png",
        FIGURE_DIR / "figure5_postcovid_generalization_landscape.png",
    ]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError(f"Missing required files: {missing}")
    build_manuscript()
    build_script()
    print(OUTPUT_MANUSCRIPT)
    print(OUTPUT_SCRIPT)


if __name__ == "__main__":
    main()
