from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"F:\作业\aibola")
SOURCE_ROOT = ROOT / "reports" / "nc_submission_revision_20260825_wape_four_path"
SOURCE_MAIN = SOURCE_ROOT / "DiseaseX_Nexus_NC_Manuscript_v11_WAPE_FourPath_20260825.docx"
SOURCE_SUPP = SOURCE_ROOT / "Supplementary_Information_v6_WAPE_FourPath_20260825.docx"
BOX_ROOT = ROOT / "reports" / "boxplot_wape_review_20260826"
BOX_FIGURE = BOX_ROOT / "Candidate_Fig_WAPE_Boxplots_Four_Path.png"
BOX_STATS = BOX_ROOT / "Candidate_Fig_WAPE_Boxplots_Five_Number_Summary.csv"
BOX_AUDIT = BOX_ROOT / "Candidate_Fig_WAPE_Boxplots_Audit.json"
BOX_SCRIPT = BOX_ROOT / "generate_wape_boxplot.py"

OUTPUT_ROOT = ROOT / "reports" / "nc_submission_revision_20260826_wape_boxplot"
OUTPUT_MAIN = OUTPUT_ROOT / "DiseaseX_Nexus_NC_Manuscript_v12_WAPE_Boxplot_20260826.docx"
OUTPUT_SUPP = OUTPUT_ROOT / "Supplementary_Information_v7_WAPE_Boxplot_20260826.docx"

FIGURE_CAPTION = (
    "Supplementary Fig. 8 | Error distributions and WAPE robustness in the four-path "
    "Hong Kong influenza benchmark. a, Box plots of row-level normalized absolute error "
    "for 1,104 matched H0-H1 rows per method. Errors were divided by the common mean "
    "target, so the group mean equals WAPE and remains defined near zero; the y axis is "
    "logarithmic. b, Box plots of the nine horizon-specific WAPE values from H0 to H8, "
    "with every horizon shown and diamonds marking aggregate WAPE. Boxes span Q1-Q3, "
    "centre lines show medians and whiskers extend to 1.5 x IQR. Lower values are better. "
    "The KG merged route uses paper AWBE at H1-H8 and is not paper-free."
)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def find_paragraph(doc: Document, prefix: str) -> Paragraph:
    matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph beginning {prefix!r}; found {len(matches)}")
    return matches[0]


def insert_after(paragraph: Paragraph, *, text: str = "", style: str | None = None) -> Paragraph:
    new_paragraph = paragraph._parent.add_paragraph()
    paragraph._p.addnext(new_paragraph._p)
    if style is not None:
        new_paragraph.style = style
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def add_alt_text(paragraph: Paragraph, *, title: str, description: str) -> None:
    doc_pr_nodes = paragraph._p.xpath(".//wp:docPr")
    if len(doc_pr_nodes) != 1:
        raise ValueError(f"Expected one wp:docPr node, found {len(doc_pr_nodes)}")
    doc_pr = doc_pr_nodes[0]
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def update_main() -> dict[str, object]:
    doc = Document(SOURCE_MAIN)
    results = find_paragraph(doc, "Figure 6 consolidates the four paths using WAPE")
    results.add_run(
        " The corresponding five-number error distributions and all horizon-specific "
        "WAPE points are shown in Supplementary Fig. 8."
    )

    methods = find_paragraph(doc, "The primary point-forecast metric was weighted absolute percentage error")
    methods.add_run(
        " For Supplementary Fig. 8a, row-level absolute errors were divided by the common "
        "mean target across the 1,104 matched H0-H1 rows; this zero-safe normalization "
        "makes the group mean equal to WAPE. Supplementary Fig. 8b treats the nine "
        "horizon-specific WAPE values from H0 to H8 as the displayed distribution. "
        "Boxes span the first to third quartiles, centre lines indicate medians and "
        "whiskers extend to 1.5 times the interquartile range; all horizon points are shown."
    )

    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_MAIN)
    return {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "images": len(doc.inline_shapes),
        "supplementary_fig_8_mentions": sum("Supplementary Fig. 8" in p.text for p in doc.paragraphs),
    }


def update_supplementary() -> dict[str, object]:
    doc = Document(SOURCE_SUPP)
    note_result = find_paragraph(doc, "For H0")
    note_result.add_run(
        " Supplementary Fig. 8 reports the paired error distributions and the spread of "
        "horizon-specific WAPE values without replacing the exact aggregate estimates."
    )

    note_fig3 = next(p for p in doc.paragraphs if p.text.startswith("Supplementary Fig. 3 | Post-COVID external WAPE"))
    note_fig5 = insert_after(
        note_fig3,
        text=(
            "Supplementary Fig. 8 | Four-path box plots of H0-H1 row-level normalized "
            "absolute error and H0-H8 horizon-specific WAPE. Boxes show Q1-Q3, medians "
            "and 1.5 x IQR whiskers; lower values are better."
        ),
        style=note_fig3.style.name,
    )
    note_fig5.paragraph_format.keep_together = True

    fig4_caption = find_paragraph(doc, "Supplementary Fig. 4 | Evaluation identity changes")
    image_paragraph = insert_after(fig4_caption, style="Normal")
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.page_break_before = True
    image_paragraph.paragraph_format.keep_with_next = True
    image_paragraph.paragraph_format.space_before = 0
    image_paragraph.paragraph_format.space_after = 0
    run = image_paragraph.add_run()
    run.add_picture(str(BOX_FIGURE), width=Inches(6.45))
    add_alt_text(
        image_paragraph,
        title="Supplementary Figure 8: Four-path WAPE box plots",
        description=(
            "Two-panel box-plot comparison of Baseline, independent DiseaseX, paper AWBE "
            "and the KG merged route. Panel A shows H0-H1 row-level normalized absolute "
            "errors; panel B shows horizon-specific WAPE from H0 to H8."
        ),
    )

    archive_caption = insert_after(image_paragraph, text=FIGURE_CAPTION, style="Caption")
    archive_caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
    archive_caption.paragraph_format.keep_together = True
    archive_caption.paragraph_format.space_before = 0
    archive_caption.paragraph_format.space_after = Pt(6)
    archive_caption.paragraph_format.line_spacing = 1.0
    for caption_run in archive_caption.runs:
        caption_run.font.name = "Arial"
        caption_run.font.size = Pt(8.5)
        caption_run.font.bold = False
        caption_run.font.color.rgb = RGBColor(31, 41, 55)

    table_archive = find_paragraph(doc, "Consolidated table archive")
    table_archive.paragraph_format.page_break_before = True

    doc.save(OUTPUT_SUPP)
    return {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "images": len(doc.inline_shapes),
        "supplementary_fig_8_mentions": sum("Supplementary Fig. 8" in p.text for p in doc.paragraphs),
    }


def copy_evidence() -> None:
    figures = OUTPUT_ROOT / "Figures"
    evidence = OUTPUT_ROOT / "Evidence"
    if figures.exists() or evidence.exists():
        raise FileExistsError("Output subdirectories already exist; refusing to overwrite")
    shutil.copytree(SOURCE_ROOT / "Figures", figures)
    shutil.copytree(SOURCE_ROOT / "Evidence", evidence)
    shutil.copy2(BOX_FIGURE, figures / "Supplementary_Fig_8_WAPE_Boxplot.png")
    shutil.copy2(BOX_STATS, evidence / BOX_STATS.name)
    shutil.copy2(BOX_AUDIT, evidence / BOX_AUDIT.name)
    shutil.copy2(BOX_SCRIPT, evidence / BOX_SCRIPT.name)


def write_manifest() -> None:
    manifest = OUTPUT_ROOT / "MANIFEST.sha256"
    lines = []
    for path in sorted(p for p in OUTPUT_ROOT.rglob("*") if p.is_file() and p != manifest):
        rel = path.relative_to(OUTPUT_ROOT).as_posix()
        lines.append(f"{sha256(path)}  {rel}")
    manifest.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    required = [SOURCE_MAIN, SOURCE_SUPP, BOX_FIGURE, BOX_STATS, BOX_AUDIT, BOX_SCRIPT]
    missing = [str(path) for path in required if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"Missing required inputs: {missing}")
    if OUTPUT_ROOT.exists():
        raise FileExistsError(f"Output directory already exists: {OUTPUT_ROOT}")

    source_hashes = {path.name: sha256(path) for path in required}
    OUTPUT_ROOT.mkdir(parents=True)
    copy_evidence()
    main_stats = update_main()
    supp_stats = update_supplementary()

    validation = {
        "status": "WAPE_BOXPLOT_INSERTED_PENDING_RENDER_QA",
        "source_hashes": source_hashes,
        "source_main_unchanged": sha256(SOURCE_MAIN) == source_hashes[SOURCE_MAIN.name],
        "source_supplementary_unchanged": sha256(SOURCE_SUPP) == source_hashes[SOURCE_SUPP.name],
        "boxplot_sha256": sha256(BOX_FIGURE),
        "main": main_stats,
        "supplementary": supp_stats,
        "expected_main_images": 11,
        "expected_supplementary_images": 15,
        "paper_scope": "comparison_experiment_only",
    }
    (OUTPUT_ROOT / "FINAL_VALIDATION.json").write_text(
        json.dumps(validation, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (OUTPUT_ROOT / "CHANGELOG.md").write_text(
        "# WAPE box-plot insertion\n\n"
        "- Preserved the existing main Figure 6 exact WAPE comparison.\n"
        "- Added a main-text reference to Supplementary Fig. 8.\n"
        "- Added the box-plot statistical definition to the comparison Methods section.\n"
        "- Inserted the Python-generated WAPE box plot into the Supplementary figure archive.\n"
        "- Added the reproducible Python script, five-number summary and audit JSON.\n"
        "- No non-comparison scientific claims were changed.\n",
        encoding="utf-8",
    )
    write_manifest()
    print(json.dumps(validation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
