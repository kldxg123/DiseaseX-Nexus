from __future__ import annotations

import hashlib
import json
import shutil
import zipfile
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt


ROOT = Path(r"F:\作业\aibola")
SOURCE_MAIN = ROOT / "DiseaseX_Nexus_NC_Manuscript_v10_English_Polished.docx"
SOURCE_SUPP = (
    ROOT
    / "reports"
    / "nc_submission_revision_20260824"
    / "NC_Submission_Final_v5_20260824_ComplexSupplementaryFigures"
    / "Supplementary"
    / "Supplementary_Information_Updated_v5_20260824.docx"
)
OLD_PRED = ROOT / "work" / "natcomm_wape_optimization_20260825" / "outputs" / "test_paired_predictions.csv"
OLD_ROUTE = ROOT / "work" / "natcomm_wape_optimization_20260825" / "outputs" / "router_paired_predictions.csv"
POSTCOVID = ROOT / "work" / "natcomm_wape_optimization_20260825" / "outputs" / "postcovid_paired_predictions.csv"
NEW_DIR = (
    ROOT
    / "reports"
    / "hk_flu_independent_short_horizon_20260825T112248Z"
    / "extracted_final"
    / "hk_flu_independent_short_horizon_20260825T112248Z"
)
NEW_PRED = NEW_DIR / "04_H0_H8_PAIRED_PREDICTIONS.csv"

OUT = ROOT / "reports" / "nc_submission_revision_20260825_wape_four_path"
FIG_DIR = OUT / "Figures"
EVIDENCE_DIR = OUT / "Evidence"
MAIN_OUT = OUT / "DiseaseX_Nexus_NC_Manuscript_v11_WAPE_FourPath_20260825.docx"
SUPP_OUT = OUT / "Supplementary_Information_v6_WAPE_FourPath_20260825.docx"

COLORS = {
    "Baseline": "#7A869A",
    "DiseaseX": "#E69F00",
    "Paper AWBE": "#1F77B4",
    "KG merged route": "#7B2CBF",
}
ORDER = list(COLORS)


def metric_frame(df: pd.DataFrame, prediction_col: str) -> dict[str, float | int]:
    paired = df[["true", prediction_col]].dropna()
    err = paired["true"] - paired[prediction_col]
    denom = paired["true"].abs().sum()
    wape = float(err.abs().sum() / denom)
    return {
        "n": int(len(paired)),
        "wape": wape,
        "agreement_pct": max(0.0, 1.0 - wape) * 100.0,
        "mae": float(err.abs().mean()),
        "rmse": float(np.sqrt(np.mean(np.square(err)))),
    }


def prepare_data() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    old = pd.read_csv(OLD_PRED)
    route = pd.read_csv(OLD_ROUTE)
    new = pd.read_csv(NEW_PRED)
    keys = ["origin", "date", "week_ahead"]
    required_old = keys + ["true", "Baseline", "AWBE", "KG_ADAPTIVE_ENSEMBLE"]
    required_new = keys + ["true", "recursive_hist_poisson"]
    merged = old[required_old].merge(
        route[keys + ["KG_HORIZON_ROUTER"]], on=keys, validate="one_to_one"
    ).merge(
        new[required_new], on=keys, suffixes=("", "_new"), validate="one_to_one"
    )
    if not np.allclose(merged["true"], merged["true_new"], equal_nan=True):
        raise RuntimeError("Truth columns do not match across archived paired panels")
    merged = merged.drop(columns=["true_new"])
    merged = merged.rename(
        columns={
            "Baseline": "Baseline",
            "recursive_hist_poisson": "DiseaseX",
            "AWBE": "Paper AWBE",
            "KG_HORIZON_ROUTER": "KG merged route",
        }
    )

    rows: list[dict[str, float | int | str]] = []
    scopes = [("H0", 0, 0), ("H1", 1, 1), ("H0-H1", 0, 1), ("H0-H8", 0, 8)]
    for scope, lo, hi in scopes:
        subset = merged[merged["week_ahead"].between(lo, hi)]
        for method in ORDER:
            rows.append({"scope": scope, "method": method, **metric_frame(subset, method)})
    metrics = pd.DataFrame(rows)

    by_horizon: list[dict[str, float | int | str]] = []
    for horizon in range(9):
        subset = merged[merged["week_ahead"] == horizon]
        for method in ORDER:
            by_horizon.append(
                {"horizon": horizon, "method": method, **metric_frame(subset, method)}
            )
    horizon_metrics = pd.DataFrame(by_horizon)

    post = pd.read_csv(POSTCOVID)
    post_rows = []
    post_mapping = {
        "Baseline": "Baseline",
        "DiseaseX": "KG_ADAPTIVE_ENSEMBLE",
        "Paper AWBE": "AWBE",
        "KG merged route": "KG_HORIZON_ROUTER",
    }
    for label, col in post_mapping.items():
        post_rows.append({"scope": "post-COVID", "method": label, **metric_frame(post, col)})
    post_metrics = pd.DataFrame(post_rows)
    return merged, metrics, horizon_metrics, post_metrics


def set_plot_style() -> None:
    sns.set_theme(style="whitegrid", context="talk")
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "axes.titleweight": "bold",
            "axes.labelsize": 14,
            "axes.titlesize": 17,
            "xtick.labelsize": 12,
            "ytick.labelsize": 12,
            "legend.fontsize": 11,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
        }
    )


def annotate_bars(ax, bars, suffix="%") -> None:
    for bar in bars:
        value = bar.get_width()
        ax.text(
            value + max(ax.get_xlim()[1] * 0.012, 0.4),
            bar.get_y() + bar.get_height() / 2,
            f"{value:.2f}{suffix}",
            va="center",
            ha="left",
            fontsize=12,
            fontweight="bold",
        )


def make_main_figure(metrics: pd.DataFrame, path: Path) -> None:
    set_plot_style()
    fig, axes = plt.subplots(1, 2, figsize=(20, 7.84), gridspec_kw={"wspace": 0.32})
    for ax, scope, title, xmax in [
        (axes[0], "H0-H1", "A  Primary short-horizon comparison (H0-H1)", 42),
        (axes[1], "H0-H8", "B  Mandatory full-horizon disclosure (H0-H8)", 86),
    ]:
        data = metrics[metrics.scope == scope].set_index("method").loc[ORDER].reset_index()
        vals = data.wape.to_numpy() * 100
        bars = ax.barh(ORDER[::-1], vals[::-1], color=[COLORS[x] for x in ORDER[::-1]], height=0.58)
        annotate_bars(ax, bars)
        ax.axvline(30, color="#B22222", ls="--", lw=1.8)
        ax.text(30.4, 3.42, "30% engineering gate", color="#8B1A1A", fontsize=10, va="top")
        ax.set_xlim(0, xmax)
        ax.set_xlabel("WAPE (%) — lower is better")
        ax.set_title(title, loc="left")
        ax.grid(axis="x", alpha=0.28)
        ax.grid(axis="y", visible=False)
        for spine in ["top", "right", "left"]:
            ax.spines[spine].set_visible(False)
    fig.suptitle(
        "Hong Kong influenza benchmark: four-path WAPE comparison",
        x=0.5,
        y=0.995,
        fontsize=21,
        fontweight="bold",
    )
    fig.text(
        0.5,
        0.015,
        "All methods use identical origin-horizon-target rows. DiseaseX is the independent surveillance-history model. "
        "The KG merged route uses a DiseaseX ensemble at H0 and paper AWBE at H1-H8; it is not an independent paper-free model.",
        ha="center",
        fontsize=10.5,
    )
    fig.subplots_adjust(left=0.13, right=0.975, top=0.88, bottom=0.16)
    fig.savefig(path, dpi=300, facecolor="white")
    plt.close(fig)


def make_supp_figure_1(metrics: pd.DataFrame, horizon: pd.DataFrame, path: Path) -> None:
    set_plot_style()
    fig = plt.figure(figsize=(16.5, 11))
    gs = fig.add_gridspec(2, 2, height_ratios=[1, 1.18], hspace=0.34, wspace=0.28)
    ax1 = fig.add_subplot(gs[0, 0])
    ax2 = fig.add_subplot(gs[0, 1])
    ax3 = fig.add_subplot(gs[1, :])

    short = metrics[metrics.scope.isin(["H0", "H1", "H0-H1"])].copy()
    sns.barplot(data=short, x="scope", y=short.wape * 100, hue="method", hue_order=ORDER, palette=COLORS, ax=ax1)
    ax1.axhline(30, color="#B22222", ls="--", lw=1.5)
    ax1.set_title("A  Short-horizon WAPE")
    ax1.set_xlabel("Forecast scope")
    ax1.set_ylabel("WAPE (%)")
    ax1.legend(frameon=False, fontsize=9, loc="upper left")

    full = metrics[metrics.scope == "H0-H8"].set_index("method").loc[ORDER].reset_index()
    bars = ax2.barh(ORDER[::-1], (full.wape.to_numpy() * 100)[::-1], color=[COLORS[x] for x in ORDER[::-1]])
    ax2.axvline(30, color="#B22222", ls="--", lw=1.5)
    ax2.set_xlim(0, 86)
    ax2.set_title("B  Full-horizon WAPE")
    ax2.set_xlabel("WAPE (%)")
    annotate_bars(ax2, bars)
    ax2.grid(axis="y", visible=False)

    for method in ORDER:
        d = horizon[horizon.method == method]
        ax3.plot(d.horizon, d.wape * 100, marker="o", ms=5, lw=2.2, color=COLORS[method], label=method)
    ax3.axhline(30, color="#B22222", ls="--", lw=1.5)
    ax3.set_xticks(range(9), [f"H{x}" for x in range(9)])
    ax3.set_ylabel("WAPE (%)")
    ax3.set_xlabel("Forecast horizon")
    ax3.set_title("C  Horizon-specific WAPE reveals the independent long-horizon gap")
    ax3.legend(ncol=4, frameon=False, loc="upper left")
    ax3.set_ylim(0, max(115, horizon.wape.max() * 105))
    fig.suptitle("Four-path paired benchmark with WAPE as the primary metric", fontsize=21, fontweight="bold", y=0.995)
    fig.text(
        0.5,
        0.012,
        "Retrospective author-output benchmark; lower WAPE is better. Agreement = max(0, 1 − WAPE) × 100% and is not classification accuracy.",
        ha="center",
        fontsize=10.5,
    )
    fig.subplots_adjust(left=0.10, right=0.97, top=0.92, bottom=0.12)
    fig.savefig(path, dpi=200, facecolor="white")
    plt.close(fig)


def make_supp_figure_2(horizon: pd.DataFrame, path: Path) -> None:
    set_plot_style()
    matrix = (
        horizon.assign(wape_pct=horizon.wape * 100)
        .pivot(index="method", columns="horizon", values="wape_pct")
        .loc[ORDER]
    )
    fig, axes = plt.subplots(2, 1, figsize=(16.5, 11), gridspec_kw={"height_ratios": [1.15, 1], "hspace": 0.36})
    sns.heatmap(
        matrix,
        annot=True,
        fmt=".1f",
        cmap="YlOrRd",
        cbar_kws={"label": "WAPE (%)"},
        linewidths=0.6,
        linecolor="white",
        ax=axes[0],
    )
    axes[0].set_title("A  Horizon-specific paired WAPE matrix")
    axes[0].set_xlabel("Forecast horizon")
    axes[0].set_ylabel("")

    paper = matrix.loc["Paper AWBE"]
    gaps = matrix.subtract(paper, axis=1)
    for method in ["Baseline", "DiseaseX", "KG merged route"]:
        axes[1].plot(
            range(9), gaps.loc[method], marker="o", lw=2.3, color=COLORS[method], label=f"{method} − paper AWBE"
        )
    axes[1].axhline(0, color="black", lw=1)
    axes[1].set_xticks(range(9), [f"H{x}" for x in range(9)])
    axes[1].set_ylabel("Absolute WAPE gap (percentage points)")
    axes[1].set_xlabel("Forecast horizon")
    axes[1].set_title("B  WAPE gap relative to paper AWBE")
    axes[1].legend(frameon=False, ncol=3, loc="upper left")
    fig.suptitle("WAPE decomposition across horizons 0–8", fontsize=21, fontweight="bold", y=0.995)
    fig.subplots_adjust(left=0.10, right=0.97, top=0.92, bottom=0.12)
    fig.savefig(path, dpi=200, facecolor="white")
    plt.close(fig)


def make_supp_figure_3(post: pd.DataFrame, path: Path) -> None:
    set_plot_style()
    data = post.set_index("method").loc[ORDER].reset_index()
    fig, axes = plt.subplots(1, 2, figsize=(16.5, 11), gridspec_kw={"width_ratios": [1.15, 1], "wspace": 0.30})
    bars = axes[0].barh(ORDER[::-1], (data.wape.to_numpy() * 100)[::-1], color=[COLORS[x] for x in ORDER[::-1]])
    annotate_bars(axes[0], bars)
    axes[0].set_xlim(0, 96)
    axes[0].set_xlabel("WAPE (%)")
    axes[0].set_title("A  Post-COVID external WAPE")
    axes[0].grid(axis="y", visible=False)

    gaps = data.copy()
    paper_wape = float(gaps.loc[gaps.method == "Paper AWBE", "wape"].iloc[0])
    gaps["gap"] = (gaps.wape - paper_wape) * 100
    bars2 = axes[1].barh(
        ORDER[::-1],
        gaps.set_index("method").loc[ORDER[::-1], "gap"],
        color=[COLORS[x] for x in ORDER[::-1]],
    )
    axes[1].axvline(0, color="black", lw=1)
    axes[1].set_xlabel("WAPE gap versus paper AWBE (percentage points)")
    axes[1].set_title("B  External generalization gap")
    for bar in bars2:
        val = bar.get_width()
        axes[1].text(val + (0.8 if val >= 0 else -0.8), bar.get_y() + bar.get_height() / 2, f"{val:+.2f}", va="center", ha="left" if val >= 0 else "right", fontsize=11)
    fig.suptitle("Post-COVID external evaluation of the legacy four-path benchmark", fontsize=20, fontweight="bold", y=0.995)
    fig.text(
        0.5,
        0.02,
        "The new surveillance-only recursive DiseaseX model was not evaluated in this external window; the DiseaseX row here is the legacy adaptive ensemble.",
        ha="center",
        fontsize=10.5,
    )
    fig.subplots_adjust(left=0.10, right=0.97, top=0.92, bottom=0.12)
    fig.savefig(path, dpi=200, facecolor="white")
    plt.close(fig)


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def replace_paragraph(paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")


def replace_image_paragraph(paragraph, image_path: Path) -> None:
    extent = paragraph._p.xpath(".//wp:inline/wp:extent")
    if extent:
        width = Emu(int(extent[0].get("cx")))
        height = Emu(int(extent[0].get("cy")))
    else:
        width, height = Emu(6_500_000), Emu(3_600_000)
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=width, height=height)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def format_table(table, header_fill: str = "D9E2F3") -> None:
    table.autofit = True
    for row_i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_i == 0:
                shade_cell(cell, header_fill)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
                    run.font.size = Pt(8.5 if len(table.columns) >= 6 else 9)
                    if row_i == 0:
                        run.bold = True


def format_caption_cell(doc: Document, cell) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.style = doc.styles["Normal"]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            run.font.name = "Arial"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
            run.font.size = Pt(8.5)
            run.bold = False


def set_table_data(table, headers: list[str], rows: list[list[str]]) -> None:
    while len(table.columns) < len(headers):
        table.add_column(Emu(1_200_000))
    while len(table.rows) < len(rows) + 1:
        table.add_row()
    while len(table.rows) > len(rows) + 1:
        table._tbl.remove(table.rows[-1]._tr)
    for c, value in enumerate(headers):
        table.cell(0, c).text = value
    for r, values in enumerate(rows, 1):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    format_table(table)


def metric_lookup(metrics: pd.DataFrame, scope: str, method: str, key: str) -> float:
    return float(metrics[(metrics.scope == scope) & (metrics.method == method)][key].iloc[0])


def update_main(metrics: pd.DataFrame, fig_main: Path, fig1: Path, fig2: Path, fig3: Path) -> None:
    doc = Document(SOURCE_MAIN)
    replacements = {
        2: (
            "Forecasting systems for infectious diseases must adapt to changes in pathogens, reporting regimes and data availability. "
            "However, model selection and updating are often fixed, poorly documented or difficult to audit. We present DiseaseX Nexus, "
            "a local-first framework in which LLM agents retrieve evidence and propose bounded model configurations; a versioned knowledge graph "
            "links diseases, data profiles, models and parameter provenance; and deterministic code executes backtests and forecasts only after approval "
            "through an explicit authorization contract. The implemented snapshot comprised 1,574 full-text versions, 4,502 evidence spans, 3,791 assertions, "
            "26 data profiles and 17 execution contracts. A Brazil dengue case study demonstrates data profiling, evidence-scored model admission and a deliberately "
            "deferred execution state. In a matched-pair Hong Kong influenza benchmark, an independently trained DiseaseX model achieved a short-horizon H0–H1 WAPE "
            "of 30.43%, compared with 30.22% for the paper AWBE, while the validation-locked KG merged route achieved 29.15%. Across H0–H8, however, the independent "
            "DiseaseX WAPE increased to 57.22%, compared with 34.18% for AWBE and 33.94% for the merged route. The merged route reuses AWBE beyond H0 and is therefore "
            "not an independent paper-free model. Governance-aware evaluation further showed that the only fresh blind test failed the prespecified error gates. "
            "Together, these results establish an auditable framework for agent orchestration and feedback while preserving explicit limits on predictive claims."
        ),
        12: (
            "Here, we evaluate the claims that the current implementation can support. We first quantify the knowledge and contract assets and then trace a non-executed "
            "Brazil dengue recommendation from live observations to an authorization gate. Separately, we evaluate a Hong Kong influenza benchmark against immutable "
            "outputs released by the authors of a 2024 Nature Communications study while preserving the existing governed evaluation tracks. The comparison is reported "
            "as four paths—paper constant baseline, an independent DiseaseX surveillance-history model, paper AWBE and a validation-locked KG merged route—with WAPE as the "
            "primary metric. The independent path reached short-horizon parity but not full-horizon parity; the merged route produced a small retrospective gain because it "
            "routes later horizons to AWBE. The available evidence therefore does not establish state-of-the-art superiority, general multi-pathogen accuracy, calibrated "
            "uncertainty or operational early-warning benefit."
        ),
        29: "A WAPE-centred Hong Kong influenza benchmark separates the independent DiseaseX and KG merged paths",
        30: (
            "We evaluated an executed benchmark, independent of the Brazil case, using the official data and immutable forecast outputs associated with Tsang et al. 9, "
            "a 2024 Nature Communications study of adaptive-weight influenza ensembles in Hong Kong. The primary archive comprised 4,968 common origin–horizon–target rows "
            "over horizons 0–8 from 7 November 2007 to 14 July 2019, excluding the documented surveillance-regime transition. For the new independent-method analysis, H0–H1 "
            "was prespecified as the supported short-horizon scope (1,104 paired rows); H0–H8 was retained as a mandatory secondary disclosure. A second locked external test "
            "comprised 423 matched rows from 5 March 2023 to 21 January 2024."
        ),
        31: (
            "Four paths were compared on identical target keys. The paper constant baseline and paper AWBE were immutable author-released outputs. The independent DiseaseX "
            "model was a recursively refitted seasonal Poisson gradient-boosting comparator trained only on the official weekly influenza-activity history, with deterministic "
            "lag, rolling and seasonal features; no AWBE, SAE, NBE, AWAE or other paper forecast output was used as a feature or candidate. Separately, the KG merged route retained "
            "the earlier development-locked decision: a DiseaseX adaptive ensemble at H0 and paper AWBE at H1–H8. The merged route is therefore an orchestration result rather than "
            "an independent reproduction of the paper method. Test labels were not used to fit the independent model or to alter the archived merged route."
        ),
        32: (
            "WAPE was the primary comparison metric. Over H0–H1, the constant baseline, independent DiseaseX, paper AWBE and KG merged route achieved WAPE values of 34.61%, "
            "30.43%, 30.22% and 29.15%, respectively, on the same 1,104 rows. Thus, the independent DiseaseX model differed from AWBE by only 0.21 percentage points in the "
            "prespecified short-horizon scope; it was better at H0 (24.56% versus 25.98%) and worse at H1 (36.31% versus 34.47%). Across H0–H8, the corresponding WAPE values were "
            "75.94%, 57.22%, 34.18% and 33.94%. The independent path therefore did not retain parity beyond H1, whereas the merged route remained close to AWBE because AWBE supplies "
            "its H1–H8 forecasts. RMSE values for H0–H1 were 2.9039, 2.5868, 2.6230 and 2.5190, respectively; full-horizon RMSE values were 5.9452, 4.4576, 2.7472 and 2.7255. "
            "Agreement is defined as max(0, 1 − WAPE) × 100 and is not classification accuracy (Fig. 6, Supplementary Figs. 1–2 and Supplementary Table 2)."
        ),
        34: (
            "Supplementary Fig. 1 Hong Kong influenza four-path WAPE benchmark. a, WAPE at H0, H1 and H0–H1. b, Mandatory H0–H8 WAPE disclosure. c, Horizon-specific WAPE. "
            "DiseaseX denotes the independently trained surveillance-history model; the KG merged route uses a DiseaseX ensemble at H0 and paper AWBE at H1–H8. Lower values are better."
        ),
        36: (
            "Supplementary Fig. 2 Horizon-specific WAPE matrix and absolute WAPE gaps relative to paper AWBE. The independent DiseaseX method remains competitive only at short horizons; "
            "the KG merged route inherits AWBE at H1–H8. All values use identical paired target keys."
        ),
        38: (
            "Supplementary Fig. 3 Post-COVID external WAPE evaluation on 423 matched rows. This panel retains the legacy adaptive-ensemble comparison because the new independent recursive "
            "DiseaseX model was not evaluated in the external period. The locked KG merged route did not outperform paper AWBE externally."
        ),
        40: (
            "The experiment identity was RETROSPECTIVE_AUTHOR_OUTPUT_BENCHMARK, and the promotion status was SHADOW_ONLY_NOT_PROMOTED. Deterministic replay and the future-label sentinel "
            "passed, whereas database writes, forecast-run creations, promotions and active-pointer changes were all zero. The short-horizon independent result demonstrates that a paper-free "
            "surveillance-history model can approach AWBE at H0–H1; the mandatory H0–H8 disclosure shows that this does not generalize to longer horizons. The merged result demonstrates "
            "horizon-aware orchestration but cannot be interpreted as an independent model because AWBE supplies H1–H8."
        ),
        51: "WAPE-centred comparison of baseline, independent DiseaseX, paper AWBE and the KG merged route",
        52: (
            "Figure 6 consolidates the four paths using WAPE on identical matched rows. In the supported H0–H1 scope, independent DiseaseX (30.43%) was within 0.21 percentage points of paper "
            "AWBE (30.22%), while the KG merged route achieved 29.15%. The full H0–H8 disclosure materially changed the interpretation: independent DiseaseX increased to 57.22%, whereas paper "
            "AWBE and the KG merged route achieved 34.18% and 33.94%, respectively. This figure therefore reports both the short-horizon near-parity result and the retained long-horizon gap. "
            "The KG merged path is labelled separately because it uses AWBE beyond H0; it is not evidence that an entirely independent DiseaseX model surpassed the paper."
        ),
        58: (
            "The Hong Kong influenza benchmark provides a calibrated assessment of predictive capability. The new independent DiseaseX surveillance-history model approached paper AWBE over "
            "H0–H1 (30.43% versus 30.22% WAPE) but diverged over H0–H8 (57.22% versus 34.18%). The KG merged route achieved 29.15% at H0–H1 and 33.94% over H0–H8, but it routes H1–H8 to AWBE "
            "and is therefore a system-level orchestration result rather than an independent alternative to the paper method. The legacy merged route did not retain its small advantage in the "
            "post-COVID external window. These contrasts show why method independence, forecast horizon and evaluation identity must be reported together."
        ),
        60: (
            "This study has five main limitations. First, the primary Brazil case ends at the recommendation stage and does not evaluate forecasts. Second, only one fresh blind test was available, "
            "and it failed. Third, the direct Hong Kong influenza comparison used immutable prediction outputs released by the authors rather than a full independent rerun of the original R/Windows "
            "stack. Although the new paper-free DiseaseX comparator approached AWBE at H0–H1, it did not retain parity at H2–H8; conversely, the KG merged route uses AWBE beyond H0 and did not "
            "generalize in the post-COVID window. Fourth, the knowledge assets constitute a verified snapshot because incremental production construction was incomplete. Fifth, the current evaluation "
            "does not estimate decision utility, fairness, cross-site transportability, probabilistic calibration or prospective early-warning lead time. These limitations preclude claims of clinical or "
            "public-health effectiveness."
        ),
        75: "Hong Kong influenza four-path WAPE benchmark and horizon-aware orchestration",
        76: (
            "We evaluated the Hong Kong influenza benchmark using the official surveillance data and immutable forecast outputs released in association with Tsang et al. 9,25. The archived release "
            "was fixed before analysis, and its integrity was verified against a version-controlled provenance record (Supplementary Note 4 and Supplementary Table 14). The full evaluation comprised "
            "4,968 common origin–horizon–target observations across horizons 0–8; the prespecified short-horizon analysis comprised the H0–H1 subset of 1,104 observations. The external post-COVID "
            "evaluation comprised 423 common observations with target dates from 5 March 2023 to 21 January 2024. Missing forecasts were retained as missing and were not replaced with zero."
        ),
        77: (
            "The independent DiseaseX comparator was recursively refitted using only lagged official weekly influenza activity, seasonal lags, rolling summaries and deterministic calendar terms. It did not "
            "consume paper AWBE or component-model forecasts. Its supported evaluation scope was locked to H0–H1, while H0–H8 was reported as a mandatory limitation analysis. The separate KG merged "
            "route retained the existing development-locked policy: a DiseaseX adaptive ensemble at H0 and author-released AWBE at H1–H8. All four paths—constant baseline, independent DiseaseX, paper "
            "AWBE and KG merged route—were evaluated on identical target keys. The repository revision, artifact checksums, row-level predictions and immutable route records are reported in the "
            "supplementary reproducibility inventory."
        ),
        88: (
            "The Hong Kong influenza benchmark inputs, released predictions and source scripts are available from the official repository associated with Tsang et al. 25. The processed common-pair panels, "
            "independent DiseaseX feature and model pipelines, locked route records, WAPE tables, figures and SHA-256 manifest are included in the reproducibility package of the authors. Brazil dengue "
            "observations were retrieved under the source key who_dengue_global_validated and are represented in the archived data-profile object. [PLACEHOLDER: deposit the exact processed datasets, "
            "source-response manifests and non-sensitive audit objects in a persistent repository; insert the accession identifier and DOI. State any source-specific restrictions on redistribution.]"
        ),
    }
    for idx, text in replacements.items():
        replace_paragraph(doc.paragraphs[idx], text)

    replace_image_paragraph(doc.paragraphs[33], fig1)
    replace_image_paragraph(doc.paragraphs[35], fig2)
    replace_image_paragraph(doc.paragraphs[37], fig3)
    replace_image_paragraph(doc.paragraphs[53], fig_main)
    doc.tables[0].cell(0, 0).text = (
        "Figure 6 | WAPE-centred Hong Kong influenza comparison. The four rows distinguish the paper constant baseline, the independently trained DiseaseX surveillance-history model, paper AWBE and the validation-locked KG merged route. "
        "Panel a reports the prespecified H0–H1 scope; panel b reports the mandatory H0–H8 disclosure. The independent DiseaseX path approaches AWBE only at H0–H1. The merged route uses a DiseaseX ensemble at H0 and AWBE at H1–H8 and is not an independent paper-free model. "
        "All metrics use identical origin–horizon–target rows. Lower WAPE is better; agreement derived from WAPE is not classification accuracy."
    )
    format_caption_cell(doc, doc.tables[0].cell(0, 0))
    doc.save(MAIN_OUT)


def update_supp(metrics: pd.DataFrame, fig_main: Path, fig1: Path, fig2: Path, fig3: Path) -> None:
    doc = Document(SOURCE_SUPP)
    replacements = {
        1: "Updated Hong Kong influenza four-path WAPE benchmark and complete figure/table archive — 25 August 2026",
        2: (
            "Evidence status: retrospective author-output benchmark; shadow only; external generalization not confirmed. The independent DiseaseX model uses surveillance history only and is supported for H0–H1 in this benchmark. "
            "The KG merged route uses paper AWBE at H1–H8. Agreement values are display transformations of WAPE, not classification accuracy."
        ),
        26: "Supplementary Note 4. Nature Communications Hong Kong influenza four-path WAPE benchmark",
        27: (
            "The benchmark was based on Tsang et al. (Nature Communications 15, 8625, 2024; DOI 10.1038/s41467-024-52504-1) and a fixed revision of the official author repository. "
            "The official surveillance data, author-released predictions and repository revision were fixed before analysis. The new independent DiseaseX comparator used only the surveillance-history series; "
            "it did not use AWBE or any author-released component prediction as a feature. File-level checksums and immutable selection records are retained in Supplementary Table 14 and the reproducibility manifest."
        ),
        29: (
            "Four paths were evaluated on identical target keys: the paper constant baseline, an independent DiseaseX recursive seasonal Poisson gradient-boosting comparator, paper AWBE and the existing KG merged route. "
            "The independent model used deterministic lag, rolling and seasonal features and was recursively refitted from past observations only. Its supported scope was H0–H1. The KG merged route retained the earlier development-locked decision: a DiseaseX adaptive ensemble at H0 and AWBE at H1–H8. "
            "Accordingly, the merged route is a system-level orchestration result and not an independent paper-free comparator."
        ),
        30: (
            "For H0–H1, WAPE was 34.61% for the constant baseline, 30.43% for independent DiseaseX, 30.22% for paper AWBE and 29.15% for the KG merged route. Independent DiseaseX was within 0.21 percentage points of AWBE. "
            "For H0–H8, the corresponding WAPE values were 75.94%, 57.22%, 34.18% and 33.94%, showing that independent parity did not persist beyond H1. On the external post-COVID period, the legacy KG merged route remained slightly worse than AWBE (49.35% versus 48.88% WAPE). "
            "The archived state is HK_FLU_INDEPENDENT_H0_H1_PARITY_VERIFIED_FULL_HORIZON_GAP_RETAINED; database writes, forecast-run creation, promotion and active-pointer changes were all zero."
        ),
        31: "Supplementary Table 4 | Four-path paired WAPE metrics for the prespecified H0–H1 scope and mandatory H0–H8 disclosure.",
        32: "Supplementary Fig. 1 | Four-path paired WAPE summary. Panel a reports H0, H1 and H0–H1; panel b reports H0–H8; panel c reports horizon-specific WAPE. Lower is better.",
        33: "Supplementary Fig. 2 | Horizon-specific WAPE matrix and WAPE gaps relative to paper AWBE. Lower values are better; negative gaps favour the comparator.",
        34: "Supplementary Fig. 3 | Post-COVID external WAPE for the legacy four-path benchmark. The new surveillance-only recursive DiseaseX model was not evaluated in this external window; the DiseaseX row denotes the prior adaptive ensemble.",
        71: "Supplementary Fig. 1 | Four-path paired WAPE summary for the main Hong Kong influenza benchmark.",
        74: "Supplementary Fig. 2 | Horizon-specific WAPE matrix and WAPE gaps relative to paper AWBE.",
        77: (
            "Supplementary Fig. 3 | Post-COVID external WAPE for the legacy four-path benchmark. The new surveillance-only recursive DiseaseX model was not evaluated in this external window; the DiseaseX row denotes the prior adaptive ensemble."
        ),
        91: (
            "Figure 6 | WAPE-centred Hong Kong influenza comparison. The four methods are the paper constant baseline, independent DiseaseX, paper AWBE and the validation-locked KG merged route. H0–H1 is the supported short-horizon comparison; H0–H8 is the mandatory limitation disclosure. "
            "The KG merged route uses AWBE at H1–H8 and is not an independent paper-free model."
        ),
    }
    for idx, text in replacements.items():
        replace_paragraph(doc.paragraphs[idx], text)

    replace_image_paragraph(doc.paragraphs[70], fig1)
    replace_image_paragraph(doc.paragraphs[73], fig2)
    replace_image_paragraph(doc.paragraphs[76], fig3)
    replace_image_paragraph(doc.paragraphs[89], fig_main)

    summary_rows: list[list[str]] = []
    for method in ORDER:
        for scope in ["H0-H1", "H0-H8"]:
            r = metrics[(metrics.scope == scope) & (metrics.method == method)].iloc[0]
            summary_rows.append(
                [
                    method,
                    scope,
                    f"{int(r['n']):,}",
                    f"{float(r['wape']) * 100:.2f}",
                    f"{float(r['agreement_pct']):.2f}",
                    f"{float(r['rmse']):.4f}",
                ]
            )
    for table_index in [3, 18]:
        set_table_data(
            doc.tables[table_index],
            ["Method", "Scope", "n", "WAPE (%)", "Agreement (%)", "RMSE"],
            summary_rows,
        )

    wape_rows: list[list[str]] = []
    for scope in ["H0", "H1", "H0-H1", "H0-H8"]:
        wape_rows.append(
            [scope]
            + [f"{metric_lookup(metrics, scope, method, 'wape') * 100:.2f}" for method in ORDER]
        )
    for table_index in [10, 24]:
        set_table_data(
            doc.tables[table_index],
            ["Scope", "Baseline", "DiseaseX", "Paper AWBE", "KG merged route"],
            wape_rows,
        )

    doc.tables[14].cell(0, 0).text = (
        "Figure 6 | Four-path WAPE comparison. Baseline, independent DiseaseX, paper AWBE and KG merged route are reported on common target keys for H0–H1 and H0–H8. Independent DiseaseX approaches AWBE only at H0–H1; the KG merged route uses AWBE at H1–H8."
    )
    format_caption_cell(doc, doc.tables[14].cell(0, 0))
    doc.save(SUPP_OUT)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_evidence(merged: pd.DataFrame, metrics: pd.DataFrame, horizon: pd.DataFrame, post: pd.DataFrame) -> None:
    EVIDENCE_DIR.mkdir(parents=True, exist_ok=True)
    metrics.to_csv(EVIDENCE_DIR / "FOUR_PATH_METRICS.csv", index=False)
    horizon.to_csv(EVIDENCE_DIR / "FOUR_PATH_HORIZON_METRICS.csv", index=False)
    post.to_csv(EVIDENCE_DIR / "POSTCOVID_FOUR_PATH_METRICS.csv", index=False)
    merged.to_csv(EVIDENCE_DIR / "FOUR_PATH_PAIRED_PREDICTIONS_H0_H8.csv", index=False)
    primary = metrics[metrics.scope == "H0-H1"].set_index("method")
    full = metrics[metrics.scope == "H0-H8"].set_index("method")
    summary = {
        "primary_scope": "H0-H1",
        "primary_n": int(primary.iloc[0]["n"]),
        "primary_wape_pct": {m: round(float(primary.loc[m, "wape"]) * 100, 6) for m in ORDER},
        "full_scope": "H0-H8",
        "full_n": int(full.iloc[0]["n"]),
        "full_wape_pct": {m: round(float(full.loc[m, "wape"]) * 100, 6) for m in ORDER},
        "independent_short_horizon_gap_vs_paper_pp": round(
            (float(primary.loc["DiseaseX", "wape"]) - float(primary.loc["Paper AWBE", "wape"])) * 100,
            6,
        ),
        "method_identity": {
            "Baseline": "paper constant baseline",
            "DiseaseX": "independent recursive seasonal Poisson gradient boosting; surveillance history only",
            "Paper AWBE": "immutable author-released paper forecast",
            "KG merged route": "DiseaseX adaptive ensemble at H0; paper AWBE at H1-H8",
        },
        "claim_boundary": "Independent parity is supported only for H0-H1; KG merged route is not paper-free.",
    }
    (EVIDENCE_DIR / "FOUR_PATH_SUMMARY.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "CHANGELOG.md").write_text(
        "# Comparison-only manuscript revision\n\n"
        "- Updated only Hong Kong influenza comparison-related abstract, Results, Discussion, Methods, Data availability, figure captions and supplementary comparison tables/figures.\n"
        "- Added the independent DiseaseX H0-H1 result (WAPE 30.43%) beside paper AWBE (30.22%).\n"
        "- Retained the mandatory H0-H8 limitation (DiseaseX 57.22%; AWBE 34.18%).\n"
        "- Reported the KG merged route separately (29.15% at H0-H1; 33.94% at H0-H8) and explicitly stated that it uses AWBE at H1-H8.\n"
        "- Replaced Figure 6 and Supplementary Figs. 1-3 with English Python-generated WAPE figures.\n"
        "- Rebuilt Supplementary Tables 2 and 4 around four-path common-pair WAPE.\n"
        "- All non-comparison architecture, knowledge-graph and case-study figures were preserved.\n",
        encoding="utf-8",
    )


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    merged, metrics, horizon, post = prepare_data()
    write_evidence(merged, metrics, horizon, post)

    fig_main = FIG_DIR / "Figure_6_WAPE_FourPath.png"
    fig1 = FIG_DIR / "Supplementary_Fig_1_WAPE_FourPath.png"
    fig2 = FIG_DIR / "Supplementary_Fig_2_Horizon_WAPE.png"
    fig3 = FIG_DIR / "Supplementary_Fig_3_PostCOVID_WAPE.png"
    make_main_figure(metrics, fig_main)
    make_supp_figure_1(metrics, horizon, fig1)
    make_supp_figure_2(horizon, fig2)
    make_supp_figure_3(post, fig3)
    update_main(metrics, fig_main, fig1, fig2, fig3)
    update_supp(metrics, fig_main, fig1, fig2, fig3)

    manifest_lines = []
    for path in sorted(p for p in OUT.rglob("*") if p.is_file()):
        if path.name in {"MANIFEST.sha256", "NC_Submission_WAPE_FourPath_20260825.zip"}:
            continue
        manifest_lines.append(f"{sha256(path)}  {path.relative_to(OUT).as_posix()}")
    (OUT / "MANIFEST.sha256").write_text("\n".join(manifest_lines) + "\n", encoding="utf-8")
    print(json.dumps({"main": str(MAIN_OUT), "supp": str(SUPP_OUT), "files": len(manifest_lines)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()




