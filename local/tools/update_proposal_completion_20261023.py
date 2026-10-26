"""Append executed vm201 experiments without replacing frozen earlier results."""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile,ZIP_DEFLATED
from io import BytesIO
import csv,json,re,hashlib,statistics
from lxml import etree as E
from docx import Document
from docx.shared import Inches
from reframe_proposal_baseline_selection_20261023 import styled,table,blocks_to_rows,W,NS,w,serialize,text,replace

ROOT=Path(__file__).resolve().parents[1];DATA=ROOT/'reports/evidence_completion_20261023';OUT=ROOT/'output/proposal_completion_20261023'
SOURCE=ROOT/'output/proposal_real_llm_holdout_20261023/安浩然_开题报告_三病种新增时段与真实LLM实验修订稿_20261023.docx'
DEST=OUT/'安浩然_开题报告_vm201实验补充完整修订稿_20261023.docx'
DISEASES={'Dengue Fever':'登革热','Campylobacterenterosis':'弯曲菌肠炎','Mumps':'腮腺炎'}
ARMS={'profile':'仅画像','graph':'画像与图谱','llm':'画像与大模型','graph_llm':'图谱与大模型联动'}
def load(name):return json.loads((DATA/name).read_text(encoding='utf-8'))
def pct(value):return '—' if value is None else f'{100*value:.2f}%'
def interval(r):return f"[{r['ci95_pp'][0]:.3f}, {r['ci95_pp'][1]:.3f}]"
def caption(s):
    p=styled(s,'caption');p.find(f'{{{W}}}pPr').append(w('keepNext'));return p
def heading(s):
    p=styled(s,'h2');p.find(f'{{{W}}}pPr').append(w('keepNext'));return p

def main():
    OUT.mkdir(exist_ok=True);complete=load('completion.json');assert complete['completed']
    assert load('test_result.json')['success'] and load('ledger_audit.json')['all_1872_complete']
    metrics=load('factorial_metrics.json');cis=load('factorial_confidence_intervals.json');proposals=load('factorial_proposals.json');costs=load('factorial_costs.json')
    availability=load('availability_summary_v2.json');wall=load('wall_budget_results.json')
    assert len(wall)==18 and len(metrics)==36 and len(availability)==27
    def metric(disease,arm,budget=13):return next(r for r in metrics if (r['disease'],r['arm'],r['validation_origins'])==(disease,arm,budget))
    def ci(disease):return next(r for r in cis if r['disease']==disease and r['comparison']=='graph_llm minus profile' and r['block_origins']==8)
    package=Document(SOURCE);p=package.add_paragraph();p.alignment=1;p.add_run().add_picture(str(OUT/'授权状态迁移.png'),width=Inches(5.6))
    buffer=BytesIO();package.save(buffer)
    with ZipFile(buffer) as z:parts={n:z.read(n) for n in z.namelist()}
    doc=E.fromstring(parts['word/document.xml']);body=doc.find(f'{{{W}}}body')
    figure=next(p for p in body if p.tag==f'{{{W}}}p' and p.xpath('.//w:drawing',namespaces=NS));body.remove(figure)
    original_ps=doc.xpath('//w:p',namespaces=NS);refs=[text(p) for p in original_ps if re.match(r'^\[\d+\]',text(p))]
    equations=[E.tostring(e) for e in doc.xpath('//m:oMath',namespaces=NS)]
    teacher=next(p for p in original_ps if text(p).startswith('学术导师对该生开题报告的简要评语'))
    teacher_text=[text(p) for p in teacher.getparent().xpath('.//w:p',namespaces=NS)]
    edits=[]
    def edit(p,s):
        if text(p)!=s:edits.append({'before':text(p),'after':s});replace(p,s)
    for p in original_ps:
        s=text(p)
        if re.search(r'表\s*14(?!\d)',s):edit(p,re.sub(r'表\s*14(?!\d)','表15',s))
        elif s.startswith('本课题研究面向多病种疫情监测的可信智能预测系统。'):
            edit(p,'本课题研究面向多病种疫情监测的可信智能预测系统。系统依据病种数据条件生成候选，将合格基线与推荐模型共同纳入历史滚动评价；简单基线更好时保留基线。原型保存来源、输入版本、模型参数及复核记录。登革热、弯曲菌肠炎和腮腺炎在2020—2022年回顾评价中的WAPE分别为22.09%、28.96%和46.70%。2024年度历史留出及真实大模型实验覆盖三病种，但登革热仍劣于最优简单基线。新增图谱与大模型联合对照、数据可用率敏感性、墙钟预算及持久化授权测试，检验候选校验、拒绝预测与运行追溯。现有结果未证明稳定的图谱或大模型精度增益，也不支持未知病原预测有效或统计等效。预期交付本地系统、代码差分及复核材料，另开展人工审核与前瞻评价。')
        elif s.startswith('本人于2026年10月23日确认项目代码由本人完成，'):
            edit(p,'本人于2026年10月23日确认项目代码由本人完成，后续按模块、接口、修改前后哈希、差分和测试日志登记；验证实现和文字修订使用AI辅助，第三方库、数据及引用算法保留来源。缺少可用Git历史，追溯依据为交付快照与逐行差分。现已完成三病种回顾与历史留出、真实大模型对照、图谱联合消融、可用率与计算预算试验，并新增持久化授权原型。图谱双人审核、真实用户评价及导师确认另行安排；隔离原型测试不代表生产部署验收。')
        elif s.startswith('已完成检索留档、代码与数据身份核查、三病种回顾评价、'):
            edit(p,'已完成检索留档、代码与数据身份核查、三病种回顾评价、配对差值区间、年度嵌套回测、预算原型消融、准入敏感性、222项新增组件测试及42项既有工作流回归。2.8节补充2024年度历史留出、12次服务器真实大模型推理和12项输出校验；2.9节另列6次真实提议、36次联合对照、28项持久化与认证集成测试，以及实际墙钟预算试验。12周影子评价、真人审核和用户评价仍按后续计划执行。生产权限与Docker启动未在本轮验收，不将隔离测试扩大为生产系统结论。')
        elif s.startswith('图谱双人断言审核、目标用户任务评价、负责人或导师确认需要真人参与；'):
            edit(p,'图谱双人断言审核、目标用户任务评价及负责人或导师确认需要真人参与；12周影子评价需要连续新发布资料及真实日历。本轮补充了图谱与大模型联合数值流程、可用率敏感性、墙钟预算、持久化执行边界及状态迁移图，详见2.9节和附录D。生产权限系统接入、Docker启动和用途容差依据仍须分别验收。历史留出不等同于前瞻运行，导师原评语保留，不表示本次修订已经获得新的签字确认。')
        elif s.startswith('本轮完成了新增年度历史留出和真实模型候选提议对照，'):
            edit(p,'本节完成新增年度历史留出和首版真实模型候选提议对照。之后补充的图谱联合消融与持久化执行原型见2.9节和附录D，不替换本节协议、失败提议或指标。两轮均不等于完整生产系统的端到端用户验证，双人断言审核、12周影子评价、真实用户任务研究及导师贡献确认仍分别保留。')
    mainrows=[]
    for disease,zh in DISEASES.items():mainrows.append([zh]+[pct(metric(disease,arm)['wape']) for arm in ARMS]+[f"{ci(disease)['delta_wape_pp']:+.3f}",interval(ci(disease))])
    additions=[heading('2.9 图谱联合对照与持久化执行验证'),
        styled('在已经查看的2024年历史资料上，进一步比较仅画像、画像与图谱、画像与大模型、图谱与大模型联动四组。各组均保留三种常驻基线，并使用两个附加候选；统一采用13个已完整观察的历史原点选择模型。新增提示词明确排除常驻基线，仅针对候选格式作版本化修订，不能据此把旧测试窗口重新称为盲测。6次新模型调用的原始输出均通过候选校验，原12次调用及其8次回退保留不变。'),
        caption('表14 图谱与大模型联合对照的2024年WAPE'),
        table(['病种','仅画像','加图谱','加大模型','图谱与大模型','差值/百分点','差值95%区间'],mainrows,[1200,950,950,1100,1200,1100,1500]),
        styled('注：每组每病种均为52个预测原点、202个相关目标。差值为联动组减仅画像组，区间沿用5,000次原点块自助法，主块长8。当前图谱仅命中登革热的一条条件规则；另外两病种不借用登革热知识，未提供直接适用规则时保留画像策略。三病种的联动组与仅大模型组点预测相同，不能据此断言图谱增益。各差值区间均跨零，也不构成稳定的大模型精度优势。'),
        styled('补充实验完成36次独立数值运行，记录7,272条预测和1,872个已完成执行契约。不同组不共用数值缓存，选择原点数3、8、13分别作为预算诊断，完整成绩与耗时见附录D。准入试验对最近100个历史输入遮蔽30、20或10个值，比较70%、80%、90%阈值；缺失值不补零或插值，不满足条件时拒绝输出。这是受控缺失诊断，不是实际发布滞后资料验证，也不据此在测试集上确定最优阈值。'),
        styled('新持久化原型通过28项测试，覆盖认证请求、角色越权、八项绑定篡改、重启读取、多进程同键执行及进程中断。8个进程对同一请求仅执行一次回调；执行前先保存状态，中断后保留运行中状态并停止自动重试。审核身份由测试夹具提供，不表示真人批准。1、5、30秒硬墙钟预算试验分别记录完成与超时任务，冷启动加载与推理均计时；该路径只生成点预测，未包括概率区间和生产发布，不与完整系统延迟混用。')]
    marker=next(p for p in doc.xpath('//w:p',namespaces=NS) if text(p)=='3 论文提纲');cell=marker.getparent();row=cell.getparent();outer=row.getparent();tcpr=deepcopy(cell.find(f'{{{W}}}tcPr'))
    content=[x for x in cell if x.tag!=f'{{{W}}}tcPr'];idx=content.index(marker);position=list(outer).index(row);outer.remove(row)
    for i,r in enumerate(blocks_to_rows(content[:idx]+additions+content[idx:],tcpr)):outer.insert(position+i,r)
    appendix=[heading('附录D 联合消融与执行边界补充实验'),heading('D.1 协议 图谱和实际模型输出'),
        styled('本轮各实验均在vm201新隔离目录运行，未修改在线服务、模型权重、全局环境或生产数据库。联合对照在运行前登记协议，使用既有4B模型、CPU float32和4线程，温度0、固定种子、最多96个输出token，每病种分别调用有图谱和无图谱两次。由于年度结果已查看且提示词在既有输出问题基础上修订，本轮属于探索性后续验证。图谱采用当前版本，不冒充历史时点可用知识，也未增加未经复核的病种关系。'),
        styled('结构核查覆盖54个节点、86条边、12条规则和19个证据来源；节点标识无重复，边端点与规则证据标识均能解析。该检查只证明结构一致，不代替来源真实性与适用性审核。登革热命中无外部气候输入规则，但规则推荐的是模型家族，并未把这些家族擅自映射为本实验的附加算法。弯曲菌肠炎和腮腺炎没有直接规则，返回空提示，仍允许基线及画像候选参与比较。'),
        caption('表D1 六次新真实调用的输出与图谱命中')]
    rawrows=[]
    for p in proposals:
        if 'raw_response' not in p:continue
        rawrows.append([DISEASES[p['disease']],'有' if p['arm']=='graph_llm' else '无',', '.join(p['ids']),'通过' if p['valid'] else '回退',f"{p['seconds']:.2f}"])
    appendix += [table(['病种','图谱提示','输出的两个附加候选','契约','推理/秒'],rawrows,[1400,1000,3100,900,1100]),
        styled('主选择窗下，有图谱与无图谱大模型的逐项点预测相同，但登革热的附加候选分别为drift4与mean13、drift4与mean8；两种池不等同于逐项选择一定不同。所有候选仍经历史误差选择，不能把候选变化当作精度提升。图谱信息缺少适用规则或可执行映射时，系统保持数值流程而不虚构证据。'),
        heading('D.2 选择预算与实际墙钟预算'),caption('表D2 三种选择原点预算下的四组WAPE')]
    budgetrows=[]
    for disease,zh in DISEASES.items():
        for budget in (3,8,13):budgetrows.append([zh,str(budget)]+[pct(metric(disease,arm,budget)['wape']) for arm in ARMS])
    appendix += [table(['病种','选择原点数','仅画像','加图谱','加大模型','联动组'],budgetrows,[1450,1200,1300,1300,1400,1350]),
        styled('三种配置每个预测原点分别执行15、40、65次选择拟合，随后执行一次选中模型预测；每组均为五候选且保留三种基线。每个年度任务另有260次校准拟合。不同预算的模型提议共用同一次预先输出，数值拟合不跨组缓存；汇总的提议耗时是分别归属成本，不能将三个预算相加解释为三次真实推理，也不把13原点之外较好的成绩事后替换主结果。完整选择轨迹、拟合数及独立数值耗时随证据包保存。'),
        caption('表D3 相同硬墙钟预算下的点预测任务完成情况')]
    wallrows=[]
    for disease,zh in DISEASES.items():
        for cap in (1,5,30):
            a=next(r for r in wall if (r['disease'],r['arm'],r['cap_seconds'])==(disease,'profile',cap));b=next(r for r in wall if (r['disease'],r['arm'],r['cap_seconds'])==(disease,'graph_llm',cap))
            wallrows.append([zh,str(cap),'202' if a['completed_within_budget'] else '超时',f"{a['elapsed_seconds']:.2f}",'202' if b['completed_within_budget'] else '超时',f"{b['elapsed_seconds']:.2f}"])
    appendix += [table(['病种','上限/秒','画像完成目标','画像耗时/秒','联动完成目标','联动耗时/秒'],wallrows,[1400,1000,1500,1300,1500,1300]),
        styled('每个格子使用新子进程，计入导入、冷模型加载、一次真实提议和完整52个原点的点预测；两组使用相同墙钟上限及最多4个CPU线程，不使用GPU。超时即终止该任务的整个子进程组，清理开销可能使观测耗时略超过上限。1秒与5秒下联动组均无完整输出，不能填零误差；30秒下三个病种均完成202个目标。画像组在三个预算均完成。这说明小型数值流程与冷启动大模型的实际成本不同，不支持联动组节省整体时间的结论。每格仅一次、仅一台主机，不据此推广一般吞吐能力；该试验不含区间校准、持久化契约或生产发布。'),
        heading('D.3 数据滞后可用率的受控敏感性'),
        styled('对各病种2024年的49个完整四步原点，遮蔽最近100个输入中的30、20、10个值；五组固定随机排列使三种可用率嵌套。对每个输入再比较70%、80%、90%准入阈值，共6,615条原点—掩码—阈值记录。最新值不可用时拒绝输出；各适配器仅用已观察输入，AR(2)使用完整滞后行拟合，均值与季节适配器缺少必要输入时退出。通过条件的适配器作预先指定等权平均，不按未来误差挑选。'),
        caption('表D4 可用率等于准入阈值时的输出与条件误差')]
    avrows=[]
    for disease,zh in DISEASES.items():
        for available in (70,80,90):
            r=next(r for r in availability if (r['disease'],r['available_pct'],r['threshold_pct'])==(disease,available,available))
            avrows.append([zh,f'{available}% / {available}%',f"{r['accepted']}/{r['tasks']}",pct(r['rejection_rate']),pct(r['conditional_wape']),pct(r['same_target_complete_input_wape'])])
    appendix += [table(['病种','可用率/阈值','完成任务','拒绝率','条件WAPE','同键完整输入WAPE'],avrows,[1350,1500,1050,1200,1400,1500]),
        styled('完整27组配置均保留：可用率低于阈值时全部拒绝；阈值不高于可用率时，还需检查最新值与适配器输入。表D4每行245条记录来自49个相关原点与五个重复掩码，不是245个独立流行病任务。两列误差使用相同成功目标、相同适配器集合；跨行成功目标不同，不能据条件WAPE直接选择最佳阈值。拒绝任务不填零误差。该试验模拟输入缺失，不声称复现真实报告延迟、修订机制或未知病原。原始成本字段曾漏记被最终拒绝任务的适配器计算，已保留原记录并另存成本修订v2；全部准入判断、掩码与预测误差复核一致。'),
        heading('D.4 认证 持久化和多进程执行'),
        styled('新增文件持久化执行原型，将数据、任务、图谱、模型、参数、原点、输出范围及复核要求共同绑定到契约摘要。认证测试使用进程内生成的临时不透明令牌，角色由服务端测试注册表确定，HTTP请求体不能自行指定角色。状态以原子替换、文件同步及Linux进程锁保存；该设计只适用于同机协作进程，不声称替代跨机数据库事务或生产身份服务。'),
        caption('图D1 持久化授权契约的状态迁移'),figure,
        styled('运行状态在回调前落盘。进程中断后保留RUNNING并拒绝自动重试，避免将不确定的副作用重复执行；恢复需要另行审核，尚未实现自动恢复。已完成任务的同键请求返回保存结果，但仍需通过有效期和绑定校验。撤销与过期是终止状态，非法请求不进入数值回调。'),
        caption('表D5 授权与持久化测试的实际覆盖'),
        table(['测试组','测试项数','实际结果'],[
            ['生命周期、重启、并发、中断与返回副本','14','全部通过；8个独立进程仅执行1次同键回调'],
            ['八项绑定内容篡改','8','全部拒绝，数值回调未执行'],
            ['认证HTTP、角色越权与已有认证依赖','6','全部通过；缺失或错误凭据拒绝，合法请求链完成']], [3300,1200,3500]),
        styled('28项测试在vm201执行，其中六项为隔离ASGI请求测试，不启动公开监听。已有服务器Bearer认证依赖使用随机测试值单独核验，未读取在线令牌。真实数值运行经过新原型，36个持久化账本共1,872条契约全部完成，产生7,488条状态事件，并保留逐账本哈希。测试账号仅代表夹具身份，不是受试者或真人批准；新原型未接入现有生产API、单点登录或多租户边界，不能宣布生产权限验收完成。'),
        heading('D.5 批注处理范围与剩余条件'),
        styled('本轮补充了老师关于真实大模型对照、图谱联合消融、70%—90%可用率、资源归一墙钟比较、状态迁移图及持久化并发边界的可执行证据。图谱结构检查不代替双人质量抽检；既有代码归属声明和差分不代替负责人核验；历史资料回测不代替未来12周影子运行。未观察到显著增益的结果保留，不将实现通过写成预测方法优越。'),
        styled('服务器项目规则禁止在当前容器内使用Docker，因此本轮未执行Docker启动，也未为完成验收而改动宿主环境。后续应在获准的本地Docker环境验证Compose，再将隔离权限原型接入生产流程并做验收。用于实际预警的误差容差仍需依据任务用途预先确定，不能由本次回测成绩倒推。资源负责人批准、文献与资源许可的最终逐项复核、真人审核及真实用户评价分别登记，不以自动实验替代。')]
    appendix[0].find(f'{{{W}}}pPr').append(w('pageBreakBefore'))
    teacher_row=teacher.getparent().getparent();outer=teacher_row.getparent();position=list(outer).index(teacher_row)
    appendix_rows=blocks_to_rows(appendix,tcpr)
    # A heading at the end of a narrative row must travel with its first table.
    d2=next(p for row in appendix_rows for p in row.xpath('.//w:p',namespaces=NS) if text(p)=='D.2 选择预算与实际墙钟预算')
    d2.getparent().remove(d2)
    first_budget=next(row for row in appendix_rows if any(text(p).startswith('表D2 ') for p in row.xpath('.//w:p',namespaces=NS)))
    target_cell=first_budget.find(f'{{{W}}}tc');target_cell.insert(1,d2)
    for i,r in enumerate(appendix_rows):outer.insert(position+i,r)
    assert refs==[text(p) for p in doc.xpath('//w:p',namespaces=NS) if re.match(r'^\[\d+\]',text(p))]
    assert equations==[E.tostring(e) for e in doc.xpath('//m:oMath',namespaces=NS)]
    assert teacher_text==[text(p) for p in teacher.getparent().xpath('.//w:p',namespaces=NS)]
    parts['word/document.xml']=serialize(doc)
    with ZipFile(DEST,'w',ZIP_DEFLATED) as z:
        for name,payload in parts.items():z.writestr(name,payload)
    (OUT/'文档修改核查.json').write_text(json.dumps({'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'output_sha256':hashlib.sha256(DEST.read_bytes()).hexdigest(),
        'references_preserved':len(refs),'native_equations_preserved':len(equations),'teacher_text_preserved':True,'edits':edits,'new_main_table':mainrows,
        'original_protocols_and_unfavorable_results_preserved':True},ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'document_created':True,'references':len(refs),'equations':len(equations)},ensure_ascii=True))

if __name__=='__main__':main()
