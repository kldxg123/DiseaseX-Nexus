from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENTATION, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"F:\作业\aibola\论文材料包_20260813")
OUT = ROOT / "01_论文说明文档" / "疫情自动获取系统_实验架构与预测结果论文材料.docx"
ARCH = ROOT / "02_架构图"
RESULT = ROOT / "03_实验结果"
SHOT = ROOT / "04_系统截图"

NAVY = "162033"
BLUE = "2563EB"
PURPLE = "7C3AED"
TEAL = "0F766E"
RED = "DC2626"
LIGHT_BLUE = "EAF2FF"
LIGHT_GREY = "F4F7FB"
MID_GREY = "D9E2EF"
TEXT = "172235"
MUTED = "526076"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = MID_GREY, size: int = 6) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    r = p.add_run(text)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor.from_string(MUTED)


def add_image(doc: Document, path: Path, width_cm: float, caption: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    if caption:
        add_caption(doc, caption)


def add_callout(doc: Document, title: str, body: str, fill: str = LIGHT_BLUE, accent: str = BLUE) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(0.22)
    table.columns[1].width = Cm(15.4)
    left, right = table.rows[0].cells
    set_cell_shading(left, accent)
    set_cell_shading(right, fill)
    set_cell_border(left, accent, 0)
    set_cell_border(right, fill, 0)
    p = right.paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title + "\n")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(TEXT)
    r.font.size = Pt(11)
    r2 = p.add_run(body)
    r2.font.color.rgb = RGBColor.from_string(TEXT)
    r2.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_metric_tiles(doc: Document, items: list[tuple[str, str]], cols: int = 3) -> None:
    rows = (len(items) + cols - 1) // cols
    table = doc.add_table(rows=rows, cols=cols)
    table.autofit = False
    for i, (label, value) in enumerate(items):
        cell = table.cell(i // cols, i % cols)
        set_cell_shading(cell, LIGHT_GREY)
        set_cell_border(cell, "D8E1EE", 6)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(5)
        r1 = p.add_run(label + "\n")
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = RGBColor.from_string(MUTED)
        r2 = p.add_run(value)
        r2.font.size = Pt(17)
        r2.font.bold = True
        r2.font.color.rgb = RGBColor.from_string(BLUE)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = 1
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_border(cell, "FFFFFF", 4)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9)
        if widths:
            cell.width = Cm(widths[i])
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_shading(cells[i], "FFFFFF" if row_idx % 2 == 0 else LIGHT_GREY)
            set_cell_border(cells[i], MID_GREY, 4)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor.from_string(TEXT)
            if widths:
                cells[i].width = Cm(widths[i])


def add_bullets(doc: Document, bullets: list[str]) -> None:
    for text in bullets:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(text)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)

    for name, size, color in [("Title", 31, NAVY), ("Heading 1", 19, NAVY), ("Heading 2", 14, BLUE), ("Heading 3", 11.5, TEAL)]:
        st = styles[name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)

    for sec in doc.sections:
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hr = hp.add_run("EPIDEMIC INTELLIGENCE PLATFORM  /  THESIS MATERIALS")
        hr.font.name = "Aptos"
        hr.font.size = Pt(8)
        hr.font.bold = True
        hr.font.color.rgb = RGBColor.from_string(MUTED)
        add_page_number(sec.footer.paragraphs[0])


def new_landscape_section(doc: Document):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENTATION.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    return section


def new_portrait_section(doc: Document):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENTATION.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    return section


def build() -> None:
    doc = Document()
    configure_document(doc)

    # Cover
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("疫情自动获取系统")
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(32)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("实验设计、系统架构、知识图谱与预测结果论文材料")
    r2.font.size = Pt(20)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor.from_string(BLUE)
    doc.add_paragraph()
    add_callout(doc, "材料定位", "用于硕士论文的方法设计、系统实现与实验结果章节。所有数值均来自现有审计工件；四轨道结果图明确区分历史回放、重复评估与新鲜盲测。", "EEF4FF", BLUE)
    doc.add_paragraph()
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in ["版本日期：2026-08-14", "数据与代码状态：新增第三条 70% 以上历史回放链路；未触发正式发布", "用途：论文撰写 / 中期汇报 / 答辩展示"]:
        rr = p3.add_run(line + "\n")
        rr.font.size = Pt(11)
        rr.font.color.rgb = RGBColor.from_string(MUTED)
    doc.add_page_break()

    doc.add_heading("摘要", level=1)
    doc.add_paragraph(
        "本材料汇总了“疫情自动获取系统”的整体工程架构、预测模型知识图谱、四条预测评估链路，以及一组基于 2024 年新论文方法的可复现实验对比。系统围绕可信官方数据、来源溯源、人工复核、知识图谱推荐、确定性回测和受控 Qwen 调参构建。知识图谱已形成 1,574 个唯一全文版本、4,502 个证据片段、3,791 条断言和 26 条数据画像；其中 17 条推荐具备完整执行合同。对比实验采用新加坡政府 2012—2019 年周度登革热数据，在统一切分和 rolling one-step 协议下，复现论文 AR(2) 方法并与项目模型、验证集选出的组合模型和 Last-value 基线比较。项目组合模型测试 WAPE 为 12.80%、预测吻合度为 87.20%，接近论文 AR(2) 的 11.34% 和 88.66%，但未超过。轨道 A、B 与新增 E 均达到 70% 以上吻合度，但分别属于历史回放、重复评估和论文对齐历史回放；C 是一次性消费的新鲜盲测，却未达到 70% 门禁。"
    )
    doc.add_heading("本包包含什么", level=2)
    add_bullets(doc, [
        "可直接引用的 DOCX/PDF 说明文档；",
        "项目整体架构、知识图谱逻辑架构、对比实验架构的 PNG/SVG/PDF/drawio；",
        "2024 论文方法对比结果图、逐行预测 CSV 与参数搜索结果；",
        "四条预测链路的论文结果图，以及现有真实系统页面截图；",
        "知识图谱资产构建、浏览器验收与可复现脚本。",
    ])
    add_callout(doc, "重要解释", "本文中的“预测吻合度”定义为 max(0, 1−WAPE)×100%，用于把误差转成易读百分比；它不是分类准确率，也不等于未来部署精度。", "FFF7E8", "F59E0B")

    doc.add_heading("1 项目总体设计", level=1)
    doc.add_paragraph(
        "系统采用 local-first 与分层治理设计。官方数据、论文证据和人工候选首先进入可追溯的原始层；清洗规范化后分别进入 Raw、Trusted、Preview 和 Forecast 逻辑边界。知识图谱把论文身份、证据定位、疾病—模型关系、参数来源和数据画像连接起来，再生成具备风险门禁的 ExecutionContract。预测执行由确定性代码完成，Qwen 只可在白名单模型与参数空间内提出方案。"
    )
    new_landscape_section(doc)
    doc.add_heading("图 1  项目整体架构", level=1)
    add_image(doc, ARCH / "图1_项目整体架构.png", 25.5, "图 1  疫情自动获取系统整体架构（本轮用户提供的权威替换版）")

    new_portrait_section(doc)
    doc.add_heading("2 预测模型知识图谱", level=1)
    doc.add_paragraph(
        "本项目的知识图谱可概括为三层：第一层是数据存储与证据层，保存论文、PDF 版本、页级 locator、双 SHA 和数据画像；第二层是关系与语义层，表达疾病、模型、参数、适用范围、证据分和执行合同；第三层是问答、推荐与审计层，接收国家、病种和指标，返回候选模型、中文理由与论文溯源。人工批准某条疾病—模型 Claim 不等于允许正式预测，二者由不同门禁控制。"
    )
    add_table(doc, ["资产", "规模", "论文含义"], [
        ["唯一全文版本", "1,574", "支撑论文方法与参数证据的原始载体"],
        ["EvidenceSpan", "4,502", "页码、offset、excerpt 与 SHA 可反向验证"],
        ["Assertion", "3,791", "疾病—模型—适用条件等语义关系"],
        ["DataProfile", "26", "国家、病种、指标、频率和质量画像"],
        ["ExecutionContract", "17", "可由执行器消费的模型、参数、特征和风险门禁"],
    ], [4.0, 2.0, 9.2])

    new_landscape_section(doc)
    doc.add_heading("图 2  知识图谱逻辑架构", level=1)
    add_image(doc, ARCH / "图2_知识图谱逻辑架构.png", 25.0, "图 2  知识图谱的证据存储—关系推理—问答推荐三层结构")

    new_portrait_section(doc)
    doc.add_heading("知识图谱页面与论文溯源", level=2)
    add_image(doc, SHOT / "知识图谱_数据画像与推荐.png", 15.3, "图 3  知识图谱页面：数据画像、模型推荐与参数来源（既有真实页面截图）")
    add_image(doc, SHOT / "知识图谱_论文证据链.png", 15.3, "图 4  知识图谱页面：论文证据链与 locator（既有真实页面截图）")
    doc.add_paragraph(
        "页面的推荐不是静态模型清单。系统先读取数据画像，再根据证据分、参数可信度、频率、样本长度、稀疏性和目标语义筛选候选；用户可以从推荐项反向查看 DOI、原文片段和定位信息。"
    )

    doc.add_heading("3 对比实验设计", level=1)
    doc.add_heading("3.1 研究问题", level=2)
    doc.add_paragraph(
        "研究问题是：在与 2024 年论文一致的单变量周度感染人数场景中，本项目的滞后计数提升模型和验证集组合策略，能否在相同数据、相同切分和相同一步滚动协议下达到接近论文 AR(2) 方法的预测误差。"
    )
    doc.add_heading("3.2 数据与切分", level=2)
    add_table(doc, ["项目", "设置"], [
        ["数据来源", "新加坡政府官方周度登革热数据目录"],
        ["时间范围", "2012-W01 至 2019-W48"],
        ["规范化点数", "412；另保留 1 个 lag warm-up 点"],
        ["训练 / 验证 / 测试", "311 / 50 / 50 个目标周"],
        ["预测协议", "Rolling one-step；预测后才揭示当前周实际值"],
        ["比较方法", "论文 AR(2)、项目模型、项目组合模型、Last-value"],
    ], [4.0, 11.2])
    doc.add_heading("3.3 公平性与防泄漏门禁", level=2)
    add_bullets(doc, [
        "论文 AR(2) 使用公开系数：intercept=8.841、lag1=1.024、lag2=−0.062；",
        "项目模型参数只在 Validation 上选择；",
        "组合权重在 0.05 步长单纯形网格上只依据 Validation WAPE 确定；",
        "Test 标签在最终一次评估前保持未打开；",
        "所有方法按相同 target week 配对；独立重放 4/4 预测哈希一致。",
    ])

    new_landscape_section(doc)
    doc.add_heading("图 5  对比实验架构", level=1)
    add_image(doc, ARCH / "图3_2024论文对比实验架构.png", 25.5, "图 5  2024 年新论文方法对比实验架构")

    new_portrait_section(doc)
    doc.add_heading("4 对比实验结果", level=1)
    add_image(doc, RESULT / "图4_2024论文对比实验结果.png", 16.0, "图 6  测试集 WAPE、预测吻合度和 RMSE 对比")
    add_table(doc, ["方法", "Test WAPE ↓", "吻合度 ↑", "RMSE ↓", "结论"], [
        ["论文 AR(2)", "11.34%", "88.66%", "43.53", "最佳"],
        ["项目原始模型", "13.82%", "86.18%", "51.03", "未超过论文"],
        ["项目组合模型", "12.80%", "87.20%", "47.22", "接近论文"],
        ["Last-value 基线", "11.95%", "88.05%", "44.72", "强基线"],
    ], [4.0, 2.5, 2.5, 2.5, 4.0])
    add_callout(doc, "结果解读", "组合模型相对项目原始模型显著降低测试 WAPE，但仍比论文 AR(2) 高约 12.82%（相对差异）。这说明该数据具有很强的短期自相关，简单 AR(2) 与 Last-value 都是必须保留的强基线。论文写作中应表述为“达到接近性能层级”，而不是“超过论文”。", "FFF7E8", "F59E0B")
    doc.add_heading("4.1 复现实验的限制", level=2)
    add_bullets(doc, [
        "官方目录重建得到的序列在时间范围与点数上与论文对齐，但不是作者发布的字节级原始数据工件；",
        "为对齐 412 个点，重建过程中排除了 2014-W53；该决定和原始来源均被保留在 manifest 中；",
        "论文方法以公开方程和系数复现，不能替代作者完整训练环境；",
        "WAPE 对高值周权重更大，应与 RMSE、MAE、区间覆盖率和图形诊断联合解释。",
    ])

    doc.add_heading("5 四条预测评估链路", level=1)
    doc.add_paragraph(
        "四条链路均完成了知识图谱推荐或论文对齐模型选择、确定性回测与审计对账，但评估身份不同。A、B 与 E 的吻合度超过 70%，却不是新的 Fresh 盲测；C 是新鲜盲测，却没有达到 70% 认证阈值。该区别是系统治理设计的重要组成部分。"
    )
    add_table(doc, ["轨道", "国家 / 病种", "身份", "WAPE", "吻合度", "区间覆盖", "认证结论"], [
        ["A", "AUT / Legionella", "历史回放", "22.7%", "77.3%", "100%", "不构成新鲜盲测"],
        ["B", "FIN / Hantavirus", "重复评估", "28.0%", "72.0%", "97.9%", "重分类为非 fresh"],
        ["C", "CHL / Influenza", "新鲜盲测", "71.7%", "28.3%", "89.6%", "门禁失败，不认证"],
        ["E", "SGP / Dengue", "论文对齐历史回放", "12.80%", "87.20%", "未作为概率链评估", "达到展示门槛，非 Fresh"],
    ], [1.2, 3.0, 2.2, 1.7, 1.7, 1.8, 3.8])

    for idx, (img, caption) in enumerate([
        ("预测链路A_奥地利军团菌_结果图.png", "图 7  链路 A：奥地利军团菌历史回放结果"),
        ("预测链路B_芬兰汉坦病毒_结果图.png", "图 8  链路 B：芬兰汉坦病毒重复评估结果"),
        ("预测链路C_智利流感_结果图.png", "图 9  链路 C：智利流感新鲜盲测结果"),
        ("预测链路E_新加坡登革热_结果图.png", "图 10  链路 E：新加坡登革热论文对齐历史回放结果"),
    ]):
        if idx:
            doc.add_page_break()
        add_image(doc, RESULT / img, 16.0, caption)
        if idx == 0:
            doc.add_paragraph("链路 A 用于证明完整机制可运行；其测试窗口已用于历史回放，结果不可当作独立前瞻验证。")
        elif idx == 1:
            doc.add_paragraph("链路 B 的同样本结果优于基线，但旧窗口曾被多个 run 打开，治理迁移已将其标记为 RETROSPECTIVE_REPLAY_NOT_FRESH。")
        elif idx == 2:
            doc.add_paragraph("链路 C 的区间覆盖率通过，但 WAPE≤0.30 和吻合度≥70% 两项门禁失败。系统保留结果并拒绝认证。")
        else:
            doc.add_paragraph("链路 E 在同一官方数据、同一切分和同一逐周协议下达到 87.20% 吻合度；组合权重只由 Validation 选择。它是可复算的论文对齐历史回放，不构成新的独立盲测认证。")

    doc.add_heading("6 已上线的双数据集演示与反馈闭环", level=1)
    doc.add_paragraph(
        "除 A/B/C 治理轨道外，系统还保存了奥地利军团菌和秘鲁流感两条 Dashboard 演示链路。两条链路可展示知识图谱推荐、模型执行、基线覆盖与反馈记录，但其预测结果仍属于 Demo/Shadow，不应表述为正式公共卫生预测。"
    )
    add_image(doc, SHOT / "dual_loop_austria.png", 16.0, "图 11  奥地利军团菌预测闭环页面（既有真实页面截图）")
    add_image(doc, SHOT / "dual_loop_peru.png", 16.0, "图 12  秘鲁流感预测闭环页面（既有真实页面截图）")
    add_image(doc, SHOT / "dual_loop_feedback.png", 16.0, "图 13  模型反馈与学习周期页面（既有真实页面截图）")

    doc.add_heading("7 可信数据可视化", level=1)
    add_image(doc, SHOT / "可信数据三联图_巴西登革热.png", 16.0, "图 14  可信数据三联图：巴西登革热（既有真实页面截图）")
    doc.add_paragraph(
        "三联图展示可信数据的时间趋势、地理/指标摘要和数据表，并保留来源、时间范围与快照身份。该页面证明知识图谱与预测模块并非脱离数据治理层单独运行。"
    )

    doc.add_heading("8 论文写作建议", level=1)
    add_table(doc, ["论文部分", "建议使用材料", "建议表述"], [
        ["系统设计", "图 1、图 2", "强调分层治理、证据溯源和确定性执行"],
        ["方法设计", "图 5、对比实验协议", "强调同数据、同切分、同样本与防泄漏"],
        ["实验结果", "图 6、表 4", "如实说明项目接近但未超过 AR(2)"],
        ["工程验证", "图 7–12", "区分历史回放、重复评估和新鲜盲测"],
        ["可信性与局限", "第 4.1 节", "说明数据重建假设与吻合度定义"],
    ], [3.0, 4.0, 8.2])
    add_callout(doc, "禁止夸大", "不得将 77.3%/72.0%/87.20% 直接称为“新鲜盲测准确率”，也不得将 28.3% 的 Track C 写成成功认证。预测吻合度是 1−WAPE 的展示指标。", "FFF1F2", RED)

    doc.add_heading("参考文献与数据来源", level=1)
    refs = [
        "[1] Machine Learning Models for Dengue Forecasting in Singapore. arXiv:2407.00332, 2024. https://arxiv.org/abs/2407.00332",
        "[2] Singapore Government Data Catalogue. Weekly number of dengue and dengue haemorrhagic fever cases. https://data.gov.sg/datasets/d_ca168b2cb763640d72c4600a68f9909e/view",
        "[3] 本项目 2024 论文方法对比运行证据：comparison_2024_method.json、paired_predictions_2024_official_overlap.csv。",
        "[4] 本项目 Qwen held-out governance 运行证据与 browser acceptance JSON。",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.add_run(ref)

    doc.add_heading("附录：材料目录与复现入口", level=1)
    add_bullets(doc, [
        "02_架构图：每张图同时提供 drawio、SVG、PDF 和 PNG；",
        "03_实验结果：对比 JSON、参数网格、逐行预测 CSV、四轨道结果图；",
        "04_系统截图：知识图谱、双数据集预测闭环、可信数据三联图；",
        "05_知识图谱证据：资产构建报告与浏览器验收 JSON；",
        "06_可复现材料：数据 manifest、数据准备脚本和方法对比脚本；",
        "MANIFEST.sha256：对包内全部文件进行完整性校验。",
    ])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
