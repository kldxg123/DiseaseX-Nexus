from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"F:\作业\aibola\reports\top_venue_paper_benchmark_summary_20260820")
DOCX = OUT / "正式论文对比实验结果总结_RS0S_2024.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "17324D"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
LIGHT_GREEN = "EAF5EE"
LIGHT_GOLD = "FFF4D6"
RED = "9B1C1C"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int], indent: int = 120) -> None:
    total = sum(widths)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None, italic: bool | None = None) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_para(p, before: float = 0, after: float = 6, line: float = 1.10, keep: bool = False) -> None:
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep


def add_text(doc: Document, text: str, *, bold_lead: str | None = None, color: str | None = None, after: float = 6) -> None:
    p = doc.add_paragraph()
    set_para(p, after=after)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, bold=True, color=color or INK)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2, color=color or None)
    else:
        r = p.add_run(text)
        set_font(r, color=color)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)


def add_callout(doc: Document, title: str, text: str, fill: str = LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    set_para(p, after=2)
    r = p.add_run(title)
    set_font(r, bold=True, color=INK)
    p2 = cell.add_paragraph()
    set_para(p2, after=0)
    r2 = p2.add_run(text)
    set_font(r2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_kv_table(doc: Document, rows: list[tuple[str, str]], widths: list[int] | None = None) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for idx, (label, value) in enumerate(rows):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        set_cell_shading(cells[0], LIGHT_GRAY)
        for run in cells[0].paragraphs[0].runs:
            set_font(run, bold=True, color=INK)
        for run in cells[1].paragraphs[0].runs:
            set_font(run)
        for cell in cells:
            for p in cell.paragraphs:
                set_para(p, after=0, line=1.05)
    set_table_geometry(table, widths or [2700, 6660])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_results_table(doc: Document) -> None:
    headers = ["方法", "类别", "WAPE", "预测吻合度", "MAE", "RMSE", "95%区间覆盖率"]
    rows = [
        ("paper_arima", "论文方法", "22.31%", "77.69%", "68.61", "86.32", "80.41%"),
        ("last_value", "基线", "22.52%", "77.48%", "69.27", "87.76", "67.01%"),
        ("project_kg_ensemble", "项目组合", "22.89%", "77.11%", "70.41", "88.46", "90.72%"),
        ("project_qwen_tuned", "项目调优", "23.22%", "76.78%", "71.42", "90.08", "94.33%"),
        ("project_kg_primary", "项目主模型", "24.86%", "75.14%", "76.48", "97.54", "96.91%"),
        ("rolling_mean_4", "基线", "30.97%", "69.03%", "95.26", "115.63", "62.89%"),
        ("paper_slr", "论文方法", "35.55%", "64.45%", "109.35", "139.29", "55.20%"),
        ("paper_subepidemic_top", "论文方法近似", "44.10%", "55.90%", "135.63", "160.17", "47.90%"),
        ("seasonal_naive", "基线", "81.39%", "18.61%", "250.34", "293.60", "81.96%"),
        ("paper_gam", "论文方法", "896.99%", "0.00%", "2758.95", "8976.14", "17.00%"),
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        set_cell_shading(cell, BLUE)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_para(p, after=0, line=1.0)
            for run in p.runs:
                set_font(run, size=8.5, bold=True, color=WHITE)
    set_repeat_table_header(table.rows[0])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            if ridx == 0:
                set_cell_shading(cells[i], LIGHT_GREEN)
            elif ridx == 2:
                set_cell_shading(cells[i], LIGHT_BLUE)
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i < 2 else WD_ALIGN_PARAGRAPH.CENTER
                set_para(p, after=0, line=1.0)
                for run in p.runs:
                    set_font(run, size=8.2, bold=(ridx in (0, 2) and i in (0, 2, 3)))
    set_table_geometry(table, [1950, 1250, 850, 1050, 850, 850, 1550])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    run2 = paragraph.add_run(" 页")
    set_font(run2, size=9, color=MUTED)


def build() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_r = header_p.add_run("疫情自动获取系统｜正式论文对比实验")
    set_font(header_r, size=9, color=MUTED)
    add_page_number(section.footer.paragraphs[0])

    p = doc.add_paragraph()
    set_para(p, before=10, after=4)
    r = p.add_run("正式论文对比实验结果总结")
    set_font(r, size=24, bold=True, color=INK)
    p2 = doc.add_paragraph()
    set_para(p2, after=14)
    r2 = p2.add_run("Bleichrodt 等（2024）RSOS 方法迁移实验与项目模型对照")
    set_font(r2, size=13.5, color=MUTED)

    add_kv_table(doc, [
        ("文档日期", "2026-08-20"),
        ("实验状态", "TOP_VENUE_PAPER_BENCHMARK_PARTIAL_WITH_EXPLICIT_GAPS"),
        ("复现身份", "METHOD_TRANSFER_REPLICATION（方法迁移，不是论文原数据精确复现）"),
        ("评估数据", "新加坡官方周度登革热病例，2012-W01 至 2019-W48"),
    ])

    add_callout(
        doc,
        "核心结论",
        "新论文不是会议论文，而是 Royal Society Open Science 正式同行评审期刊论文。项目组合模型测试集预测吻合度为 77.11%，已超过 70% 演示阈值，但论文 ARIMA 仍以 77.69% 位列第一；因此不能表述为项目击败所有论文方法。",
        LIGHT_BLUE,
    )

    add_heading(doc, "1. 论文身份与发表等级", 1)
    add_kv_table(doc, [
        ("论文题目", "Evaluating the forecasting performance of ensemble sub-epidemic frameworks and other time series models for the 2022–2023 mpox epidemic"),
        ("作者", "Amanda Bleichrodt, Ruiyan Luo, Alexander Kirpich, Gerardo Chowell"),
        ("发表载体", "Royal Society Open Science, 2024, 11(7):240248"),
        ("DOI", "10.1098/rsos.240248"),
        ("论文类型", "正式同行评审期刊论文，不是会议论文，也不是预印本"),
        ("数据库收录", "SCIE"),
        ("JCR 2024", "影响因子 3.0；Multidisciplinary Sciences 43/140；Q2"),
        ("JCR 2023", "影响因子 2.9；Multidisciplinary Sciences 31/134；Q1"),
        ("CCF 等级", "不适用。CCF A/B/C 不能直接套用于该综合性期刊"),
        ("中科院分区", "本次证据未使用，不能与 JCR 分区混写"),
    ])
    add_text(doc, "等级判断：该论文满足“正式同行评审、SCIE 收录、JCR 一区或二区”的候选门槛。最稳妥的表述是“SCIE 期刊论文，JCR 2024 Q2；上一年度 JCR 2023 Q1”，不应简称为 CCF A/B 论文。", bold_lead="等级判断：")

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_heading(doc, "2. 实验设计与可比性", 1)
    add_text(doc, "论文原任务：使用 2022—2023 年 mpox 周度确诊病例，对 ARIMA、GAM、简单线性回归和子流行病模型等方法进行 1—4 周预测比较。", bold_lead="论文原任务：")
    add_text(doc, "项目实验任务：将论文方法迁移到新加坡 2012—2019 年周度登革热官方病例序列，与项目 KG 推荐模型、Qwen 调优模型、组合模型及三个基线进行统一协议比较。", bold_lead="项目实验任务：")
    add_kv_table(doc, [
        ("输入变量", "单变量周度病例数，不引入气象、搜索、移动性、基因组或临床外生变量"),
        ("数据规模", "412 个连续规范化周；无缺失，不补零"),
        ("训练段", "312 个历史点，其中 311 个训练标签和 1 个滞后预热点"),
        ("验证段", "50 个 rolling origins，共 200 个 H1—H4 配对样本"),
        ("测试段", "50 个 rolling origins；H1/H2/H3/H4 分别 50/49/48/47，共 194 个配对样本"),
        ("选参规则", "Qwen 参数和组合权重仅由验证集选择，测试集 exactly-once"),
        ("公平性", "所有方法使用相同的 origin、horizon、target week 和 actual 配对集合"),
        ("治理身份", "历史方法迁移比较，不是 Fresh held-out 独立盲测"),
    ])
    add_callout(doc, "指标解释", "WAPE 越低越好。文档中的“预测吻合度”定义为 max(0, 1-WAPE)×100%，用于直观表达相对误差，并不是分类任务中的 Accuracy。", LIGHT_GOLD)

    add_heading(doc, "3. 测试集结果", 1)
    add_results_table(doc)
    psrc = doc.add_paragraph()
    set_para(psrc, before=4, after=4)
    rr = psrc.add_run("注：表中数值来自冻结的 194 条同配对测试记录；WAPE 与误差越低越好，覆盖率需结合区间宽度判断。")
    set_font(rr, size=9, italic=True, color=MUTED)

    add_heading(doc, "4. 结果解读", 1)
    add_text(doc, "总体冠军：paper_arima 的 WAPE 为 22.31%，预测吻合度为 77.69%，是本轮全部方法中的最优点预测结果。", bold_lead="总体冠军：")
    add_text(doc, "项目冠军：project_kg_ensemble 的 WAPE 为 22.89%，预测吻合度为 77.11%。其 WAPE 相对 paper_arima 约落后 2.6%，相对 last_value 约落后 1.6%。", bold_lead="项目冠军：")
    add_text(doc, "Qwen 调优效果：项目主模型 WAPE 从 24.86% 降至 23.22%，相对改善约 6.6%；验证集选择的组合进一步降至 22.89%。这证明受控调优有效，但未改变 ARIMA 仍占优的事实。", bold_lead="Qwen 调优效果：")
    add_text(doc, "区间质量：项目组合模型覆盖率为 90.72%，高于 ARIMA 的 80.41% 和 last-value 的 67.01%。但项目组合平均区间宽度约 259.53，ARIMA 约 255.47，必须同时呈现覆盖率与宽度，不能只用覆盖率宣称全面领先。", bold_lead="区间质量：")

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_heading(doc, "5. 为什么项目模型没有超过 ARIMA", 1)
    add_text(doc, "数据阶段差异：测试窗口处于相对平稳阶段，短窗 ARIMA 和 last-value 更容易利用近期惯性；项目 boosting 使用更长历史，包含高发年份，可能在低发阶段产生轻微高估。", bold_lead="数据阶段差异：")
    add_text(doc, "方法目标差异：论文子流行病框架主要面向单次暴发过程，而新加坡登革热序列具有多年季节性；将 mpox 方法迁移到 dengue 本身存在疾病机制和时间结构差异。", bold_lead="方法目标差异：")
    add_text(doc, "误差与不确定性的取舍：项目模型并非单纯追求最低点预测误差，还保留了更高的区间覆盖率。该优势有实际审慎意义，但不能替代点预测排名。", bold_lead="误差与不确定性的取舍：")

    add_heading(doc, "6. 尚未闭合的方法缺口", 1)
    gaps = [
        ("Prophet", "未实现，当前隔离环境缺少可用 cmdstan 运行链"),
        ("完整子流行病集成", "论文的多子波（n≤5）Cthr 加权/不加权组合未实现，当前仅有 top 单波近似"),
        ("区间生成", "论文参数 bootstrap 尚未复现，当前使用残差经验分位数区间"),
        ("ARIMA 与 GAM", "auto.arima 和 mgcv 使用确定性网格或样条近似，不构成原工具链逐字节复现"),
        ("疾病迁移", "论文原对象为 mpox，本项目对比对象为新加坡 dengue，属于方法迁移而非数值复现"),
    ]
    add_kv_table(doc, gaps, [2600, 6760])
    add_callout(doc, "复现身份限制", "在上述缺口补齐前，实验只能标记为 METHOD_TRANSFER_REPLICATION，不能写成“完整复现论文”或“项目已超过论文全部方法”。", LIGHT_GOLD)

    add_heading(doc, "7. 当前工程状态", 1)
    add_kv_table(doc, [
        ("实验与证据包", "已生成，ZIP SHA-256 已核对"),
        ("页面补丁", "已在本地 page_patch 生成，但尚未部署到 vm201 正式项目"),
        ("API 文件对账", "报告记录为通过"),
        ("浏览器验收", "未通过：2 个 addEventListener 空节点异常，多个 RSOS DOM 区块为空"),
        ("数据库业务数据", "未修改"),
        ("Trusted / Promotion / Active pointer", "零变更"),
        ("当前准确状态", "TOP_VENUE_PAPER_BENCHMARK_PARTIAL_WITH_UI_ACCEPTANCE_FAILURE_AND_METHOD_GAPS"),
    ])

    add_heading(doc, "8. 可直接用于论文的结果表述", 1)
    add_callout(
        doc,
        "建议正文",
        "本研究选取发表于 Royal Society Open Science 的 Bleichrodt 等（2024）作为正式同行评审对比文献，并在同一新加坡周度登革热病例序列、相同 rolling-origin 切分和 H1—H4 配对样本上开展方法迁移实验。论文 ARIMA 方法取得 22.31% 的 WAPE 和 77.69% 的预测吻合度，为全部方法最优；本项目的知识图谱组合模型取得 22.89% 的 WAPE 和 77.11% 的预测吻合度，点预测误差略高于 ARIMA，但 95% 预测区间覆盖率由 80.41% 提升至 90.72%。Qwen 受控调优使项目主模型 WAPE 相对下降约 6.6%，说明知识图谱约束下的参数优化具有有效性。由于实验属于跨疾病方法迁移，且 Prophet、完整多子波集成和参数 bootstrap 尚未完全复现，本文将该结果界定为方法迁移比较，而非论文原结果的精确复现。",
        LIGHT_GREEN,
    )

    add_heading(doc, "9. 来源与证据", 1)
    add_text(doc, "论文：Bleichrodt A, Luo R, Kirpich A, Chowell G. Royal Society Open Science, 2024, 11(7):240248. DOI: https://doi.org/10.1098/rsos.240248")
    add_text(doc, "等级证据：Royal Society 官方期刊 metrics 页面记录的 JCR 年度指标；本报告仅采用 JCR 分区，不使用中科院分区替代。")
    add_text(doc, "实验依据：top_venue_paper_benchmark_20260819.zip，SHA-256 1f53fb1e6c63d5de5994486735666d1ba08da33211099b8224f7161589aa120f。")

    core = doc.core_properties
    core.title = "正式论文对比实验结果总结"
    core.subject = "疫情自动获取系统：RSOS 2024 方法迁移与项目模型比较"
    core.author = "epidemic-openclaw project"
    core.keywords = "infectious disease forecasting, knowledge graph, WAPE, ARIMA, Qwen"
    core.comments = "Generated from verified benchmark evidence package."

    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    build()
