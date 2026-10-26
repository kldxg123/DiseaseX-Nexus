from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import time
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import parse_qsl, quote, urlencode, urlsplit, urlunsplit
from urllib.request import Request, urlopen


PROJECT = Path(os.environ.get("PROJECT", ".")).resolve()
REPORT = PROJECT / "reports" / "p1_16r_ebola_source_acquisition"
HTTP_EVIDENCE = REPORT / "http_evidence"
SCHEMA_PROBE = REPORT / "schema_probe"
ARTIFACTS = PROJECT / "storage" / "p1_16r" / "ebola_source_acquisition" / "artifacts"

IRIS = "https://iris.who.int"
HDX = "https://data.humdata.org"
EXISTING_ARTIFACT_BY_URL: dict[str, dict[str, Any]] = {}
SENSITIVE_URL_QUERY_KEYS = {
    "AWSAccessKeyId",
    "Signature",
    "x-amz-security-token",
    "X-Amz-Credential",
    "X-Amz-Signature",
    "X-Amz-Security-Token",
    "Expires",
}


def redact_url(url: str | None) -> str | None:
    if not url:
        return url
    try:
        parts = urlsplit(url)
        query_pairs = parse_qsl(parts.query, keep_blank_values=True)
        kept = [(k, v) for k, v in query_pairs if k not in SENSITIVE_URL_QUERY_KEYS and k.lower() not in {x.lower() for x in SENSITIVE_URL_QUERY_KEYS}]
        removed = len(kept) != len(query_pairs)
        if removed:
            kept.append(("redacted_signed_query", "true"))
        return urlunsplit((parts.scheme, parts.netloc, parts.path, parts.query and urlencode(kept), parts.fragment))
    except Exception:
        return "[unparseable-url-redacted]"


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_json(path: Path, obj: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def append_jsonl(path: Path, obj: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as f:
        f.write(json.dumps(obj, ensure_ascii=False, sort_keys=True) + "\n")


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if fieldnames is None:
        keys: list[str] = []
        for row in rows:
            for key in row:
                if key not in keys:
                    keys.append(key)
        fieldnames = keys
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def safe_name(text: str, max_len: int = 90) -> str:
    text = re.sub(r"[^A-Za-z0-9._-]+", "-", text).strip("-")
    return (text or "artifact")[:max_len]


def http_get(url: str, accept: str | None = None, timeout: int = 45) -> dict[str, Any]:
    headers = {
        "User-Agent": "epidemic-openclaw-p1-16r-a-source-audit/1.0",
    }
    if accept:
        headers["Accept"] = accept
    req = Request(url, headers=headers)
    started = time.time()
    try:
        with urlopen(req, timeout=timeout) as resp:
            body = resp.read()
            return {
                "url": url,
                "final_url": redact_url(resp.geturl()),
                "status": int(resp.status),
                "headers": {k.lower(): v for k, v in resp.headers.items()},
                "body": body,
                "elapsed_ms": int((time.time() - started) * 1000),
                "error": None,
                "retrieved_at": now(),
            }
    except HTTPError as e:
        body = e.read() if hasattr(e, "read") else b""
        return {
            "url": url,
            "final_url": redact_url(getattr(e, "url", url)),
            "status": int(e.code),
            "headers": {k.lower(): v for k, v in e.headers.items()} if e.headers else {},
            "body": body,
            "elapsed_ms": int((time.time() - started) * 1000),
            "error": f"HTTPError:{e.code}",
            "retrieved_at": now(),
        }
    except (URLError, TimeoutError, OSError) as e:
        return {
            "url": url,
            "final_url": url,
            "status": None,
            "headers": {},
            "body": b"",
            "elapsed_ms": int((time.time() - started) * 1000),
            "error": f"{type(e).__name__}:{e}",
            "retrieved_at": now(),
        }


def save_http_evidence(key: str, resp: dict[str, Any], include_body_sha_only: bool = True) -> None:
    body = resp.pop("body", b"")
    evidence = dict(resp)
    evidence["body_sha256"] = sha256_bytes(body) if body else None
    evidence["body_bytes"] = len(body)
    if not include_body_sha_only and body:
        evidence["body_preview"] = body[:512].decode("utf-8", "replace")
    write_json(HTTP_EVIDENCE / f"{safe_name(key)}.json", evidence)
    resp["body"] = body


def persist_artifact(kind: str, label: str, resp: dict[str, Any], expected_ext: str) -> dict[str, Any]:
    body = resp["body"]
    digest = sha256_bytes(body)
    rel = Path(digest[:2]) / f"{safe_name(label)}-{digest[:12]}{expected_ext}"
    path = ARTIFACTS / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(body)
    content_type = resp.get("headers", {}).get("content-type", "")
    magic = body[:8]
    record = {
        "kind": kind,
        "label": label,
        "url": resp["url"],
        "final_url": resp["final_url"],
        "status": resp["status"],
        "content_type": content_type,
        "bytes": len(body),
        "sha256": digest,
        "storage_uri": str(path.relative_to(PROJECT)),
        "magic_hex": magic.hex(),
        "retrieved_at": resp["retrieved_at"],
        "elapsed_ms": resp["elapsed_ms"],
    }
    append_jsonl(REPORT / "artifact_manifest.jsonl", record)
    return record


def load_existing_artifact_cache() -> dict[str, dict[str, Any]]:
    cache: dict[str, dict[str, Any]] = {}
    path = REPORT / "artifact_manifest.jsonl"
    if not path.exists():
        return cache
    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        if not line.strip():
            continue
        try:
            rec = json.loads(line)
        except json.JSONDecodeError:
            continue
        url = rec.get("url")
        storage_uri = rec.get("storage_uri")
        if url and storage_uri and (PROJECT / storage_uri).exists():
            rec["url"] = redact_url(str(url))
            rec["final_url"] = redact_url(rec.get("final_url"))
            cache[str(url)] = rec
    return cache


def load_json_body(resp: dict[str, Any]) -> dict[str, Any]:
    try:
        return json.loads(resp["body"].decode("utf-8"))
    except Exception as e:  # noqa: BLE001
        return {"_json_error": f"{type(e).__name__}:{e}"}


def get_meta_values(item: dict[str, Any], key: str) -> list[str]:
    out: list[str] = []
    for ent in item.get("metadata", {}).get(key, []) or []:
        if isinstance(ent, dict) and ent.get("value"):
            out.append(str(ent["value"]))
    return out


def item_title(item: dict[str, Any]) -> str:
    vals = get_meta_values(item, "dc.title")
    return vals[0] if vals else str(item.get("name") or item.get("uuid") or item.get("id") or "untitled")


def item_date(item: dict[str, Any]) -> str | None:
    for key in ["dc.date.issued", "dc.date.available", "dc.date.accessioned"]:
        vals = get_meta_values(item, key)
        if vals:
            return vals[0]
    return None


def item_handle(item: dict[str, Any]) -> str | None:
    vals = get_meta_values(item, "dc.identifier.uri")
    for v in vals:
        if "handle" in v:
            return v
    return item.get("handle")


def extract_items(search_json: dict[str, Any]) -> list[dict[str, Any]]:
    objs = (
        search_json.get("_embedded", {})
        .get("searchResult", {})
        .get("_embedded", {})
        .get("objects", [])
    )
    items: list[dict[str, Any]] = []
    for obj in objs:
        candidate = obj.get("_embedded", {}).get("indexableObject") or obj.get("indexableObject") or obj
        if isinstance(candidate, dict):
            uuid = candidate.get("uuid") or candidate.get("id")
            if uuid:
                items.append(candidate)
    return items


def iris_search(query: str, page: int = 0, size: int = 100) -> tuple[str, dict[str, Any]]:
    params = urlencode({"query": query, "page": page, "size": size})
    url = f"{IRIS}/server/api/discover/search/objects?{params}"
    resp = http_get(url, "application/json", timeout=60)
    save_http_evidence(f"iris-search-{query}-{page}", resp, include_body_sha_only=False)
    return url, load_json_body(resp)


def iris_get_json(url: str, key: str) -> tuple[dict[str, Any], dict[str, Any]]:
    resp = http_get(url, "application/json", timeout=60)
    save_http_evidence(key, resp)
    return load_json_body(resp), resp


def iris_bundle_url(item_uuid: str) -> str:
    return f"{IRIS}/server/api/core/items/{item_uuid}/bundles"


def bundle_bitstreams_url(bundle_uuid: str) -> str:
    return f"{IRIS}/server/api/core/bundles/{bundle_uuid}/bitstreams"


def parse_embedded_list(obj: dict[str, Any], *names: str) -> list[dict[str, Any]]:
    cur: Any = obj.get("_embedded", {})
    for name in names[:-1]:
        cur = cur.get(name, {}).get("_embedded", {}) if isinstance(cur, dict) else {}
    last = names[-1] if names else ""
    val = cur.get(last, []) if isinstance(cur, dict) else []
    return val if isinstance(val, list) else []


def extract_text_from_pdf(path: Path) -> str:
    pdftotext = shutil.which("pdftotext")
    if pdftotext:
        proc = subprocess.run(
            [pdftotext, "-layout", str(path), "-"],
            cwd=str(PROJECT),
            text=True,
            capture_output=True,
            timeout=90,
        )
        if proc.returncode == 0 and proc.stdout.strip():
            return proc.stdout
    try:
        from pypdf import PdfReader  # type: ignore

        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages[:8])
    except Exception:
        return ""


def infer_report_number(title: str, text: str) -> int | None:
    hay = f"{title}\n{text[:5000]}"
    patterns = [
        r"(?:Situation\s+Report|External\s+Situation\s+Report|Sitrep|Report)\D{0,25}(?:No\.?|#)?\s*(\d{1,3})",
        r"\bNo\.?\s*(\d{1,3})\b",
    ]
    for pat in patterns:
        m = re.search(pat, hay, flags=re.I)
        if m:
            n = int(m.group(1))
            if 1 <= n <= 200:
                return n
    return None


DATE_RE = r"(\d{1,2}\s+[A-Za-z]+\s+\d{4}|\d{4}-\d{2}-\d{2})"


def infer_dates(text: str) -> dict[str, str | None]:
    out: dict[str, str | None] = {"date_of_issue": None, "data_as_of": None}
    m = re.search(r"(?:Date of issue|Issue date)\s*:?\s*" + DATE_RE, text[:6000], flags=re.I)
    if m:
        out["date_of_issue"] = m.group(1)
    m = re.search(r"(?:Data as reported by|Data as of|as of)\s*:?\s*" + DATE_RE, text[:6000], flags=re.I)
    if m:
        out["data_as_of"] = m.group(1)
    return out


def extract_candidate_rows_from_text(text: str) -> list[str]:
    rows: list[str] = []
    lines = text.splitlines()
    in_table = False
    for raw in lines:
        line = re.sub(r"\s+", " ", raw).strip()
        if not line:
            continue
        if re.search(r"\bTable\s+1\b", line, flags=re.I):
            in_table = True
        if in_table and re.search(r"\bTable\s+[2-9]\b", line, flags=re.I):
            in_table = False
        numeric_count = len(re.findall(r"\b\d{1,6}\b", line))
        letter_count = len(re.findall(r"[A-Za-z]", line))
        if (in_table or re.search(r"\b(health zone|province|territory|confirmed|probable|North Kivu|Ituri)\b", line, flags=re.I)) and numeric_count >= 3 and letter_count >= 4:
            rows.append(line)
    return rows


def rowset_hash(rows: list[str]) -> str:
    canonical = "\n".join(rows)
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def parse_pdf_twice(path: Path) -> dict[str, Any]:
    text1 = extract_text_from_pdf(path)
    rows1 = extract_candidate_rows_from_text(text1)
    text2 = extract_text_from_pdf(path)
    rows2 = extract_candidate_rows_from_text(text2)
    return {
        "text_chars": len(text1),
        "candidate_table_rows": len(rows1),
        "rowset_hash_1": rowset_hash(rows1),
        "rowset_hash_2": rowset_hash(rows2),
        "double_run_consistent": rowset_hash(rows1) == rowset_hash(rows2),
        "has_table1_markers": bool(re.search(r"Table\s+1|Health Zone|confirmed|probable", text1[:20000], flags=re.I)),
        "dates": infer_dates(text1),
    }


def acquire_iris_reports() -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    seen: set[str] = set()
    source_registry: list[dict[str, Any]] = []
    west_rows: list[dict[str, Any]] = []
    drc_rows: list[dict[str, Any]] = []
    duplicate_audit_rows: list[dict[str, Any]] = []

    searches = [
        ("west-africa", 'dc.title:"Ebola response roadmap situation report"', [0], 40),
        ("west-africa", 'dc.title:"Ebola situation report"', [0], 60),
        ("drc-2018-2020", '"Ebola Virus Disease Democratic Republic of Congo External Situation Report"', [0, 1], 120),
        ("drc-2018-2020", '"Ebola Virus Disease Democratic Republic of Congo"', [0, 1], 120),
    ]

    event_items: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for event, query, pages, max_items in searches:
        for page in pages:
            _url, js = iris_search(query, page=page, size=100)
            for item in extract_items(js):
                uuid = item.get("uuid") or item.get("id")
                if not uuid or uuid in seen:
                    continue
                title = item_title(item)
                if "ebola" not in title.lower():
                    continue
                seen.add(uuid)
                event_items[event].append(item)
                if len(event_items[event]) >= max_items:
                    break

    for event, items in event_items.items():
        limit = 20 if event == "west-africa" else 10
        for item in items[:limit]:
            uuid = item.get("uuid") or item.get("id")
            title = item_title(item)
            issued = item_date(item)
            handle = item_handle(item)
            item_dir = REPORT / "iris_items" / str(uuid)
            item_dir.mkdir(parents=True, exist_ok=True)
            write_json(item_dir / "item.json", item)

            bundles_json, bundles_resp = iris_get_json(iris_bundle_url(str(uuid)), f"iris-bundles-{uuid}")
            write_json(item_dir / "bundles.json", bundles_json)
            bundles = parse_embedded_list(bundles_json, "bundles")
            original_bundles = [b for b in bundles if str(b.get("name", "")).upper() == "ORIGINAL"]
            if not original_bundles:
                original_bundles = bundles[:1]

            downloaded_any = False
            for bundle in original_bundles:
                buuid = bundle.get("uuid") or bundle.get("id")
                if not buuid:
                    continue
                bits_json, _bits_resp = iris_get_json(bundle_bitstreams_url(str(buuid)), f"iris-bitstreams-{uuid}-{buuid}")
                write_json(item_dir / f"bitstreams-{buuid}.json", bits_json)
                bitstreams = parse_embedded_list(bits_json, "bitstreams")
                pdf_bitstreams: list[dict[str, Any]] = []
                for bit in bitstreams:
                    name = bit.get("name") or bit.get("uuid") or "bitstream"
                    bit_uuid = bit.get("uuid") or bit.get("id")
                    fmt = (bit.get("bundleName") or bit.get("format") or bit.get("type") or "")
                    if not bit_uuid:
                        continue
                    content_url = bit.get("_links", {}).get("content", {}).get("href") or f"{IRIS}/server/api/core/bitstreams/{bit_uuid}/content"
                    # Prefer PDFs. IRIS extracted text is still evidence but not counted as primary PDF.
                    is_pdf_candidate = str(name).lower().endswith(".pdf") or "pdf" in str(fmt).lower()
                    if not is_pdf_candidate:
                        continue
                    bit["_p1_16r_content_url"] = content_url
                    pdf_bitstreams.append(bit)
                if not pdf_bitstreams:
                    continue

                def pdf_score(bit: dict[str, Any]) -> tuple[int, str]:
                    name_l = str(bit.get("name") or "").lower()
                    bad_lang = any(tag in name_l for tag in ["_ara", "_fre", "_fra", "_por", "_spa", "_rus", "_chi", "_zho", "_fr", "_es", "_pt", "_ar"])
                    score = 0
                    if any(tag in name_l for tag in ["_eng", "-eng", " english", "_en", "-en"]):
                        score += 10
                    if bad_lang:
                        score -= 20
                    # If there is only one plain PDF and no language suffix, keep it as likely default/English.
                    if not re.search(r"_(ara|fre|fra|por|spa|rus|chi|zho|fr|es|pt|ar)\.pdf$", name_l):
                        score += 1
                    return (score, name_l)

                preferred_bits = [sorted(pdf_bitstreams, key=pdf_score, reverse=True)[0]]
                for bit in preferred_bits:
                    name = bit.get("name") or bit.get("uuid") or "bitstream"
                    bit_uuid = bit.get("uuid") or bit.get("id")
                    content_url = str(bit.get("_p1_16r_content_url"))
                    cached = EXISTING_ARTIFACT_BY_URL.get(content_url)
                    if cached and (PROJECT / cached["storage_uri"]).exists():
                        artifact = dict(cached)
                        append_jsonl(REPORT / "artifact_manifest.jsonl", artifact)
                        body = (PROJECT / artifact["storage_uri"]).read_bytes()
                        resp = {
                            "url": content_url,
                            "final_url": artifact.get("final_url"),
                            "status": artifact.get("status"),
                            "headers": {"content-type": artifact.get("content_type")},
                            "body": body,
                            "retrieved_at": artifact.get("retrieved_at"),
                            "elapsed_ms": 0,
                        }
                    else:
                        resp = http_get(content_url, "application/pdf", timeout=90)
                        save_http_evidence(f"iris-pdf-{event}-{uuid}-{bit_uuid}", resp)
                        body = resp["body"]
                        artifact = persist_artifact("iris_pdf", f"{event}-{uuid}-{name}", resp, ".pdf") if body else None
                    pdf_ok = resp["status"] == 200 and body.startswith(b"%PDF")
                    downloaded_any = bool(downloaded_any or pdf_ok)
                    parse_info: dict[str, Any] = {}
                    report_no = None
                    if artifact and pdf_ok:
                        parse_info = parse_pdf_twice(PROJECT / artifact["storage_uri"])
                        report_no = infer_report_number(title, extract_text_from_pdf(PROJECT / artifact["storage_uri"])[:8000])
                    record = {
                        "event": event,
                        "publisher": "WHO IRIS",
                        "source": title,
                        "item_uuid": uuid,
                        "bitstream_uuid": bit_uuid,
                        "handle": handle,
                        "publication_date": issued,
                        "report_number": report_no,
                        "url": content_url,
                        "final_url": resp.get("final_url"),
                        "http_status": resp.get("status"),
                        "mime": resp.get("headers", {}).get("content-type"),
                        "bytes": len(body),
                        "sha256": artifact["sha256"] if artifact else None,
                        "pdf_magic": pdf_ok,
                        "parser_status": "DETERMINISTIC_PARSE_EXPERIMENT" if pdf_ok else "NOT_PARSED",
                        "candidate_table_rows": parse_info.get("candidate_table_rows", 0),
                        "rowset_hash": parse_info.get("rowset_hash_1"),
                        "double_run_consistent": parse_info.get("double_run_consistent"),
                        "date_of_issue": parse_info.get("dates", {}).get("date_of_issue") if parse_info else None,
                        "data_as_of": parse_info.get("dates", {}).get("data_as_of") if parse_info else None,
                        "terminal_disposition": "ACQUIRED_OFFICIAL_PRIMARY" if pdf_ok else "BLOCKED_HTTP",
                        "reason": None if pdf_ok else f"HTTP_OR_PDF_MAGIC_FAILED_{resp.get('status')}",
                    }
                    append_jsonl(REPORT / "iris_item_manifest.jsonl", record)
                    append_jsonl(REPORT / "acquisition_manifest.jsonl", record)
                    source_registry.append(record)
                    duplicate_audit_rows.append({"sha256": record["sha256"], "source": title, "url": content_url})
                    if event == "west-africa":
                        west_rows.append(record)
                    else:
                        drc_rows.append(record)
            if not downloaded_any:
                record = {
                    "event": event,
                    "publisher": "WHO IRIS",
                    "source": title,
                    "item_uuid": uuid,
                    "handle": handle,
                    "publication_date": issued,
                    "terminal_disposition": "BLOCKED_HTTP",
                    "reason": "NO_ORIGINAL_PDF_BITSTREAM_DOWNLOADED",
                }
                append_jsonl(REPORT / "iris_item_manifest.jsonl", record)
                append_jsonl(REPORT / "acquisition_manifest.jsonl", record)
                source_registry.append(record)

    return source_registry, west_rows, drc_rows


def hdx_package_show(package_id: str) -> tuple[dict[str, Any], dict[str, Any]]:
    url = f"{HDX}/api/3/action/package_show?id={quote(package_id)}"
    resp = http_get(url, "application/json", timeout=60)
    save_http_evidence(f"hdx-package-{package_id}", resp, include_body_sha_only=False)
    js = load_json_body(resp)
    write_json(REPORT / "hdx_packages" / f"{package_id}.json", js)
    return js, resp


def csv_schema_probe(path: Path, skip_hxl: bool = True) -> dict[str, Any]:
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as f:
        reader = csv.reader(f)
        rows = list(reader)
    if not rows:
        return {"rows": 0, "header": [], "data_rows": 0}
    header = rows[0]
    body = rows[1:]
    hxl_skipped = False
    if skip_hxl and body and body[0] and str(body[0][0]).startswith("#"):
        body = body[1:]
        hxl_skipped = True
    dates: list[str] = []
    countries: set[str] = set()
    regions: set[str] = set()
    for row in body[:20000]:
        rowd = {header[i]: row[i] if i < len(row) else "" for i in range(len(header))}
        for key in ["Date", "date", "report_date", "data_as_of"]:
            if rowd.get(key):
                dates.append(rowd[key])
        for key in ["Country", "country", "country_name", "Country name"]:
            if rowd.get(key):
                countries.add(rowd[key])
        for key in ["Health Zone", "health_zone", "zone_sante", "Zone de sante", "Location"]:
            if rowd.get(key):
                regions.add(rowd[key])
    return {
        "rows_including_header": len(rows),
        "data_rows": len(body),
        "header": header,
        "hxl_tag_row_skipped": hxl_skipped,
        "sample_rows": body[:3],
        "date_min_sample": min(dates) if dates else None,
        "date_max_sample": max(dates) if dates else None,
        "country_count_sample": len(countries),
        "region_count_sample": len(regions),
    }


def xlsx_schema_probe(path: Path) -> dict[str, Any]:
    out: dict[str, Any] = {"zip_magic_ok": path.read_bytes()[:2] == b"PK"}
    try:
        import openpyxl  # type: ignore

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheets = []
        for ws in wb.worksheets:
            rows = 0
            header = None
            for row in ws.iter_rows(values_only=True):
                rows += 1
                if header is None and any(x is not None for x in row):
                    header = [str(x) if x is not None else "" for x in row]
                if rows > 5:
                    break
            sheets.append({"name": ws.title, "max_row": ws.max_row, "max_column": ws.max_column, "first_header": header})
        out["openpyxl_ok"] = True
        out["sheets"] = sheets
    except Exception as e:  # noqa: BLE001
        out["openpyxl_ok"] = False
        out["error"] = f"{type(e).__name__}:{e}"
    return out


def acquire_hdx() -> tuple[list[dict[str, Any]], dict[str, Any]]:
    records: list[dict[str, Any]] = []
    probes: dict[str, Any] = {}
    targets = [
        {
            "event": "west-africa",
            "package": "ebola-cases-2014",
            "label": "HDX ebola_data_db_format CSV",
            "url": "https://data.humdata.org/dataset/0d089fa0-3567-4b01-9c03-39d340ff34e3/resource/c59b5722-ca4b-41ca-a446-472d6d824d01/download/ebola_data_db_format.csv",
            "ext": ".csv",
            "kind": "csv",
            "expected_status": "ACQUIRED_FORMAL_SECONDARY_PENDING_ROW_VALIDATION",
        },
        {
            "event": "west-africa",
            "package": "rowca-ebola-cases",
            "label": "HDX ROWCA public XLSX",
            "url": "https://data.humdata.org/dataset/f326cf9d-8e37-437d-9b85-5ea3d36e3d39/resource/9b68ab69-e0b3-4ff5-b2d8-90bf446acd52/download/data-ebola-public.xlsx",
            "ext": ".xlsx",
            "kind": "xlsx",
            "expected_status": "ACQUIRED_FORMAL_SECONDARY_PENDING_ROW_VALIDATION",
        },
        {
            "event": "drc-2018-2020",
            "package": "ebola-cases-and-deaths-drc-north-kivu",
            "label": "HDX DRC national CSV",
            "url": "https://data.humdata.org/dataset/ebola-cases-and-deaths-drc-north-kivu/resource/4532a799-ebaf-4ff6-95a5-3607973f0de4/download/ebola-epidemic-national-figures.csv",
            "ext": ".csv",
            "kind": "csv",
            "expected_status": "ACQUIRED_FORMAL_SECONDARY_PENDING_ROW_VALIDATION",
        },
        {
            "event": "drc-2018-2020",
            "package": "ebola-cases-and-deaths-drc-north-kivu",
            "label": "HDX DRC health-zone CSV",
            "url": "https://data.humdata.org/dataset/ebola-cases-and-deaths-drc-north-kivu/resource/5f5a0ce9-4e44-440f-9d04-86784ecc7b5f/download/ebola-epidemic-health-zone-figures.csv",
            "ext": ".csv",
            "kind": "csv",
            "expected_status": "ACQUIRED_FORMAL_SECONDARY_PENDING_ROW_VALIDATION",
        },
    ]
    for package in sorted({t["package"] for t in targets} | {"evd-cases-by-district"}):
        js, resp = hdx_package_show(package)
        resources = js.get("result", {}).get("resources", []) if isinstance(js, dict) else []
        probes[f"package_{package}"] = {
            "status": resp.get("status"),
            "success": js.get("success") if isinstance(js, dict) else None,
            "resource_count": len(resources),
            "resource_summary": [
                {
                    "id": r.get("id"),
                    "name": r.get("name"),
                    "format": r.get("format"),
                    "url": r.get("url"),
                    "hash": r.get("hash"),
                    "size": r.get("size"),
                    "last_modified": r.get("last_modified"),
                }
                for r in resources[:20]
            ],
        }
    for target in targets:
        resp = http_get(target["url"], timeout=90)
        save_http_evidence(f"hdx-resource-{target['label']}", resp)
        body = resp["body"]
        status = target["expected_status"]
        reason = None
        artifact = None
        schema: dict[str, Any] = {}
        if resp.get("status") == 202:
            status = "BLOCKED_HTTP"
            reason = "HTTP_202_RETRYABLE_PENDING"
        elif resp.get("status") != 200:
            status = "BLOCKED_HTTP"
            reason = f"HTTP_{resp.get('status')}_{resp.get('error') or 'NO_RESPONSE_BYTES'}"
        elif target["kind"] == "csv" and body.lstrip()[:32].lower().startswith(b"<!doctype"):
            status = "BLOCKED_SCHEMA"
            reason = "CSV_RESOURCE_RETURNED_HTML"
        elif target["kind"] == "xlsx" and not body.startswith(b"PK"):
            status = "BLOCKED_SCHEMA"
            reason = "XLSX_RESOURCE_MISSING_PK_MAGIC"
        else:
            artifact = persist_artifact("hdx_" + target["kind"], target["label"], resp, target["ext"])
            path = PROJECT / artifact["storage_uri"]
            if target["kind"] == "csv":
                schema = csv_schema_probe(path)
            else:
                schema = xlsx_schema_probe(path)
            write_json(SCHEMA_PROBE / f"{safe_name(target['label'])}.json", schema)
        rec = {
            "event": target["event"],
            "publisher": "OCHA HDX",
            "package": target["package"],
            "source": target["label"],
            "url": target["url"],
            "final_url": resp.get("final_url"),
            "http_status": resp.get("status"),
            "mime": resp.get("headers", {}).get("content-type"),
            "bytes": len(body),
            "sha256": artifact["sha256"] if artifact else None,
            "schema_probe": schema,
            "terminal_disposition": status,
            "reason": reason or "SECONDARY_INDEX_REQUIRES_ROW_LEVEL_PRIMARY_LINEAGE_BEFORE_TRUSTED_PROMOTION",
        }
        records.append(rec)
        append_jsonl(REPORT / "acquisition_manifest.jsonl", rec)
    write_json(SCHEMA_PROBE / "hdx_package_resource_probe.json", probes)
    return records, probes


def read_db_counts() -> dict[str, Any]:
    try:
        import sqlalchemy as sa  # type: ignore
        from sqlalchemy.engine import make_url  # type: ignore
    except Exception as e:  # noqa: BLE001
        return {"available": False, "error": f"{type(e).__name__}:{e}"}
    url = os.environ.get("MIGRATION_DATABASE_URL") or os.environ.get("DATABASE_URL")
    if not url:
        return {"available": False, "error": "NO_DATABASE_URL_ENV"}
    try:
        parsed = make_url(url)
        admin_user = os.environ.get("POSTGRES_ADMIN_USER")
        admin_password = os.environ.get("POSTGRES_ADMIN_PASSWORD")
        if admin_user and admin_password:
            parsed = parsed.set(username=admin_user, password=admin_password)
        url = parsed
    except Exception:
        pass
    engine = sa.create_engine(url, future=True)
    queries = {
        "normalized_observations": "select count(*) from epidemic.normalized_observations",
        "trusted_lineage": "select count(*) from epidemic_p1_12.trusted_observation_lineage",
        "ebola_trusted_rows": """
            select count(*) from epidemic_p1_12.trusted_dashboard_series_v1
            where disease_slug='ebola'
        """,
        "flunet_trusted_rows": """
            select count(*) from epidemic_p1_12.trusted_dashboard_series_v1
            where source_key='trusted-influenza-who-flunet-v1'
        """,
    }
    out: dict[str, Any] = {"available": True, "transaction_read_only": None}
    try:
        with engine.connect() as conn:
            conn.execute(sa.text("BEGIN TRANSACTION ISOLATION LEVEL REPEATABLE READ READ ONLY"))
            conn.execute(sa.text("SET LOCAL statement_timeout = '30s'"))
            conn.execute(sa.text("SET LOCAL idle_in_transaction_session_timeout = '60s'"))
            out["transaction_read_only"] = conn.execute(sa.text("SHOW transaction_read_only")).scalar()
            for key, sql in queries.items():
                sp_name = f"sp_{key}"
                conn.execute(sa.text(f"SAVEPOINT {sp_name}"))
                try:
                    out[key] = int(conn.execute(sa.text(sql)).scalar() or 0)
                    conn.execute(sa.text(f"RELEASE SAVEPOINT {sp_name}"))
                except Exception as e:  # noqa: BLE001
                    out[key] = {"error": f"{type(e).__name__}:{e}"}
                    conn.execute(sa.text(f"ROLLBACK TO SAVEPOINT {sp_name}"))
                    conn.execute(sa.text(f"RELEASE SAVEPOINT {sp_name}"))
            conn.execute(sa.text("ROLLBACK"))
    finally:
        engine.dispose()
    return out


def git_state() -> dict[str, Any]:
    def run(args: list[str]) -> tuple[int, str, str]:
        p = subprocess.run(args, cwd=str(PROJECT), text=True, capture_output=True, timeout=30)
        return p.returncode, p.stdout.strip(), p.stderr.strip()

    return {
        "head": run(["git", "rev-parse", "HEAD"])[1],
        "dirty_count": len(run(["git", "status", "--short"])[1].splitlines()),
        "dirty_sample": run(["git", "status", "--short"])[1].splitlines()[:20],
    }


def build_reports(
    db_before: dict[str, Any],
    db_after: dict[str, Any],
    registry: list[dict[str, Any]],
    west_rows: list[dict[str, Any]],
    drc_rows: list[dict[str, Any]],
    hdx_records: list[dict[str, Any]],
) -> dict[str, Any]:
    west_dates = sorted({r.get("publication_date") or r.get("data_as_of") or r.get("date_of_issue") for r in west_rows if r.get("terminal_disposition") == "ACQUIRED_OFFICIAL_PRIMARY" and (r.get("publication_date") or r.get("data_as_of") or r.get("date_of_issue"))})
    west_countries = ["Guinea", "Liberia", "Sierra Leone"]
    west_matrix: list[dict[str, Any]] = []
    for country in west_countries:
        west_matrix.append(
            {
                "country": country,
                "downloaded_who_report_timepoints": len(west_dates),
                "secondary_hdx_artifacts": sum(1 for r in hdx_records if r["event"] == "west-africa" and r["terminal_disposition"].startswith("ACQUIRED")),
                "status": "ACQUIRED_OFFICIAL_PRIMARY" if len(west_dates) >= 20 else "ACQUIRED_DISCOVERY_INDEX_ONLY",
                "blocker": None if len(west_dates) >= 20 else "LESS_THAN_20_DISTINCT_WHO_REPORT_TIMEPOINTS",
            }
        )
    write_csv(REPORT / "west_africa_coverage_matrix.csv", west_matrix)

    drc_report_rows: list[dict[str, Any]] = []
    parseable_drc = 0
    for r in sorted(drc_rows, key=lambda x: (x.get("report_number") or 999, x.get("publication_date") or "")):
        parseable = bool(r.get("candidate_table_rows", 0) and r.get("double_run_consistent"))
        parseable_drc += 1 if parseable else 0
        drc_report_rows.append(
            {
                "report_number": r.get("report_number"),
                "title": r.get("source"),
                "publication_date": r.get("publication_date"),
                "date_of_issue": r.get("date_of_issue"),
                "data_as_of": r.get("data_as_of"),
                "pdf_sha256": r.get("sha256"),
                "candidate_table_rows": r.get("candidate_table_rows", 0),
                "double_run_consistent": r.get("double_run_consistent"),
                "terminal_disposition": r.get("terminal_disposition"),
            }
        )
    write_csv(REPORT / "drc_report_series_manifest.csv", drc_report_rows)

    drc_hz = next((r for r in hdx_records if "health-zone" in r.get("source", "").lower()), None)
    drc_hz_rows: list[dict[str, Any]] = []
    if drc_hz:
        schema = drc_hz.get("schema_probe", {})
        drc_hz_rows.append(
            {
                "source": drc_hz.get("source"),
                "status": drc_hz.get("terminal_disposition"),
                "rows": schema.get("data_rows"),
                "region_count_sample": schema.get("region_count_sample"),
                "hxl_tag_row_skipped": schema.get("hxl_tag_row_skipped"),
                "sha256": drc_hz.get("sha256"),
                "trusted_promotion_eligible_next_stage": "NO_PRIMARY_ROW_LINEAGE_NOT_YET_VALIDATED",
            }
        )
    write_csv(REPORT / "drc_health_zone_coverage.csv", drc_hz_rows)

    all_acquired = [r for r in registry + hdx_records if str(r.get("terminal_disposition", "")).startswith("ACQUIRED")]
    lineage_rows = []
    for r in all_acquired:
        lineage_rows.append(
            {
                "source": r.get("source"),
                "publisher": r.get("publisher"),
                "event": r.get("event"),
                "has_sha": bool(r.get("sha256")),
                "has_final_url": bool(r.get("final_url")),
                "has_size": bool(r.get("bytes")),
                "has_mime": bool(r.get("mime")),
                "terminal_disposition": r.get("terminal_disposition"),
                "trusted_promotion_eligible_next_stage": (
                    "YES_PARSE_CANDIDATE_ONLY_REQUIRES_ROW_LEVEL_CONTRACT"
                    if r.get("publisher") == "WHO IRIS" and r.get("candidate_table_rows", 0) > 0 and r.get("double_run_consistent")
                    else "NO_ROW_LEVEL_PRIMARY_LINEAGE_NOT_YET_COMPLETE"
                ),
            }
        )
    write_csv(REPORT / "lineage_validation.csv", lineage_rows)

    sha_counts = Counter(r.get("sha256") for r in all_acquired if r.get("sha256"))
    duplicate_conflict = {
        "duplicate_artifact_sha_count": sum(1 for _sha, c in sha_counts.items() if c > 1),
        "duplicate_artifact_shas": {sha: c for sha, c in sha_counts.items() if c > 1},
        "parser_double_run_inconsistent_count": sum(1 for r in registry if r.get("double_run_consistent") is False),
        "row_level_conflict_count": 0,
        "trusted_writes": 0,
    }
    write_json(REPORT / "duplicate_conflict_audit.json", duplicate_conflict)
    temporal = {
        "missing_converted_to_zero": False,
        "event_final_spread_to_weeks": False,
        "total_components_double_counted": False,
        "reporting_week_claimed_as_onset_week": False,
        "hdx_change_trusted_without_primary_lineage": False,
        "notes": [
            "This phase records acquisition feasibility only.",
            "ROWCA/HDX records remain secondary pending row-level primary lineage validation.",
        ],
    }
    write_json(REPORT / "temporal_semantics_audit.json", temporal)

    downloaded_files = [p for p in ARTIFACTS.rglob("*") if p.is_file()]
    total_bytes = sum(p.stat().st_size for p in downloaded_files)
    status = "EBOLA_SOURCE_ACQUISITION_FEASIBLE"
    blockers: list[str] = []
    if len(west_dates) < 20:
        blockers.append("WEST_AFRICA_WHO_REPORT_TIMEPOINTS_LT_20")
    if parseable_drc < 10:
        blockers.append("DRC_TABLE1_PARSEABLE_REPORTS_LT_10")
    if not drc_hz or drc_hz.get("terminal_disposition") != "ACQUIRED_FORMAL_SECONDARY_PENDING_ROW_VALIDATION":
        blockers.append("DRC_HEALTH_ZONE_CSV_NOT_ACQUIRED")
    if blockers:
        status = "EBOLA_SOURCE_ACQUISITION_PARTIAL_EXTERNAL_BLOCKERS"

    db_writes = {
        "ebola_trusted_delta": (
            db_after.get("ebola_trusted_rows") - db_before.get("ebola_trusted_rows")
            if isinstance(db_before.get("ebola_trusted_rows"), int) and isinstance(db_after.get("ebola_trusted_rows"), int)
            else None
        ),
        "flunet_trusted_delta": (
            db_after.get("flunet_trusted_rows") - db_before.get("flunet_trusted_rows")
            if isinstance(db_before.get("flunet_trusted_rows"), int) and isinstance(db_after.get("flunet_trusted_rows"), int)
            else None
        ),
        "normalized_observations_delta": (
            db_after.get("normalized_observations") - db_before.get("normalized_observations")
            if isinstance(db_before.get("normalized_observations"), int) and isinstance(db_after.get("normalized_observations"), int)
            else None
        ),
    }
    state = {
        "status": status,
        "blockers": blockers,
        "created_at_utc": now(),
        "downloaded_file_count": len(downloaded_files),
        "downloaded_bytes": total_bytes,
        "west_africa_distinct_report_timepoints": len(west_dates),
        "west_africa_countries_covered": west_countries,
        "drc_reports_downloaded": sum(1 for r in drc_rows if r.get("terminal_disposition") == "ACQUIRED_OFFICIAL_PRIMARY"),
        "drc_table1_parseable_reports": parseable_drc,
        "drc_health_zone_csv_rows": drc_hz_rows[0].get("rows") if drc_hz_rows else None,
        "terminal_disposition_counts": Counter(str(r.get("terminal_disposition")) for r in registry + hdx_records),
        "database_before": db_before,
        "database_after": db_after,
        "database_deltas": db_writes,
        "trusted_insert_update_delete": {
            "ebola": {"insert": 0, "update": 0, "delete": 0},
            "flunet": {"insert": 0, "update": 0, "delete": 0},
        },
    }
    write_json(REPORT / "EXECUTION_STATE.json", state)

    feasibility = f"""# P1-16R-A Ebola source acquisition feasibility

Status: {status}

This phase only acquired and audited sources. It did not write Ebola Trusted, FluNet Trusted, train models, or change API/Dashboard state.

## Summary

- Downloaded artifact files: {len(downloaded_files)}
- Downloaded bytes: {total_bytes}
- West Africa WHO report timepoints acquired: {len(west_dates)}
- DRC WHO IRIS reports acquired: {state['drc_reports_downloaded']}
- DRC reports with deterministic Table 1-style parseability evidence: {parseable_drc}
- DRC HDX health-zone CSV rows: {state['drc_health_zone_csv_rows']}
- Ebola Trusted delta: {db_writes['ebola_trusted_delta']}
- FluNet Trusted delta: {db_writes['flunet_trusted_delta']}

## Feasibility

Sources with next-stage promotion potential are WHO IRIS PDFs whose deterministic parse experiment produced stable rowset hashes, and HDX/ROWCA rows only when their row-level primary WHO/MoH lineage can be proven. HDX/ROWCA remain secondary index/candidate evidence in this phase.

## Blockers

{json.dumps(blockers, ensure_ascii=False, indent=2)}

## Governance notes

- Missing reports are not converted to zero.
- Event-final values are not allocated to weeks.
- Country totals and components are not double-counted.
- HDX change fields are not trusted without primary lineage validation.
"""
    (REPORT / "SOURCE_ACQUISITION_FEASIBILITY.md").write_text(feasibility, encoding="utf-8")
    return state


def make_zip() -> dict[str, str]:
    payload: list[dict[str, Any]] = []
    for p in sorted(REPORT.rglob("*")):
        if p.is_file():
            payload.append({"path": str(p.relative_to(REPORT)), "bytes": p.stat().st_size, "sha256": sha256_file(p)})
    write_json(REPORT / "payload_sha_manifest.json", {"created_at_utc": now(), "files": payload})
    zip_path = PROJECT / "reports" / "p1_16r_ebola_source_acquisition.zip"
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for p in sorted(REPORT.rglob("*")):
            if p.is_file():
                zf.write(p, p.relative_to(REPORT.parent))
    digest = sha256_file(zip_path)
    (PROJECT / "reports" / "p1_16r_ebola_source_acquisition.zip.sha256").write_text(
        f"{digest}  {zip_path.name}\n",
        encoding="utf-8",
    )
    return {"zip": str(zip_path), "sha256": digest}


def main() -> None:
    global EXISTING_ARTIFACT_BY_URL
    REPORT.mkdir(parents=True, exist_ok=True)
    HTTP_EVIDENCE.mkdir(parents=True, exist_ok=True)
    SCHEMA_PROBE.mkdir(parents=True, exist_ok=True)
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    EXISTING_ARTIFACT_BY_URL = load_existing_artifact_cache()
    for dirname in ["http_evidence", "schema_probe", "iris_items", "hdx_packages"]:
        d = REPORT / dirname
        if d.exists():
            shutil.rmtree(d)
        d.mkdir(parents=True, exist_ok=True)
    for name in [
        "acquisition_manifest.jsonl",
        "iris_item_manifest.jsonl",
        "artifact_manifest.jsonl",
    ]:
        p = REPORT / name
        if p.exists():
            p.unlink()

    execution_plan = """# P1-16R-A execution plan

1. Record read-only database/git baseline.
2. Use IRIS DSpace REST to enumerate items, bundles and ORIGINAL PDF bitstreams.
3. Download and hash WHO IRIS PDFs.
4. Fetch HDX CKAN package_show metadata and selected CSV/XLSX resources.
5. Run deterministic parser probes and schema probes only.
6. Produce source acquisition feasibility report without Trusted writes.
"""
    (REPORT / "EXECUTION_PLAN.md").write_text(execution_plan, encoding="utf-8")
    db_before = read_db_counts()
    baseline = {"git": git_state(), "database": db_before, "created_at_utc": now()}
    write_json(REPORT / "baseline_readonly.json", baseline)

    registry, west_rows, drc_rows = acquire_iris_reports()
    hdx_records, hdx_probe = acquire_hdx()
    # GHO legacy metadata probe; no repeated resource retries.
    gho_resp = http_get(f"{HDX}/api/3/action/package_show?id=evd-cases-by-district", "application/json", timeout=60)
    save_http_evidence("hdx-package-evd-cases-by-district-second-pass", gho_resp, include_body_sha_only=False)
    gho_js = load_json_body(gho_resp)
    gho_record = {
        "event": "legacy-gho-probe",
        "publisher": "WHO/HDX",
        "source": "evd-cases-by-district",
        "url": gho_resp["url"],
        "final_url": gho_resp["final_url"],
        "http_status": gho_resp["status"],
        "mime": gho_resp.get("headers", {}).get("content-type"),
        "terminal_disposition": "ACQUIRED_DISCOVERY_INDEX_ONLY" if gho_js.get("success") else "UPSTREAM_LEGACY_RESOURCE_BYTES_UNAVAILABLE",
        "reason": "METADATA_ONLY_PROBE_NO_TRUSTED_BYTES_CLAIM",
    }
    append_jsonl(REPORT / "acquisition_manifest.jsonl", gho_record)
    registry.append(gho_record)

    write_json(REPORT / "source_registry.json", {"sources": registry + hdx_records + [gho_record]})
    db_after = read_db_counts()
    state = build_reports(db_before, db_after, registry, west_rows, drc_rows, hdx_records)
    zip_info = make_zip()
    state["final_zip_sha256"] = zip_info["sha256"]
    write_json(REPORT / "EXECUTION_STATE.json", state)
    print(json.dumps({"status": state["status"], "zip_sha256": zip_info["sha256"], "downloaded_file_count": state["downloaded_file_count"], "downloaded_bytes": state["downloaded_bytes"]}, ensure_ascii=False, sort_keys=True))


if __name__ == "__main__":
    main()
