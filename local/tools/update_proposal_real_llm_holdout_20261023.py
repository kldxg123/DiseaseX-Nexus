"""Surgical school-form update from frozen server results; preserve prior evidence."""
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from io import BytesIO
from copy import deepcopy
import csv,json,re,hashlib,statistics
from lxml import etree as E
from docx import Document
from docx.shared import Inches
from PIL import Image,ImageDraw,ImageFont
from reframe_proposal_baseline_selection_20261023 import styled,table,blocks_to_rows,W,NS,w,serialize,text,replace

ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'reports/frozen_llm_validation_20261023'
OUT=ROOT/'output/proposal_real_llm_holdout_20261023'
SOURCE=ROOT/'output/proposal_baseline_selection_20261023/安浩然_开题报告_三病种滚动择优与基线保留修订稿_20261023.docx'
DEST=OUT/'安浩然_开题报告_三病种新增时段与真实LLM实验修订稿_20261023.docx'
DISEASES={'Dengue Fever':'登革热','Campylobacterenterosis':'弯曲菌肠炎','Mumps':'腮腺炎'}
MODELS={'last_value':'持续性','mean4':'4周均值','seasonal52':'52周季节性','mean8':'8周均值','mean13':'13周均值','drift4':'4周漂移','drift13':'13周漂移','log_ar2':'对数AR(2)','log_ar4':'对数AR(4)','deterministic':'确定性提议','real_llm':'大模型提议'}

def load(name):return json.loads((DATA/name).read_text(encoding='utf-8'))
def readcsv(name):
    with (DATA/name).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def pct(x):return f'{x*100:.2f}%'
def metric(data,disease,arm,h='all'):return next(r for r in data if (r['disease'],r['arm'],r['horizon'])==(disease,arm,h))
def ci(data,disease,comparison):return next(r for r in data if (r['disease'],r['comparison'],r['block_origins'])==(disease,comparison,8))
def ci_text(r):return f"[{r['ci95_pp'][0]:.3f}, {r['ci95_pp'][1]:.3f}]"

def plot(rows):
    image=Image.new('RGB',(1760,1180),'white');draw=ImageDraw.Draw(image)
    font=ImageFont.truetype('C:/Windows/Fonts/simsun.ttc',40);small=ImageFont.truetype('C:/Windows/Fonts/simsun.ttc',32)
    svg=E.Element('svg',xmlns='http://www.w3.org/2000/svg',width='1760',height='1180',viewBox='0 0 1760 1180')
    E.SubElement(svg,'rect',width='1760',height='1180',fill='white')
    def label(x,y,s,big=False,color='#222222'):
        draw.text((x,y),s,font=font if big else small,fill=color)
        t=E.SubElement(svg,'text',x=str(x),y=str(y+(40 if big else 32)),fill=color,**{'font-family':'SimSun,serif','font-size':str(40 if big else 32)});t.text=s
    def line(points,color,width):
        draw.line(points,fill=color,width=width)
        E.SubElement(svg,'polyline',points=' '.join(f'{x:.2f},{y:.2f}' for x,y in points),fill='none',stroke=color,**{'stroke-width':str(width)})
    label(125,12,'周报告病例数',True)
    line([(1050,35),(1110,35)],'#24649A',4);label(1120,14,'实际报告')
    line([(1370,35),(1430,35)],'#D56A28',4);label(1440,14,'一步预测')
    plotted=[]
    for panel,(disease,zh) in enumerate(DISEASES.items()):
        rs=sorted([r for r in rows if r['disease']==disease and r['arm']=='real_llm' and int(r['horizon'])==1],key=lambda r:int(r['origin']))
        assert len(rs)==52
        top=110+panel*350;bottom=top+245;left=130;right=1680;maximum=max(max(float(r['actual']),float(r['point'])) for r in rs)*1.12
        label(left,top-46,f'{chr(97+panel)}  {zh}',True)
        for f in [0,.5,1]:
            y=bottom-245*f;line([(left,y),(right,y)],'#DFE3E8',1);label(24,y-16,f'{maximum*f:.0f}')
        line([(left,top),(left,bottom),(right,bottom)],'#555555',2)
        for field,color in [('actual','#24649A'),('point','#D56A28')]:
            line([(left+i*(right-left)/51,bottom-float(r[field])*245/maximum) for i,r in enumerate(rs)],color,4)
        for i in [0,13,26,39,51]:label(left+i*(right-left)/51-75,bottom+12,rs[i]['target_week'])
        plotted.extend([{k:r[k] for k in ['disease','target_week','actual','point','selected']} for r in rs])
    path=OUT/'新增时段三病种一步预测.png';image.save(path)
    (OUT/'新增时段三病种一步预测.svg').write_bytes(E.tostring(svg,encoding='utf-8',xml_declaration=True))
    with (OUT/'图2数据.csv').open('w',encoding='utf-8',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(plotted[0]));writer.writeheader();writer.writerows(plotted)
    return path

def main():
    OUT.mkdir(exist_ok=True)
    summary=load('result_summary.json');assert summary['new_window_completed'] and summary['real_llm_calls']==12 and summary['all_execution_on_vm201']
    new=load('new_window_metrics.json');historical=load('historical_llm_metrics.json');newci=load('new_window_confidence_intervals.json');oldci=load('historical_llm_confidence_intervals.json')
    ps=[load(p.name) for p in sorted(DATA.glob('proposal_*.json'))];assert len(ps)==12
    valid=sum(p['schema_valid'] for p in ps);json_valid=sum(isinstance(json.loads(p['raw_response']),dict) for p in ps)
    seconds=sum(p['seconds'] for p in ps);tokens=sum(p['usage']['completion_tokens'] for p in ps)
    figure_path=plot(readcsv('new_window_predictions.csv'))
    package=Document(SOURCE);p=package.add_paragraph();p.alignment=1;p.add_run().add_picture(str(figure_path),width=Inches(5.6))
    buffer=BytesIO();package.save(buffer)
    with ZipFile(buffer) as z:parts={n:z.read(n) for n in z.namelist()}
    doc=E.fromstring(parts['word/document.xml']);body=doc.find(f'{{{W}}}body')
    figure=next(p for p in list(body) if p.tag==f'{{{W}}}p' and p.xpath('.//w:drawing',namespaces=NS));body.remove(figure)
    original_ps=doc.xpath('//w:p',namespaces=NS)
    refs=[text(p) for p in original_ps if re.match(r'^\[\d+\]',text(p))]
    equations=[E.tostring(e) for e in doc.xpath('//m:oMath',namespaces=NS)]
    teacher=next(p for p in original_ps if text(p).startswith('学术导师对该生开题报告的简要评语'))
    teacher_text=[text(p) for p in teacher.getparent().xpath('.//w:p',namespaces=NS)]
    edits=[]
    def edit(p,newtext):
        if text(p)!=newtext:edits.append({'before':text(p),'after':newtext});replace(p,newtext)
    # Keep main captions consecutive: new Tables 12 and 13 precede the existing plan table.
    for p in original_ps:
        if re.search(r'表\s*12(?!\d)',text(p)):edit(p,re.sub(r'表\s*12(?!\d)','表14',text(p)))
    for p in original_ps:
        s=text(p)
        if s.startswith('本课题研究面向多病种疫情监测的可信智能预测系统。'):
            edit(p,'本课题研究面向多病种疫情监测的可信智能预测系统。系统依据病种数据条件生成候选，将合格基线与推荐模型共同纳入历史滚动评价；简单基线更好时保留基线。原型保存来源、输入版本、模型参数、授权与复核记录。登革热、弯曲菌肠炎和腮腺炎在2020—2022年回顾评价中的WAPE分别为22.09%、28.96%和46.70%。新增2024年度历史留出与12次真实大语言模型候选提议实验，三病种均完成预测，但登革热仍劣于最优简单基线，未观察到稳定的大模型精度增益。现有证据用于检验模型选择与可追溯运行，不支持未知病原预测有效或统计等效。预期交付本地系统、代码差分和复核材料，图谱质量及完整协作流程另行验收。')
        elif s.startswith('已用于原型、调参或结果展示的病例及年份归入开发或历史诊断资料。'):
            edit(p,'已用于原型、调参或结果展示的病例及年份归入开发或历史诊断资料。正文2.7节保留三病种原滚动结果，2.8节另列流程登记后获取的2024年度历史留出与真实大模型对照；附录B保留原型退化结果，附录C说明新增协议、候选契约和完整参考结果。新增年度表不与原622键混算。原流程先于下载登记，因完整主机核查发现既有模型而修订了执行模型，修订发生在首次推理和评分之前；不称第三方封存盲测。未来连续12个发布周的影子评价仍需实际执行。')
        elif s.startswith('本人于2026年10月23日确认项目代码由本人完成，'):
            edit(p,'本人于2026年10月23日确认项目代码由本人完成，后续按模块、接口、修改前后哈希、差分和测试日志登记；本轮验证实现和文字修订使用AI辅助，第三方库、数据及引用算法保留来源。没有可用Git历史，追溯依据为交付快照与逐行差分。已增加三病种回顾评价、年度回测、预算原型消融、2024年度历史留出和服务器真实大模型受控实验。图谱双人审核、完整模型协作与生产权限链路、Docker验收及真实用户任务评价仍需补齐，导师复核与签字另行确认。')
        elif '新窗口未冻结' in s:
            edit(p,s.replace('新窗口未冻结','新增2024年历史留出见2.8节'))
        elif s.startswith('三病种已有同源周度回顾数据，各574周，'):
            edit(p,'原三病种回顾资料各574周，2.7节保留原协议评价。新增2023—2024年各52周，现有序列各678周，新增留出结果见2.8节；表8保留早期案例状态。各协议不混算目标键。历史退化与预算受限分析见附录B，不据较好的主表改变原测试身份。同一国家监测资料仍不足以证明跨地区或未知病原适用性。')
        elif s.startswith('现已完成检索留档、代码与数据身份核查、'):
            edit(p,'已完成检索留档、代码与数据身份核查、三病种回顾评价、配对差值区间、年度嵌套回测、预算原型消融、准入敏感性、222项新增组件测试和42项既有工作流回归。本轮又完成2024年度历史留出、12次服务器真实大模型推理及12项选模与输出校验测试。图谱双人抽检、12周影子评价、完整大模型联动、生产授权、Docker验收和目标用户研究仍未完成。代码归属按本人声明登记，导师意见与签字仍需导师确认。')
        elif s.startswith('本候选池中六维、仅画像、无图谱等版本预测相同，'):
            edit(p,s.replace('真实大模型消融未执行。','该预算原型未执行真实大模型消融；另行完成的受控实验见2.8节和附录C，不混为同一协议。'))
        elif s.startswith('应用层冒烟最终通过，12个API端点均正常，'):
            edit(p,s.replace('无可用本地大模型运行条件，未执行真实LLM消融，未启动线上服务。','此前该轮检查未执行真实LLM消融；本轮另在服务器复用既有模型完成受控实验，见附录C，仍未启动或改动线上服务。'))
        elif s.startswith('图谱双人断言审核、目标用户任务评价、负责人或导师确认需要真人参与；'):
            edit(p,'图谱双人断言审核、目标用户任务评价、负责人或导师确认需要真人参与；12周影子评价需要新发布数据及真实日历。本轮新增时段与受控大模型实验见2.8节，不代替图谱增量、完整协作流程、状态迁移图、生产权限、Docker启动或用途容差验证。主线调整不取消这些要求，历史留出也不等同于前瞻运行。代码归属按本人声明登记，导师原评语保留，不代表本次修订已经获得新的签字确认。')
    # New numeric results, with the unfavourable primary comparison kept in the main text.
    main_rows=[];old_rows=[];baseline_lookup={}
    for disease,zh in DISEASES.items():
        reference=min([metric(new,disease,m) for m in ['last_value','mean4','seasonal52']],key=lambda r:r['wape']);baseline_lookup[disease]=reference['arm']
        llm=metric(new,disease,'real_llm');det=metric(new,disease,'deterministic');r=ci(newci,disease,'real_llm minus '+reference['arm'])
        main_rows.append([zh,pct(det['wape']),pct(llm['wape']),MODELS[reference['arm']],pct(reference['wape']),f"{r['delta_wape_pp']:+.3f}",ci_text(r)])
        r=ci(oldci,disease,'real_llm minus deterministic')
        old_rows.append([zh,pct(metric(historical,disease,'deterministic')['wape']),pct(metric(historical,disease,'real_llm')['wape']),f"{r['delta_wape_pp']:+.3f}",ci_text(r)])
    blocks=[styled('2.8 新增历史留出与真实大模型消融','h2'),
      styled('在原有回顾实验之外，下载官方2023年与2024年年度汇总表[36,37]，分别作为追加历史及校准资料、全年留出评价资料。每年每病种52个周度观测，无缺失值插补；2024年H1—H4分别有52、51、50、49个目标键，共202个相关配对。原流程于2026年10月23日先于年度表下载登记；随后因主机核查发现服务器既有模型而修订了模型执行方案，保留首版协议和修订记录。审查范围内未发现这些2024年周度表此前用于开发的记录，但未做第三方封存，不能据此称为严格盲测或前瞻验证。'),
      styled('真实推理复用服务器既有Qwen3-4B-Instruct-2507[38]，以CPU、float32和4线程运行，不改动原模型、环境或在线服务。模型仅接收历史样本数、零值比例、变异系数、自相关与近期变化比，不接收病名、地区、年份、日期或未来标签。两组均保留持续性、4周均值和52周季节性基线，再由确定性规则或大模型提供两个附加候选。每年首次预测前确定候选池，随后逐周用最近13个已完整观察的1—4周历史原点评价择优。各组候选拟合预算相同，大模型额外推理耗时单列；其预训练数据是否包含相关公开资料无法完全排除。'),
      styled('表12 2024年度历史留出预测与简单基线的比较','caption'),
      table(['病种','确定性WAPE','大模型WAPE','参考基线','基线WAPE','差值/百分点','差值95%区间'],main_rows,[1150,1000,1000,1200,950,1000,1600]),
      styled('注：参考基线为同一2024年目标集上三种常驻简单基线中表现最好的方法，仅作事后参考；差值为大模型组减参考基线。区间采用5,000次原点块自助法，主块长8，4与13作敏感性分析。附加候选未必全部入池，完整九候选的事后参考见附录C，不以测试期最优模型替代实际选择器。'),
      styled('三病种在新时段均完成预测。大模型组的WAPE为19.64%、36.54%和42.69%，相对最优简单基线的差值分别为+3.944、+3.432和−3.489个百分点。登革热差值区间高于零，不能据此声称各病种均达到最优基线水平；弯曲菌肠炎差值区间跨零，腮腺炎区间低于零，但比较仍属于有限任务的探索性分析。保留基线只是保障其参与历史比较，不保证基线在未来测试期的排名能够被提前识别。'),
      styled('图2 2024年度三病种完整一步预测','caption'),figure,
      styled('蓝线为年度表中的周报告值，橙线为实际大模型组选择器的一步预测。每个病种展示全部52周，不截取有利窗口；H2—H4及概率指标见附录C。曲线对应包含契约校验与回退的完整实验组，不表示每次候选都由大模型有效提供。'),
      styled('表13 2020—2022年真实大模型提议与确定性提议的受控比较','caption'),
      table(['病种','确定性WAPE','大模型WAPE','差值/百分点','差值95%区间'],old_rows,[1450,1450,1450,1450,2100]),
      styled('表13每病种含三个年度、157个原点和610个目标键，校准与候选按年度冻结，不能与2.7节622键的原四候选方案混算。大模型组相对确定性组未显示稳定的点预测优势。新增2024年窗口中，登革热和弯曲菌肠炎两组预测相同；腮腺炎WAPE减少1.116个百分点，差值区间为[−2.083,0.571]，仍不足以断言稳定增益。'),
      styled(f'共执行12次真实模型调用，输出均可解析为JSON，其中{valid}次满足候选契约；其余8次推荐了已常驻的4周均值，不符合“两个附加候选”的约束，按预定规则回退。没有执行白名单外模型或模型生成代码。推理总耗时{seconds:.2f}秒，单次中位数{statistics.median(p["seconds"] for p in ps):.2f}秒；该耗时不含模型加载和数值回测，不能当作端到端延迟。结果支持候选校验与基线回退可运行，尚不支持大模型必然提高精度或完整图谱协作已经验证。'),
      styled('年度表属于回顾汇总，不保证等同于各历史发布时点的数值。2024年登革热逐周求和为13,606，年末周报累计为13,610；保留4例差异，未按累计数改写周值，也未确定差异原因。其余两个病种的逐周和与同一年末报告累计数一致。数据版本核对、断言双人审核、真实用户评价及未来12周影子运行仍需分别完成。')]
    heading=next(p for p in doc.xpath('//w:p',namespaces=NS) if text(p)=='3 论文提纲')
    cell=heading.getparent();row=cell.getparent();outer=row.getparent();tcpr=deepcopy(cell.find(f'{{{W}}}tcPr'))
    content=[x for x in cell if x.tag!=f'{{{W}}}tcPr'];index=content.index(heading);position=list(outer).index(row);outer.remove(row)
    for j,r in enumerate(blocks_to_rows(content[:index]+blocks+content[index:],tcpr)):outer.insert(position+j,r)
    # Add source entries after the preserved original 35 references.
    last_ref=next(p for p in doc.xpath('//w:p',namespaces=NS) if text(p).startswith('[35]'))
    reference_additions=[
      '[36] COMMUNICABLE DISEASES AGENCY. Weekly Infectious Diseases Bulletin 2023[DB/OL]. [2026-10-23]. https://www.cda.gov.sg/resources/weekly-infectious-diseases-bulletin-2023/.',
      '[37] COMMUNICABLE DISEASES AGENCY. Weekly Infectious Diseases Bulletin 2024[DB/OL]. [2026-10-23]. https://www.cda.gov.sg/resources/weekly-infectious-diseases-bulletin-2024/.',
      '[38] QWEN TEAM. Qwen3-4B-Instruct-2507[EB/OL]. [2026-10-23]. https://huggingface.co/Qwen/Qwen3-4B-Instruct-2507.']
    for s in reference_additions:
        p=deepcopy(last_ref);replace(p,s);last_ref.addnext(p);last_ref=p
    # Detailed evidence in the appendix, without deleting any previous appendix.
    ap=[styled('附录C 新增时段与真实大模型实验记录','h2'),
      styled('C.1 协议与来源身份','h2'),
      styled('首版协议登记时间为2026年10月23日08:25 UTC，预定2024年全年留出、2023年校准、三种常驻基线、六种附加候选和每任务一次提议。新年度表在其后获取。首版未产生模型推理或预测成绩，因主机核查找到既有Qwen3权重与兼容环境，改用该权重并在首次推理前保存第二版协议。模型变化不是基于测试损失作出的选择，但实际模型协议晚于年度表获取，因此不称全部标签不可见条件下的外部预注册。两版协议、代码、权重与分词器指纹均保留。'),
      styled('年度表的原始标签为Dengue、Campylobacter enteritis与Mumps，分别映射到原三病种名称；DHF单列，未与登革热列合并。原2023年下载文件虽以.xls结尾，实际为XLSX压缩容器，按文件内容识别后读取，未改写原工作簿。312条规范观测保留年度、周次、期间文本、原标签及原单元格位置。审查仅覆盖已登记的相关资料与文件，不能证明所有旧归档或人工记忆均未接触该时期。'),
      styled('两个年度文件来自官方明确提供的静态下载。文件托管站robots地址返回403，未据此推断一般爬取许可；按用户明确请求仅直接获取两份文件，实际文件返回成功，未使用登录、验证码或访问控制绕过。原表未作为开放许可资料重新发布，复核包提供解析脚本、来源链接、哈希与研究所用数值记录。该处理不代替负责人对后续再分发权限的确认。'),
      styled('C.2 实际提议与回退','h2')]
    proposal_rows=[]
    for p in ps:
        raw=json.loads(p['raw_response'])['candidates']
        proposal_rows.append([DISEASES[p['disease']],str(p['year']),', '.join(raw),'通过' if p['schema_valid'] else '常驻基线重复','否' if p['schema_valid'] else '是'])
    ap += [styled('表C1 十二次真实模型输出的候选契约检查','caption'),table(['病种','年度','实际输出候选','契约结果','回退'],proposal_rows,[1250,800,3100,1700,1050]),
      styled('12次输出均为可解析JSON，但只有4次给出两个有效且互异的附加候选；其余8次包含mean4。该名称并非未注册模型，但已在常驻基线池中，不能占用两个附加候选的名额。保留原始输出和回退，不在查看预测成绩后放宽规则或重新提问。这里检验的是受约束候选JSON，不将其称为原生API tool_calls。'),
      styled('C.3 分预测期与完整候选参考','h2'),styled('表C2 2024年两组预测的分期WAPE','caption')]
    horizon_rows=[]
    for disease,zh in DISEASES.items():
        for arm in ['deterministic','real_llm']:horizon_rows.append([zh,MODELS[arm]]+[pct(metric(new,disease,arm,h)['wape']) for h in range(1,5)])
    ap += [table(['病种','提议方式','H1','H2','H3','H4'],horizon_rows,[1400,1500,1250,1250,1250,1250]),
      styled('下列九个固定候选的成绩使用相同2024年目标键复算，只作完整事后参考；实际两组始终只有三种基线与两个附加候选。固定候选成绩不等同于当时能够选中的最优方法。WIS和区间宽度受病种病例尺度影响，不跨病种直接排序。')]
    for number,(disease,zh) in enumerate(DISEASES.items(),3):
        rows=[]
        for arm in [*MODELS]:
            r=metric(new,disease,arm);rows.append([MODELS[arm],pct(r['wape']),f"{r['wis']:.3f}",pct(r['coverage95']),f"{r['width95']:.2f}"])
        ap += [styled(f'表C{number} {zh}的完整九候选与实际选择器结果','caption'),table(['方法','WAPE','WIS','95%覆盖率','95%平均宽度/病例数'],rows,[2000,1200,1300,1400,2000])]
    ap += [styled('C.4 差值区间与适用边界','h2')]
    cirows=[]
    for disease,zh in DISEASES.items():
        best=min([metric(new,disease,m) for m in MODELS if m not in ['deterministic','real_llm']],key=lambda r:r['wape'])
        for reference in [baseline_lookup[disease],best['arm']]:
            r=ci(newci,disease,'real_llm minus '+reference)
            cirows.append([zh,MODELS[reference],f"{r['delta_wape_pp']:+.3f}",ci_text(r)])
    ap += [styled('表C6 大模型组相对简单基线和完整候选参考的差值','caption'),table(['病种','参考方法','WAPE差值/百分点','95%区间'],cirows,[1450,1900,1800,2750]),
      styled('区间使用相同原点与目标键配对，2024年仅52个预测原点，202个键不是202个独立样本。大模型组相对完整九候选中的事后最优，在登革热和弯曲菌肠炎上仍有误差差距；历史2020—2022年中，登革热完整参考对数AR(4)的WAPE为20.13%，也优于两组的22.55%。这说明候选范围与历史选择均可能限制结果，不能只比较较弱参考或把历史择优写成测试期最优保障。完整预测、原点选择轨迹及块长4/8/13的全部区间随材料提供，未作多重比较校正。'),
      styled('新窗口的大模型组95%区间覆盖率在登革热、弯曲菌肠炎和腮腺炎上分别为90.59%、93.56%和84.65%，不能认为名义95%区间已经普遍校准。区间沿用2023年残差，未根据2024年表现重新校准。WAPE达到某数值或接近简单基线，不自动表示满足临床或预警用途；用途容差仍须预先确定。'),
      styled('C.5 运行与复核材料','h2'),
      styled(f'推理与数值回测全部在vm201执行，复用既有兼容环境，不修改全局依赖，不启动新监听服务，不占用GPU。共12次调用，生成{tokens}个输出token，推理合计{seconds:.2f}秒；模型加载另计。两组每个原点均执行5×13=65次逻辑历史拟合预算，实际实现共用缓存，分组缓存耗时不是独立端到端延迟，不能据此宣布整体效率提升。12项单元测试在原环境通过，本地复核同一测试集不重复计为24项独立检查。'),
      styled('本轮完成了新增年度历史留出和真实模型候选提议对照，不是完整知识图谱、大模型、生产授权与可视化链路的端到端用户验证。图谱增量、双人断言审核、12周影子评价、真实用户任务研究和导师贡献确认仍分别保留。原始提议、回退记录、实际选模轨迹和逐项预测与协议共同归档。')]
    ap[0].find(f'{{{W}}}pPr').extend([w('pageBreakBefore'),w('keepNext')])
    teacher_row=teacher.getparent().getparent();outer=teacher_row.getparent();position=list(outer).index(teacher_row)
    # Keep the C.3 heading with its first table, and avoid a mostly empty comparison page.
    begin=next(i for i,x in enumerate(ap) if text(x)=='C.3 分预测期与完整候选参考')
    last_caption=next(i for i,x in enumerate(ap) if text(x).startswith('表C3 '))
    stop=last_caption+2
    combined=w('tr');combined.append(w('trPr'));combined[0].append(w('cantSplit'))
    combined_cell=w('tc');combined_cell.append(deepcopy(tcpr));combined_cell.extend(ap[begin:stop]);combined_cell.append(styled(''));combined.append(combined_cell)
    c4=next(i for i,x in enumerate(ap) if text(x)=='C.4 差值区间与适用边界')
    c6=next(i for i,x in enumerate(ap) if text(x).startswith('表C6 '))+2
    grouped=w('tr');grouped.append(w('trPr'));grouped[0].append(w('cantSplit'))
    grouped_cell=w('tc');grouped_cell.append(deepcopy(tcpr));grouped_cell.extend(ap[c4:c6]);grouped_cell.append(styled(''));grouped.append(grouped_cell)
    appendix_rows=blocks_to_rows(ap[:begin],tcpr)+[combined]+blocks_to_rows(ap[stop:c4],tcpr)+[grouped]+blocks_to_rows(ap[c6:],tcpr)
    for j,r in enumerate(appendix_rows):outer.insert(position+j,r)
    assert refs==[text(p) for p in doc.xpath('//w:p',namespaces=NS) if re.match(r'^\[\d+\]',text(p))][:35]
    assert equations==[E.tostring(e) for e in doc.xpath('//m:oMath',namespaces=NS)]
    assert teacher_text==[text(p) for p in teacher.getparent().xpath('.//w:p',namespaces=NS)]
    parts['word/document.xml']=serialize(doc)
    with ZipFile(DEST,'w',ZIP_DEFLATED) as z:
        for n,b in parts.items():z.writestr(n,b)
    (OUT/'修改核查.json').write_text(json.dumps({'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'output':DEST.name,'original_references_preserved':35,'new_references':3,'native_equations_preserved':len(equations),'teacher_review_text_unchanged':True,'new_window_main_rows':main_rows,'historical_llm_main_rows':old_rows,'real_model_calls':len(ps),'valid_json':json_valid,'valid_optional_candidate_contracts':valid,'inference_seconds':seconds,'edits':edits},ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'output':DEST.name,'new_main_tables':2,'new_figures':1,'calls':len(ps)},ensure_ascii=True))

if __name__=='__main__':main()
