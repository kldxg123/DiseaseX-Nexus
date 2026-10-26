"""Reframe the retained proposal around rolling selection with baseline retention."""
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from copy import deepcopy
from io import BytesIO
import csv,hashlib,json,re
from lxml import etree as E
from docx import Document
from docx.shared import Inches
from PIL import Image,ImageDraw,ImageFont
from build_proposal_revision_20261022 import create_table,paragraph,w,W,NS,serialize
from polish_proposal_20261023 import replace,text

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/proposal_baseline_selection_20261023'
SOURCE=ROOT/'output/proposal_submission_20261023/安浩然_开题报告_实验补充与中文润色定稿_20261023.docx'
DEST=OUT/'安浩然_开题报告_三病种滚动择优与基线保留修订稿_20261023.docx'
DATA=ROOT/'reports/three_disease_chronological_20261023'
EXTRA=ROOT/'reports/teacher_requested_experiments_20261023'
LABELS={'Dengue Fever':'登革热','Campylobacterenterosis':'弯曲菌肠炎','Mumps':'腮腺炎'}
MODEL={'last_value':'持续性','mean4':'4周均值','seasonal52':'52周季节性','log_ar2':'对数AR(2)'}

def load(path):return json.loads(path.read_text(encoding='utf-8'))
def csvread(path):
    with path.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def styled(s,kind='p'):
    p=paragraph(s,kind);pr=p.find(f'{{{W}}}pPr')
    pr.append(w('spacing',line=240,lineRule='auto'))
    if kind=='p':pr.append(w('ind',firstLine=420))
    for run in p.findall(f'{{{W}}}r'):
        rp=run.find(f'{{{W}}}rPr')
        if rp is None:rp=w('rPr');run.insert(0,rp)
        rp.extend([w('rFonts',ascii='Times New Roman',hAnsi='Times New Roman',eastAsia='黑体' if kind in ('h1','h2','h3') else '宋体'),w('sz',val=24 if kind=='h2' else 21),w('color',val='000000')])
    return p

def table(headers,rows,widths):
    t=create_table(headers,rows,widths)
    for p in t.xpath('.//w:p',namespaces=NS):
        pr=p.find(f'{{{W}}}pPr');pr.append(w('spacing',line=240,lineRule='auto'))
        for r in p.findall(f'{{{W}}}r'):
            rp=r.find(f'{{{W}}}rPr')
            if rp is None:rp=w('rPr');r.insert(0,rp)
            rp.extend([w('rFonts',ascii='Times New Roman',hAnsi='Times New Roman',eastAsia='宋体'),w('sz',val=21),w('color',val='000000')])
    return t

def blocks_to_rows(blocks,tcpr):
    groups=[];buf=[]
    for x in blocks:
        visual=x.tag==f'{{{W}}}tbl' or bool(x.xpath('.//w:drawing',namespaces=NS))
        if visual:
            caption=buf.pop() if buf and text(buf[-1]).startswith(('表','图')) else None
            if buf:groups.append((False,buf));buf=[]
            groups.append((True,([caption] if caption is not None else [])+[x,styled('')]))
        else:buf.append(x)
    if buf:groups.append((False,buf))
    rows=[]
    for keep,elements in groups:
        row=w('tr');pr=w('trPr')
        if keep:pr.append(w('cantSplit'))
        row.append(pr);cell=w('tc');cell.append(deepcopy(tcpr));cell.extend(elements);row.append(cell);rows.append(row)
    return rows

def plot(pred):
    """Native SVG paths and matching high-resolution Pillow plot from frozen H1 data."""
    image=Image.new('RGB',(1760,1180),'white');draw=ImageDraw.Draw(image)
    font=ImageFont.truetype('C:/Windows/Fonts/simsun.ttc',42);small=ImageFont.truetype('C:/Windows/Fonts/simsun.ttc',36)
    svg=E.Element('svg',xmlns='http://www.w3.org/2000/svg',width='1760',height='1180',viewBox='0 0 1760 1180')
    E.SubElement(svg,'rect',width='1760',height='1180',fill='white')
    def label(x,y,s,size=27,color='#222222'):
        draw.text((x,y),s,font=font if size==32 else small,fill=color)
        actual_size=42 if size==32 else 36
        t=E.SubElement(svg,'text',x=str(x),y=str(y+actual_size),fill=color,**{'font-size':str(actual_size),'font-family':'SimSun,serif'});t.text=s
    def line(points,color,width=3):
        draw.line(points,fill=color,width=width)
        E.SubElement(svg,'polyline',points=' '.join(f'{x:.2f},{y:.2f}' for x,y in points),fill='none',stroke=color,**{'stroke-width':str(width)})
    label(130,10,'周报告病例数',32)
    line([(1020,35),(1080,35)],'#24649A',4);label(1090,13,'实际报告',27)
    line([(1330,35),(1390,35)],'#D56A28',4);label(1400,13,'一步预测',27)
    plot_rows=[]
    for panel,(disease,zh) in enumerate(LABELS.items()):
        rows=sorted([dict(r,target_week=r.get('target_week',r.get('week'))) for r in pred if r['disease']==disease and int(r['horizon'])==1],key=lambda r:int(r['target']))
        assert len(rows)==157
        top=110+panel*350;left=135;right=1680;bottom=top+245
        label(left,top-47,f'{chr(97+panel)}  {zh}',32)
        maximum=max(max(float(r['actual']),float(r['point'])) for r in rows)*1.12
        for fraction in (0,.5,1):
            yy=bottom-245*fraction
            line([(left,yy),(right,yy)],'#DFE3E8',1);label(24,yy-16,f'{maximum*fraction:.0f}')
        line([(left,top),(left,bottom),(right,bottom)],'#444444',2)
        for field,color in (('actual','#24649A'),('point','#D56A28')):
            points=[(left+i*(right-left)/156,bottom-float(r[field])/maximum*245) for i,r in enumerate(rows)]
            line(points,color,4)
        for i in (0,53,105,156):
            xx=left+i*(right-left)/156;label(xx-100,bottom+10,rows[i]['target_week'])
        plot_rows.extend([{k:r[k] for k in ('disease','target_week','actual','point','chosen_component')} for r in rows])
    image.save(OUT/'三病种一步滚动预测.png')
    (OUT/'三病种一步滚动预测.svg').write_bytes(E.tostring(svg,encoding='utf-8',xml_declaration=True))
    with (OUT/'图1数据.csv').open('w',encoding='utf-8',newline='') as f:
        wr=csv.DictWriter(f,fieldnames=list(plot_rows[0]));wr.writeheader();wr.writerows(plot_rows)

def main():
    OUT.mkdir(exist_ok=True)
    pred=csvread(DATA/'adaptive_predictions.csv');metrics=load(DATA/'metrics.json');adaptive=load(DATA/'adaptive_metrics.json')
    cis=load(ROOT/'reports/nested_validation_20261023/paired_confidence_intervals.json')
    plot(pred)
    package=Document(SOURCE);picture=package.add_paragraph();picture.alignment=1
    picture.add_run().add_picture(str(OUT/'三病种一步滚动预测.png'),width=Inches(5.6))
    memory=BytesIO();package.save(memory)
    with ZipFile(memory) as z:parts={n:z.read(n) for n in z.namelist()}
    doc=E.fromstring(parts['word/document.xml']);body=doc.find(f'{{{W}}}body')
    figure=next(p for p in list(body) if p.tag==f'{{{W}}}p' and p.xpath('.//w:drawing',namespaces=NS));body.remove(figure)
    ps=doc.xpath('//w:p',namespaces=NS)
    refs=[text(p) for p in ps if re.match(r'^\[\d+\]',text(p))]
    equations=[E.tostring(e) for e in doc.xpath('//m:oMath',namespaces=NS)]
    old={i:text(p) for i,p in enumerate(ps)};changes=[]
    def edit(i,s):
        changes.append({'before':text(ps[i]),'after':s});replace(ps[i],s)
    # Renumber retained main tables into first-appearance order; new results are 9..11.
    for p in ps:
        s=text(p);n=re.sub(r'表([789])(?!\d)',lambda m:'表'+{'7':'12','8':'7','9':'8'}[m[1]],s)
        if n!=s:replace(p,n)
    edits={
57:'本课题研究面向多病种疫情监测的可信智能预测系统。系统根据病种数据条件生成候选方案，通过历史滚动评价选择预测模型；简单基线表现更好时保留基线，不强制采用推荐模型。在本人开发的原型上，完善来源追溯、任务画像、候选准入、授权执行及隔离反馈。评价以WAPE和WIS为主，结合预测区间、计算成本与审核记录。已有登革热、弯曲菌肠炎和腮腺炎三病种回顾实验，各含574周观测；滚动选择在2020—2022年的WAPE分别为22.09%、28.96%和46.70%，与同轮最佳固定候选的点估计接近。该结果属于探索性评价，不表示统计等效或未知病原预测有效。预期交付本地系统、代码差分及复核材料；图谱和大语言模型的增量作用单独检验。',
66:'本课题拟减少来源核对和重复选模的人工操作，使分析人员能够追查预测所用的输入、模型及审核状态。系统根据病种数据条件生成候选方案，并通过历史滚动评价选择预测模型；当简单基线表现更好时保留基线，不强制采用推荐模型。方法评价关注这一选择流程在不同病种上的表现与成本，图谱或大语言模型是否带来额外收益另行检验。采集、推荐、运行和发布分别管理权限，数据不满足要求时停止数值预测。',
195:'同类系统已经覆盖知识组织、趋势预测和交互展示。COVID-19 Surveiller结合了社交媒体动态图、趋势预测与交互平台[25]，布鲁氏菌平台将知识图谱用于文献组织和流行病学展示[26]，EpiGIS Pro的收录摘要报告了自动事件提取、多源整合及预测分析[24]。本课题据此将研究对象限定为有来源证据、保留基线的模型选择与授权运行流程。知识图谱用于表达候选的适用条件，实际预测方法由历史评价决定；其增量作用须与仅画像、无图谱等策略比较，模块集成本身不作为创新结论。',
221:'前期已完成香港流感独立模型与AWBE的配对评价，以及将文献时间序列方法迁移至新加坡登革热的实验[27]。这些材料用于说明原型基础和适用边界，不作为当前三病种滚动选择模块的完整智能体成绩。香港H0单列为当期估计，H1—H8分别报告；香港长预测期及新加坡后续年份均出现过误差退化，不能从短期点估计接近推导出全预测期相当。原始结果、逐期差值区间及失败分析保留在附录B.3。',
226:'本课题围绕三个问题组织研究。RQ1检验：在相同数据、候选与预算条件下，将合格基线与推荐模型共同纳入历史滚动比较，能否形成结果可追溯、计算成本可记录的模型选择流程；图谱约束相对仅画像规则是否另有增益。RQ2检验数据身份绑定的授权契约能否拦截无授权、过期及错配调用，并保留完整运行记录。RQ3检验同一流程能否在至少三种病种上完成预测，并分别报告与固定基线的差值、区间质量和失效条件。跨地区及未知病原仅作探索，对照及判定办法见2.4节。',
227:'拟形成的特色是将病种数据条件、候选证据和历史评价分开处理。图谱说明模型为何可能适用，滚动评价检验其是否应被采用；合格基线始终参与比较，系统可以选择持续性、均值或季节性方法。大语言模型提出候选及说明，数值模型产生预测，执行与发布另经授权。反馈在训练侧或已封存历史任务中形成新版本，记录拒绝、回退及测试隔离。研究重点为这一流程的实现和验证，不预设复杂模型、图谱或大语言模型必然提高精度。',
230:'研究目标包括可信数据管道、保留基线的模型选择和授权评估流程。数据管道保存来源、口径、版本及复核状态；选择模块依据病种画像确定准入条件，将合格基线与推荐候选共同比较；授权流程核验任务身份、有效期和执行记录。预测评价按病种与预测期给出WAPE、WIS和差值区间，不要求推荐模型必须胜出。可接受误差应在最终验收前结合任务用途确定，尚无依据的容差不作为已达标结论。',
271:'硬准入用于排除数据或接口条件不满足的模型，软评分用于排列推荐候选。通过硬准入的持续性、均值和季节性基线不因图谱证据缺少或推荐分偏低而退出比较；其准入条件仍须满足。六项评分取0至1，缺失证据记0并说明原因。阈值、距离标准化和权重仅在开发侧确定，同一输入与版本应产生相同结果。表5说明候选生成规则，最终预测选择还需经过历史滚动评价。',
296:'式中S为0至100的推荐适配分，六个权重非负且和为1。首版等权各1/6，另在训练侧嵌套验证中比较步长0.1、总和为1的权重网格，并限制试验预算。展示阈值50分、推荐候选最多3个；最高分不足50时说明证据不足，前两名差值小于5分时并列展示。该上限只约束推荐候选，不包含始终保留的合格基线。通过准入的两类方法进入同一滚动评价，以已观察目标上的WAPE择优；分数不代表预测正确概率，也不保证后续测试期排名不变。',
300:'数值层统一封装持续性、均值、季节性朴素、统计时序、梯度提升及数据量允许时的深度学习模型。候选池由合格基线与推荐候选组成，选择时使用相同历史目标和可核对的WAPE；区间质量以WIS、覆盖率和宽度单独评价，不用适配分替代误差。每次运行保存代码、种子、窗口、参数、依赖、预算与分位数。大语言模型提议与确定性提议分别比较，每任务最多5次提议并记录成本，允许关闭大模型继续运行。上述为统一流程设计；2.7节实测的是四候选数值选择模块，尚非完整大模型联动验证。',
304:'本课题将长期自学习限定为验证、复核和版本化发布后的策略更新。新观测到达后，可以触发数据检查、历史滚动比较和模型复核；当前最佳基线仍可成为下一次运行的选择。异常增长、来源中断或画像漂移另行形成复核任务。新策略在开发侧评价后经授权发布，最终测试标签不反向用于当前版本选模。人工复核优先级与经典主动学习算法分开说明，不把一般反馈更新直接称为主动学习算法创新。',
342:'已用于原型、调参或结果展示的病例及年份归入开发或历史诊断资料。正文2.7节完整展示三病种滚动选择结果，附录B保留年度选择、预算受限及历史退化分析；这一组织方式不改变数据的评价身份。本轮本地审查未确认有未参与开发的独立历史标签集，年度嵌套回测仍属历史验证。独立评价拟先登记来源、时期及哈希，冻结策略后开放标签；未来连续12个发布周的影子评价未完成时报告实际进度。',
346:'对照包括持续性、均值、季节性朴素、训练侧选定固定模型、人工规则、仅画像、随机推荐和预算匹配搜索。完整流程的基线池不受推荐Top-3软筛选影响，各策略使用相同特征、窗口、硬件和目标键。深度模型在样本及资源满足时纳入。选择阶段拟合上限与校准、外层预测拟合分别登记，再核算总成本；30次拟合及最多5次大模型提议仅作为选择阶段预算起点，不与全流程调用数混同。测试期事后最佳固定候选只作参考，不参与当时选模。',
356:'RQ1首先评价候选生成、基线保留、历史选择和运行溯源是否按同一协议完成，再报告预测损失及计算成本。图谱相对仅画像规则的增量作用作为单独比较，不以最终选中基线为失败。Top-k评价仅针对推荐排名，k取1和3，检查是否包含固定候选池中外层损失距最优5%以内的模型；事后最优只生成评价标签，不参与推荐。相对遗憾值按推荐损失减参考最优损失再除以参考最优计算，分母为0时报告绝对差值。',
357:'核心目标是在不同病种上完成可复核预测，并量化相对基线的损失与成本，允许采用简单基线。此前Top-3提高至少10个百分点、拟合次数减少至少30%、中位WIS相对遗憾值不超过5%的数值保留为探索性工程目标，不作为已取得结果或本研究必须成立的结论。是否达到可接受误差，须在新的最终验证前明确用途、指标、容差及依据；现有历史点估计接近不能补写为正式等效或非劣。图谱没有增益或选择策略退化时，附录保留对应证据。',
362:'RQ3按病种与预测期报告预测值、WAPE、WIS及基线差值。正文展示登革热、弯曲菌肠炎和腮腺炎的全部1—4周滚动评价；三种病种共用候选适配器和选择流程，允许各自选择不同模型或基线。历史结果只检验已观察任务中的适用性，跨地区、真正留一病种和未知病原预测仍需另行评价。数据接入成功或图谱请求通过不计作预测泛化证据。',
365:'既有7个测试文件的18项离线检查及其首次失败记录保留。新增授权与推荐组件共有222项测试通过，其中215项针对授权生命周期、参数组合、回放和失败处理，7项针对推荐边界；另有13个既有工作流测试文件的42项测试通过。应用层冒烟测试跑通来源样例、清洗、特征、预测、反馈与任务链，12个API端点检查正常。上述分别属于离线组件与临时SQLite应用测试，不代表生产身份认证、容器启动或真实用户验收。首次两项旧来源数量断言及隔离环境失败的修复经过见附录B.5。',
374:'六维评分示例为D=0.8、M=1、T=0.5、E=1、P=0.5、R=1，等权适配分为80分，只用于核对计算。硬准入失败的候选不执行，合格基线不受推荐分软筛选。已完成90个真实历史截断画像上的54种准入配置诊断，涉及观测数8/16/26、季节历史104/156周、展示阈值40/50/60及候选数1/3/5；结果见附录B.4。滞后可用率70%/80%/90%的专门消融仍属后续工作，不把本次四候选诊断当作全部阈值已验证。',
395:'状态转移须在服务端核验，界面隐藏按钮不能代替权限控制。采集者提交候选数据，分析人员提出模型任务，审核者批准执行与发布，审计角色只读。表7列出状态与故障检查。新增内存执行组件已通过215项授权测试，并接入243个离线预测运行；审批角色在脚本中模拟，不是真实人工授权。生产API身份认证、持久化、多进程并发及完整状态迁移图仍需集成验证，不能将组件通过扩大为系统权限验收。',
418:'三病种已有同源周度回顾数据，各574周，指标、时期、来源和样本数见2.7节；表8保留其他来源的资料状态。正文比较的是同一数值滚动选择模块，不是不同版本结果的事后拼接。历史退化与预算受限分析见附录B，不据较好的主表改变原测试身份。同一国家监测资料也不足以证明跨地区或未知病原适用性。',
419:'现有公开序列可能经过事后修订，图谱知识也可能晚于历史预测时点，均按回顾性证据解释。香港H0与H1—H8的分期差值区间现已补齐，见附录B.3，不再以H0—H1的0.21个百分点代替向前预测结论。新加坡旧时间外推退化仍保留，其可能原因包括分布变化、季节失配和选择滞后，现有结果尚不能确定具体机制。',
421:'本人于2026年10月23日确认项目代码由本人完成，后续按模块、接口、修改前后哈希、差分和测试日志登记；本轮验证实现和文字修订使用AI辅助，第三方库、数据及引用算法保留来源。没有可用Git历史，追溯依据为交付快照与逐行差分。现已增加三病种统计、年度回测、预算原型消融及离线授权和工作流测试；图谱双人审核、未参与开发的数据验证、真实大模型与容器运行条件尚未补齐，导师复核及签字另行确认。',
596:'预期交付可本地运行的多病种疫情监测与辅助预测系统，附部署配置、接口、自动化测试及质量报告。核心流程根据病种数据生成候选，通过历史滚动评价选择方法；简单基线表现更好时保留基线。系统记录来源、模型版本与复核状态，验收覆盖功能、状态机、异常恢复、权限和追溯，页面打开或单个组件测试通过仅作为其中的检查。',
597:'形成版本化预测知识图谱、可复算候选规则及受约束的大语言模型接口，保存文献、断言和个人修改记录。数值预测的适用性、图谱的增量贡献和大模型的协作效果分别评价；不存在额外收益时按结果说明，不将复杂模型替代基线作为成果成立的条件。',
598:'以三病种完整滚动预测为主要历史实证，报告各预测期误差及概率区间质量，结合附录中的香港长预测期、新加坡外推和预算受限案例界定适用范围。独立新窗口、跨地区及未知病原仍按实际完成情况评价。预测与评估分别保存来源、参数、运行时刻、复核状态、协议及目标键，不混合历史、独立验证和前瞻证据。',
762:'现已完成检索留档、代码与数据身份核查、三病种回顾评价、配对差值区间、年度嵌套回测、预算原型消融、准入敏感性、222项新增组件测试及42项既有工作流回归。图谱双人抽检、新独立病例窗口、真实大模型消融、生产授权链路、Docker验收及目标用户研究仍未完成。现阶段结论限定为滚动选择模块与工程组件的证据，代码归属按本人声明登记，导师意见及签字仍需导师确认。'
    }
    for i,s in edits.items():edit(i,s)
    # Earlier comparisons are retained verbatim in the appendix, not erased.
    history_paras=[styled(old[i]) for i in (221,222,223,224)]
    for i in (222,223,224):ps[i].getparent().remove(ps[i])
    # Update planned thesis outline to make the adopted baseline-retention pipeline explicit.
    outline={
      '4.3 六维准入与模型推荐':'4.3 条件准入与候选模型生成','4.3.2 多病种迁移与降级策略':'4.3.2 基线保留与多病种适配',
      '6.3 独立预测方法与文献方法比较':'6.3 三病种滚动预测与基线比较','6.3.1 香港流感分预测期结果':'6.3.1 三病种分预测期误差与模型选择',
      '6.3.2 新加坡登革热方法迁移及区间评价':'6.3.2 概率区间与配对差值评价','6.4 多病种及跨时间泛化':'6.4 跨时间与候选策略适用性',
      '6.6 任务式可用性及局限分析':'6.6 任务式可用性及失败边界分析'}
    for p in ps:
        if text(p) in outline:replace(p,outline[text(p)])
    # Main result blocks: all three diseases, all four forecast horizons, intervals.
    main_blocks=[styled('2.7 三病种滚动预测与基线比较','h2'),styled(old[423]),
      styled('本节检验保留基线的数值滚动选择模块，而非完整知识图谱或大语言模型的端到端增益。四个固定候选为持续性、4周均值、52周季节性及对数AR(2)，各病种共用同一适配器与选择程序。初始训练期2012—2017年、首轮选择期2018年、区间校准期2019年、回顾测试期2020—2022年分别登记。每个预测原点比较最近13个已观察到完整1—4周目标的历史验证原点，按联合绝对误差选出方法，再预测后续目标；同一集合上的绝对误差排序与WAPE排序一致。'),
      styled('每个原点仅使用该位置之前的观测，持续性、均值和季节性候选均可被选中，不要求对数AR(2)胜出。区间沿用相应模型的2019年有符号残差，不在测试期重新校准。每病种157个原点，H1—H4分别含157、156、155和154个目标键，共622个相关配对。资料来自当前存档，不保证等同当时发布值。该滚动方案是在查看首轮固定选择结果后提出的探索性再分析，不能改称独立盲测；固定选择和其他退化版本见附录B。'),
      styled('表9 三病种滚动选择与同轮最佳固定候选的WAPE比较','caption')]
    summary=[];forecast_rows=[];interval_rows=[];best_lookup={}
    for disease,zh in LABELS.items():
        a=next(r for r in adaptive if r['disease']==disease and r['horizon']=='all')
        fixed=[r for r in metrics if r['disease']==disease and r['horizon']=='all'];best=min(fixed,key=lambda r:r['wape']);best_lookup[disease]=best
        ci=next(r for r in cis if r['experiment']=='previous_adaptive' and r['disease']==disease and r['block_origins']==8 and r['comparison']=='rolling_select minus '+best['model'])
        lo,hi=ci['wape_ci95_pp']
        summary.append([zh,f"{a['wape']*100:.2f}%",MODEL[best['model']],f"{best['wape']*100:.2f}%",f"{ci['delta_wape_pp']:+.3f}",f'[{lo:.3f}, {hi:.3f}]'])
        forecast_rows.append([zh]+[f"{next(r for r in adaptive if r['disease']==disease and str(r['horizon'])==str(h))['wape']*100:.2f}%" for h in range(1,5)])
        interval_rows.append([zh,f"{a['wis']:.3f}",f"{a['coverage95']*100:.2f}%",f"{a['width95']:.2f}"])
    main_blocks += [table(['病种','滚动WAPE','参考模型','参考WAPE','差值/百分点','差值95%区间'],summary,[1200,1050,1350,1050,1150,2100]),
      styled('注：参考模型为同一测试目标集上表现最好的固定候选，仅作事后比较，不参与当时选择。差值为滚动方案减参考模型，负值表示前者误差较低；区间来自5,000次原点块自助法，主块长8，4与13作敏感性分析，未作多重比较校正。完整候选与区间见附录B及冻结结果材料。'),
      styled('三病种滚动WAPE分别为22.09%、28.96%和46.70%，相对各自参考模型的差值为0.123、0.013和−0.043个百分点。结果说明，在这些历史任务上，允许采用简单基线的选择模块能够给出与最佳固定候选点估计接近的预测。差值区间均跨零，尚不能宣布统计等效；也不能仅因接近基线就认定绝对误差满足实际预警用途。可接受范围须在新的最终验证前依据任务要求确定。'),
      styled('表10 三病种滚动选择的完整分预测期WAPE','caption'),table(['病种','H1','H2','H3','H4'],forecast_rows,[1600,1575,1575,1575,1575]),
      styled('H1—H4均纳入评价，选模使用1—4周联合损失，未按测试期成绩改为各预测期事后择优。图1展示完整测试时期的一步预测与实际报告，便于观察不同病例尺度下的跟踪情况；H2—H4不因未画入曲线而省略评价。'),
      styled('图1 三病种一步滚动预测与实际周报告病例','caption'),figure,
      styled('图中蓝线为实际报告病例，橙线为一步预测。三图均覆盖2020—2022年的157个预测原点，不截取有利季节。曲线为既有模块的回顾性预测，非前瞻影子运行，也不证明知识图谱或大语言模型的独立贡献。'),
      styled('表11 三病种滚动选择的概率预测质量','caption'),table(['病种','WIS','名义95%区间覆盖率','95%平均宽度/病例数'],interval_rows,[1600,1500,2100,2700]),
      styled('三种疾病的WIS及区间宽度受病例规模影响，不能据数值直接跨病种排序。登革热95%区间实际覆盖率为75.72%，存在明显覆盖不足；弯曲菌肠炎与腮腺炎分别为93.89%和98.39%，仍需结合宽度及WIS判断，不把较高覆盖率单独解释为更优。点预测接近参考与概率校准达标是两项不同要求。'),
      styled('固定首轮选择、年度冻结选择与仅五个验证原点的预算原型均出现过退化，详见附录B.1—B.4。各版本的更新频率、候选准入、目标键及验证长度不同，不能把其成绩混写为一个完整智能体的结果。现有证据支持已观察三病种的数值模块适用性，尚不支持跨地区、未知病原或图谱必然增益。新增组件测试与应用层检查记录见附录B.5。')]
    # Replace the multi-row old 2.7 section, while retaining the following outline.
    first=ps[422];last=ps[499];cell0=first.getparent();cell1=last.getparent();row0=cell0.getparent();row1=cell1.getparent();outer=row0.getparent()
    assert text(last)=='3 论文提纲'
    before=[x for x in list(cell0) if x.tag!=f'{{{W}}}tcPr'][:list(cell0).index(first)-1]
    tail=list(cell1)[list(cell1).index(last):]
    tcpr=deepcopy(cell0.find(f'{{{W}}}tcPr'));i0=list(outer).index(row0);i1=list(outer).index(row1)
    for row in list(outer)[i0:i1+1]:outer.remove(row)
    for offset,row in enumerate(blocks_to_rows(before+main_blocks+tail,tcpr)):outer.insert(i0+offset,row)
    # New appendix contains the complete prior comparison data and all adverse directions.
    appendix=[styled('附录B 模型选择结果与适用边界','h1'),
      styled('本附录集中记录完整候选、不同选择版本和历史退化结果。正文以三病种同一滚动模块说明系统允许选择基线，附录用于审查该模块与其他原型的差异；移位不改变任何预测值或评价身份。',),
      styled('B.1 三病种完整固定候选与滚动选择','h2'),styled(old[424]),styled(old[425]),styled('表B1 完整候选的三病种回顾结果','caption')]
    complete=[]
    for disease,zh in LABELS.items():
        for model in ('log_ar2','mean4','last_value','seasonal52'):
            r=next(r for r in metrics if r['disease']==disease and r['horizon']=='all' and r['model']==model)
            complete.append([zh,MODEL[model],f"{r['wape']*100:.2f}%",f"{r['wis']:.3f}",f"{r['coverage95']*100:.2f}%"])
        r=next(r for r in adaptive if r['disease']==disease and r['horizon']=='all');complete.append([zh,'滚动选择（探索）',f"{r['wape']*100:.2f}%",f"{r['wis']:.3f}",f"{r['coverage95']*100:.2f}%"])
    appendix += [table(['病种','方法','WAPE','WIS','95%覆盖率'],complete,[1450,2200,1200,1450,1600]),styled(old[493]),styled(old[495]),
      styled('B.2 年度冻结选择的嵌套回测','h2'),styled(old[496]),styled(old[497]),styled(old[498]),
      styled('B.3 香港与新加坡历史对照','h2'),*history_paras,
      styled('香港档案中正文模型对应recursive_hist_poisson，已逐预测期核对原汇总表的9行WAPE与RMSE。PROJECT_RECURSIVE_ROUTE为另一组路由输出，不能替代正文方法。表B2以原方法列复算H0及H1—H8；该重算不训练新模型，也不恢复盲测身份。'),styled('表B2 香港原方法与AWBE的分期WAPE及配对区间','caption')]
    hk=load(EXTRA/'hk_canonical_horizon_confidence_intervals.json');hkrows=[]
    for r in hk:
        if r['block']!=8:continue
        h=str(r['horizon']);h=h if h.startswith('H') else 'H'+h;lo,hi=r['wape_ci95_pp']
        hkrows.append([h,str(r['paired_targets']),f"{r['project_wape_pct']:.2f}%",f"{r['awbe_wape_pct']:.2f}%",f"{r['delta_wape_pp']:+.3f}",f'[{lo:.3f}, {hi:.3f}]'])
    appendix += [table(['预测期','配对键','项目WAPE','AWBE','差值/百分点','95%区间'],hkrows,[950,950,1150,1150,1500,2200]),
      styled('H0为当期估计，不能用其较低点估计替代向前预测。H1差值区间跨零，H2—H8差值区间均高于零；H1—H8合并WAPE为项目61.31%、AWBE35.20%。旧新加坡15.69%相对12.18%的退化仍保留，两者均不与正文622键的三病种模块混合。RMSE、块长4/13及原始配对数据随复核材料提供。'),
      styled('B.4 预算受限推荐原型及准入诊断','h2'),
      styled('本轮另实现六维离线原型，并比较无图谱、去机制、去迁移、仅画像、随机、固定规则、有限搜索和冻结模型选择等版本。三病种×三个年度×九策略×三档选择预算，共243次运行，49,410条逐项预测。候选仍为四个数值适配器；每候选仅用五个历史验证原点，推荐最多三候选。它没有实现正文目标中“全部合格基线不受软筛选”的完整候选保障，故其退化用于诊断，不能当作已实现最终统一流程的成绩。'),
      styled('图谱只将明确病种规则中的基线要求视为部分适用性证据，E取0.5；不支持病种的M与E记0，不假造疾病特异证据。选择预算为5/15/30次拟合，校准及外层预测成本另计；每病种三年共610键，与正文622键不同。当前图谱知识可能晚于历史原点，实验按探索性回顾分析解释。'),styled('表B3 30次选择拟合上限下的预算原型结果','caption')]
    extra=load(EXTRA/'recommendation_aggregate_metrics.json');xr=[]
    for disease,zh in LABELS.items():
        for policy,label in (('six_dim','六维原型'),('profile_only','仅画像'),('budget_search','有限搜索')):
            r=next(r for r in extra if (r['disease'],r['policy'],r['budget'])==(disease,policy,30))
            xr.append([zh,label,f"{r['wape']*100:.2f}%",f"{r['wis']:.3f}",str(r['selection_fits']),str(r['total_fits'])])
    appendix += [table(['病种','策略','WAPE','WIS','选择拟合','总拟合'],xr,[1400,1600,1200,1300,1150,1250]),
      styled('本候选池中六维、仅画像、无图谱等版本预测相同，配对差值为0，未观察到图谱增量收益；这不证明其他候选池中的图谱均无效。Top-3为9/9，但仅四候选与九个相关任务，且仅画像相同，不能作为六维增量证据。有限搜索每任务20次选择拟合、推荐15次，选择阶段少25%，总成本还包含校准与外层预测，不能声称整体少25%。5/15/30分钟上限均未触及，不支持墙钟效率结论；真实大模型消融未执行。'),
      styled('准入诊断使用三病种、三个年度及4/7/8/15/16/25/26/52/104/156周真实截断历史，共90个画像，比较最低观测8/16/26、季节历史104/156、阈值40/50/60和候选数1/3/5，共54种配置及4,860条记录。多候选输出按预先指定等权平均，不按测试误差挑选。1,944条记录拒绝输出，占40%；这不是独立用户任务的拒绝率，拒绝任务不填0误差。完整配置保留有效任务数、WAPE、拟合数和耗时，不依据测试结果倒定可接受阈值。'),
      styled('B.5 授权组件与本地工作流检查','h2'),
      styled('新增ContractGate绑定数据快照、任务、图谱、模型、参数、原点、输出范围及风险复核状态，支持待审批、批准、运行、完成、失败、撤销和过期。222项测试全部通过：60个合法组合、150个非法组合、5项生命周期及并发等检查、7项推荐边界。非法组合均未执行预测回调，合法同键回放只执行一次；24个并发同键调用的实际回调为1次。243个预算预测运行均经过该组件。参数组合数不等于独立攻击类型数，模拟角色不是人工批准。'),
      styled('组件状态保存在本进程内存，尚未完成生产API认证、持久化任务、多进程锁和真实权限链路集成。42项既有回归测试覆盖API、看板格式和安全、worker、访问策略及Compose静态检查。首次40项通过、2项失败，原因是旧断言要求6个来源，而已有配置包含第7项CDC FluView；核对其默认禁用、审核与单请求限额后更新断言并增强保护检查，未修改或放宽访问策略。首次失败和修改前后指纹均留存。'),
      styled('应用层冒烟最终通过，12个API端点均正常，使用合成样例和临时SQLite，不将样例预测误差计作真实疾病精度。先后遇到临时目录访问及全socket禁用阻断Windows异步回环的问题，仅调整隔离环境后通过，业务代码和断言未改。外部连接禁用，只允许本机回环。本机无Docker CLI，Compose只有静态检查；无可用本地大模型运行条件，未执行真实LLM消融，未启动线上服务。'),
      styled('B.6 现阶段结论与待完成验证','h2'),
      styled('现阶段可说明三病种数值模块已完成同协议滚动评价，允许简单基线成为最终输出；点预测与同轮最佳固定候选接近，但没有等效、临床可接受或未知病原有效性的证明。六维预算原型与年度冻结选择的退化表明，更新频率、验证长度和候选保障需要在开发侧继续检验。后续以全部合格基线共同入池的冻结流程作为验证对象，不用事后最佳模型替代实际选择器。'),
      styled('图谱双人断言审核、目标用户任务评价、负责人或导师确认需要真人参与；12周影子评价需要新发布数据及真实日历。新独立窗口、图谱/LLM增量作用、状态迁移图与生产权限集成、Docker启动及临床用途容差仍有证据缺口。主线调整并不取消这些要求，也不将既有回顾结果改称最终独立验证。代码归属按本人声明登记，导师原评语保留，不代表本次修订已获新的签字确认。')]
    teacher=next(p for p in doc.xpath('//w:p',namespaces=NS) if text(p).startswith('学术导师对该生开题报告的简要评语'))
    # Start appendix and the unchanged review form on fresh pages; prevent orphan headings.
    appendix[0].find(f'{{{W}}}pPr').extend([w('pageBreakBefore'),w('keepNext')])
    teacher.find(f'{{{W}}}pPr').append(w('pageBreakBefore'))
    teacher_row=teacher.getparent().getparent();report_table=teacher_row.getparent();position=list(report_table).index(teacher_row)
    for offset,row in enumerate(blocks_to_rows(appendix,tcpr)):report_table.insert(position+offset,row)
    assert refs==[text(p) for p in doc.xpath('//w:p',namespaces=NS) if re.match(r'^\[\d+\]',text(p))]
    assert equations==[E.tostring(e) for e in doc.xpath('//m:oMath',namespaces=NS)]
    # Every retained original Table10 numeric tuple appears in the expanded full-candidate table.
    for disease,zh in LABELS.items():
        for model in ('log_ar2','mean4','last_value'):
            r=next(r for r in metrics if r['disease']==disease and r['horizon']=='all' and r['model']==model)
            assert any(row[0]==zh and row[1]==MODEL[model] and row[2]==f"{r['wape']*100:.2f}%" for row in complete)
    parts['word/document.xml']=serialize(doc)
    with ZipFile(DEST,'w',ZIP_DEFLATED) as z:
        for name,data in parts.items():z.writestr(name,data)
    (OUT/'文档修改核查.json').write_text(json.dumps({'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'source':SOURCE.name,'output':DEST.name,
        'references_preserved':len(refs),'native_equations_preserved':len(equations),'main_result_rows':summary,'forecast_horizon_rows':forecast_rows,
        'prior_failure_results_retained_in':'Appendix B','edits':changes,'data_files':{p.relative_to(ROOT).as_posix():hashlib.sha256(p.read_bytes()).hexdigest() for p in (DATA/'adaptive_predictions.csv',DATA/'metrics.json',DATA/'adaptive_metrics.json')}} ,ensure_ascii=False,indent=2),encoding='utf-8')
    print(DEST)

if __name__=='__main__':main()
