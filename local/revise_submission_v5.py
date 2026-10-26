from __future__ import annotations
from pathlib import Path
import shutil, hashlib, json, zipfile, tempfile
from docx import Document
from docx.shared import Inches
from PIL import Image

ROOT=Path(r'F:\作业\aibola\reports\nc_submission_revision_20260824\NC_Submission_Updated_20260824')
MAIN_OLD=ROOT/'DiseaseX_Nexus_NC_Manuscript_v4_20260824.docx'
MAIN_NEW=ROOT/'DiseaseX_Nexus_NC_Manuscript_v5_20260824.docx'
SI_OLD=ROOT/'Supplementary'/'Supplementary_Information_Updated_20260824.docx'
SI_NEW=ROOT/'Supplementary'/'Supplementary_Information_Updated_v2_20260824.docx'
MAN_OLD=ROOT/'Content_Manifest_Updated_20260824.docx'
MAN_NEW=ROOT/'Content_Manifest_Updated_v2_20260824.docx'
PYFIG=ROOT/'Supplementary'/'Python_Generated_Figures_20260824'
UI=ROOT/'Supplementary'/'UI_Screenshots_20260824'
EXP=ROOT/'Supplementary'/'Table4_Figure6_Experiment_20260824'
SIFIG=ROOT/'Supplementary'/'SI_Figures'
ARCH=ROOT/'Archive_Previous_20260824_Pre_Page21_CodeFigure_Revision'
ARCH.mkdir(exist_ok=True)

def sha(p:Path)->str: return hashlib.sha256(p.read_bytes()).hexdigest()
def backup(p:Path):
    if p.exists(): shutil.copy2(p,ARCH/p.name)
for p in [MAIN_OLD,SI_OLD,MAN_OLD,ROOT/'Figures'/'figure3.jpg',SIFIG/'Main_Figure_1_System_Architecture.jpeg',SIFIG/'Main_Figure_3_Brazil_Dengue.jpeg',SIFIG/'Main_Figure_6_Retrospective_Multi_Pathogen_Evaluation.png',ROOT/'Figures'/'Figure_6_Retrospective_Multi_Pathogen_Evaluation.png']:
    backup(p)

# Copy source scripts into the auditable package.
for src in [Path(r'F:\作业\aibola\generate_python_main_figures.py'),Path(r'F:\作业\aibola\generate_python_main_figures_v2.py')]: shutil.copy2(src,PYFIG/src.name)
for src in [Path(r'F:\作业\aibola\generate_revised_figure6_pillow.py')]: shutil.copy2(src,EXP/src.name)
for src in [Path(r'F:\作业\aibola\capture_submission_ui_screenshots.cjs'),Path(r'F:\作业\aibola\capture_comparison_ui_screenshot.cjs'),Path(r'F:\作业\aibola\rerun_submission_screenshots_with_csp_bypass.cjs')]: shutil.copy2(src,UI/src.name)

# Active external assets: only Python-generated figures are used for Figs. 1, 3 and 6.
fig1=PYFIG/'Main_Figure_1_System_Architecture_Python.png'
fig3=PYFIG/'Main_Figure_3_Brazil_Dengue_Python.png'
fig6=EXP/'Figure_6_Retrospective_Multi_Pathogen_Evaluation_v2.png'
def as_jpeg(src:Path,dst:Path):
    with Image.open(src) as im: im.convert('RGB').save(dst,'JPEG',quality=96,optimize=True)
as_jpeg(fig1,SIFIG/'Main_Figure_1_System_Architecture.jpeg')
as_jpeg(fig3,SIFIG/'Main_Figure_3_Brazil_Dengue.jpeg')
as_jpeg(fig3,ROOT/'Figures'/'figure3.jpg')
shutil.copy2(fig6,SIFIG/'Main_Figure_6_Retrospective_Multi_Pathogen_Evaluation.png')
shutil.copy2(fig6,ROOT/'Figures'/'Figure_6_Retrospective_Multi_Pathogen_Evaluation.png')
shutil.copy2(fig6,EXP/'Figure_6_Retrospective_Multi_Pathogen_Evaluation.png')

old_media={
 sha(ARCH/'Main_Figure_1_System_Architecture.jpeg'):fig1.read_bytes(),
 sha(ARCH/'Main_Figure_3_Brazil_Dengue.jpeg'):fig3.read_bytes(),
 sha(ARCH/'Main_Figure_6_Retrospective_Multi_Pathogen_Evaluation.png'):fig6.read_bytes(),
}
def patch_media(docx:Path)->int:
    tmp=docx.with_suffix('.media.tmp.docx'); count=0
    with zipfile.ZipFile(docx,'r') as zin,zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data=zin.read(info.filename)
            if info.filename.startswith('word/media/'):
                h=hashlib.sha256(data).hexdigest()
                if h in old_media:
                    data=old_media[h]; count+=1
            zout.writestr(info,data)
    tmp.replace(docx)
    return count

# Main Article revisions.
doc=Document(MAIN_OLD)
replacements={
 'The benchmark used the official repository for Tsang et al. 9,25 at commit b0bd4ccfd9a4a4183c07f122508056a84d271a0f. The official Data.zip SHA-256 was 27c3003babd3b4ded5cf5c37a670cacba97028062c7ff3b1e24f82777254237c and Results_point.zip SHA-256 was f7c289b81ec3d5cf9b075ca8d8fb903b1da20efba29cba312eaea982454b03de. The primary evaluation contained 4,968 common origin–horizon–target rows over horizons 0–8. The external post-COVID evaluation contained 423 common rows with targets from 5 March 2023 through 21 January 2024. Missing predictions were not imputed to zero.':
 'We evaluated the Hong Kong influenza benchmark using the official author-released surveillance data and forecast outputs associated with Tsang et al. 9,25. The archived release was fixed before analysis and its integrity was verified against a version-controlled provenance record (Supplementary Note 4 and Supplementary Table 14). The primary evaluation comprised 4,968 common origin–horizon–target observations across horizons 0–8. An external post-COVID evaluation comprised 423 common observations with target dates from 5 March 2023 to 21 January 2024. Missing forecasts were retained as missing and were not replaced by zero.',
 'DiseaseX searched 450 adaptive orchestration configurations using development RMSE only. The locked candidate nnls_w26_hl8_p0 used a 26-week window and an 8-week half-life over the released component panel. The selection-input SHA-256 was 5233f22784cb7d276cc318d400c5935bc143396a91f147015bfea7150b72b0c2 and the selection-lock SHA-256 was 1e213cfe5e938d67cc9d9a6ea8528182841a9e022dec207760585e1ad6c529e4. Horizon routing was also selected on development data: horizon 0 used the DiseaseX ensemble and horizons 1–8 used AWBE. The same lock was applied without retuning to the post-COVID period.':
 'DiseaseX evaluated 450 adaptive orchestration configurations using development-period RMSE only. The selected configuration used a 26-week rolling window and exponentially decaying weights with an 8-week half-life over the released component-forecast panel. Horizon-specific routing was likewise selected exclusively on development data: the DiseaseX ensemble was used at horizon 0 and the author-released AWBE forecast at horizons 1–8. The configuration and routing rule were then frozen and applied without retuning to the post-COVID period. Repository revision, artifact checksums and the immutable selection record are reported in the Supplementary reproducibility inventory rather than in the main text.'
}
seen=set()
for p in doc.paragraphs:
    if p.text in replacements: p.text=replacements[p.text]; seen.add(p.text)
    if p.text.startswith('Fig. 3 Brazil dengue data profile and R3C admission.'):
        p.text=('Fig. 3 Brazil dengue trusted observations and data-quality audit. Weekly new cases and deaths are shown for 74 accepted reporting periods per metric from January 2024 to July 2026. The long reporting gap remains explicit and was not interpolated or filled with zero. The side panel reports source, quality, revision and lineage metadata from the authenticated read-only history API; generating the figure did not train, predict or publish any result.')
if len(seen)!=2: raise RuntimeError(f'main paragraph replacements={len(seen)} expected 2')
doc.save(MAIN_NEW)
main_media=patch_media(MAIN_NEW)

# Unified Supplementary Information revisions.
si=Document(SI_OLD)
for p in si.paragraphs:
    if p.text.startswith('The benchmark was based on Tsang et al. (Nature Communications 15, 8625, 2024; DOI 10.1038/s41467-024-52504-1) and the official repository commit'):
        p.text=('The benchmark was based on Tsang et al. (Nature Communications 15, 8625, 2024; DOI 10.1038/s41467-024-52504-1) and a fixed revision of the official author repository. Reproducibility identifiers were recorded before evaluation: repository revision b0bd4ccfd9a4a4183c07f122508056a84d271a0f; official data-artifact SHA-256 27c3003babd3b4ded5cf5c37a670cacba97028062c7ff3b1e24f82777254237c; released point-forecast-artifact SHA-256 f7c289b81ec3d5cf9b075ca8d8fb903b1da20efba29cba312eaea982454b03de; selection-input SHA-256 5233f22784cb7d276cc318d400c5935bc143396a91f147015bfea7150b72b0c2; and immutable selection-record SHA-256 1e213cfe5e938d67cc9d9a6ea8528182841a9e022dec207760585e1ad6c529e4. Artifact roles, rather than archive filenames, are used in the scientific narrative.')
si.add_page_break(); si.add_heading('Supplementary Note 10. Code-generated figures and authenticated interface evidence',level=1)
si.add_paragraph('Main Figs. 1, 3 and 6 were rendered deterministically by Python from documented architecture elements, the authenticated read-only history export, and frozen experiment CSV files, respectively. Main Fig. 2 is the existing knowledge-graph illustration and was intentionally left unchanged. Supplementary Figs. 5–7 are English presentation renderings generated by Playwright from values verified in an authenticated read-only Dashboard session. No generative-image output is used in the revised main or supplementary figures.')
for img,caption in [
 (UI/'Supplementary_Figure_5_Authenticated_KG_Recommendation_UI.png','Supplementary Fig. 5 | Authenticated knowledge-graph recommendation interface. The page reports the selected Austria legionellosis profile, the pre-execution evidence recommendation and its decision boundary. Labels were rendered in English for publication; values were verified against the live read-only interface.'),
 (UI/'Supplementary_Figure_6_Authenticated_Model_Evaluation_UI.png','Supplementary Fig. 6 | Authenticated held-out model-evaluation interface. The Chile influenza Track C window was globally consumed once. WAPE, agreement, coverage and common-pair baseline results are shown; the fresh blind-test gates failed and no certification is claimed.'),
 (UI/'Supplementary_Figure_7_Authenticated_Comparison_Experiment_UI.png','Supplementary Fig. 7 | Authenticated same-protocol comparison interface for the Singapore dengue retrospective benchmark. The panel shows the locked data split, validation-selected ensemble and test results alongside the paper AR(2), last-value and project single-model comparators.'),
]:
    si.add_page_break(); si.add_picture(str(img),width=Inches(6.55)); cap=si.add_paragraph(caption); cap.style='Caption'
si.save(SI_NEW)
si_media=patch_media(SI_NEW)

# Updated content manifest.
man=Document(MAN_OLD)
man.paragraphs[1].text='Updated 2026-08-24 | Page-21 prose normalization, Python-generated figures and authenticated English interface evidence'
for p in man.paragraphs:
    if p.text.startswith('TABLE4_FIGURE6_RETROSPECTIVE_OPTIMIZATION_GATES_PASSED'):
        p.text=('TABLE4_FIGURE6_RETROSPECTIVE_OPTIMIZATION_GATES_PASSED. This remains a retrospective engineering evaluation, not fresh held-out or prospective certification. Main-text artifact filenames and hashes were moved to the Supplementary reproducibility inventory. Main Figs. 1, 3 and 6 are now Python-generated; Main Fig. 2 remains unchanged by instruction.')
t=man.tables[1]
for row in t.rows[1:]:
    if 'DiseaseX_Nexus_NC_Manuscript' in row.cells[0].text:
        row.cells[0].text='DiseaseX_Nexus_NC_Manuscript_v5_20260824.docx/.pdf'; row.cells[1].text='Main manuscript with conventional benchmark prose and Python-generated Figs. 1, 3 and 6'
    elif 'Supplementary_Information' in row.cells[0].text:
        row.cells[0].text='Supplementary/Supplementary_Information_Updated_v2_20260824.docx/.pdf'; row.cells[1].text='Unified SI with artifact-role reproducibility inventory and English authenticated interface evidence'
new=t.add_row().cells; new[0].text='Supplementary/Python_Generated_Figures_20260824/'; new[1].text='Python source, authenticated Brazil dengue export and deterministic Main Figs. 1 and 3'
new=t.add_row().cells; new[0].text='Supplementary/UI_Screenshots_20260824/'; new[1].text='Playwright source, three English interface renderings and zero-error browser audit'
man.save(MAN_NEW)

prov={
 'direct_generative_images_used':False,
 'knowledge_graph_figure_changed':False,
 'figures':{
  'Main Figure 1':{'generator':'generate_python_main_figures_v2.py','output':str(fig1.relative_to(ROOT)),'sha256':sha(fig1)},
  'Main Figure 2':{'generator':'unchanged existing project illustration','output':str((SIFIG/'Main_Figure_2_KG_Provenance.jpeg').relative_to(ROOT)),'sha256':sha(SIFIG/'Main_Figure_2_KG_Provenance.jpeg')},
  'Main Figure 3':{'generator':'generate_python_main_figures_v2.py','input':str((PYFIG/'Brazil_Dengue_History_Live.json').relative_to(ROOT)),'output':str(fig3.relative_to(ROOT)),'sha256':sha(fig3)},
  'Main Figure 6':{'generator':'generate_revised_figure6_pillow.py','input':str((EXP/'Figure_6_Data.csv').relative_to(ROOT)),'output':str(fig6.relative_to(ROOT)),'sha256':sha(fig6)},
  'Supplementary Figures 5-7':{'generator':'Playwright authenticated read-only rendering','audit':str((UI/'UI_SCREENSHOT_AUDIT.json').relative_to(ROOT))},
 },
 'docx_media_replacements':{'main':main_media,'supplementary':si_media}
}
(ROOT/'FIGURE_PROVENANCE_20260824.json').write_text(json.dumps(prov,indent=2,ensure_ascii=False),encoding='utf-8')
(ROOT/'NO_GENERATIVE_IMAGE_ATTESTATION.md').write_text('# Figure provenance attestation\n\nThe revised active manuscript and Supplementary Information use no direct generative-image output. Main Figs. 1, 3 and 6 are generated by Python source included in the package. Main Fig. 2 is the pre-existing knowledge-graph illustration and was left unchanged by instruction. Supplementary Figs. 5–7 are browser-rendered from an authenticated read-only session with English publication labels.\n',encoding='utf-8')
print(json.dumps({'main':str(MAIN_NEW),'si':str(SI_NEW),'manifest':str(MAN_NEW),'main_media_replaced':main_media,'si_media_replaced':si_media},ensure_ascii=False))
