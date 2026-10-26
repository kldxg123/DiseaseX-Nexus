from __future__ import annotations

import hashlib
import json
import shutil
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document


ROOT = Path(r"F:\作业\aibola\reports\nc_submission_revision_20260824")
SRC = ROOT / "NC_Submission_Updated_20260824"
DEST = ROOT / "NC_Submission_Final_v3_20260824_EnglishFiguresEmbedded"
ZIP_PATH = ROOT / "NC投稿完善_20260824_主文与Supplementary英文图终版v3.zip"

MAIN_DOCX = "DiseaseX_Nexus_NC_Manuscript_v7_20260824.docx"
MAIN_PDF = "DiseaseX_Nexus_NC_Manuscript_v7_20260824.pdf"
SI_DOCX = Path("Supplementary") / "Supplementary_Information_Updated_v4_20260824.docx"
SI_PDF = Path("Supplementary") / "Supplementary_Information_Updated_v4_20260824.pdf"
MANIFEST_DOCX = "Content_Manifest_Updated_v5_20260824.docx"
MANIFEST_PDF = "Content_Manifest_Updated_v5_20260824.pdf"

EXCLUDED_DIR_NAMES = {
    "Archive_Previous_20260823",
    "Archive_Previous_20260824_Pre_Page21_CodeFigure_Revision",
    "_qa_main_v5",
    "_qa_main_v6",
    "_qa_si_v2",
    "_qa_si_v3",
    "_qa_manifest_v2",
    "_qa_manifest_v3",
    "_qa_si_v4",
    "_qa_manifest_v4",
    "_qa_main_v7",
    "_qa_manifest_v5",
    "Python_Generated_Figures_20260824",
}

EXCLUDED_FILE_NAMES = {
    "DiseaseX_Nexus_NC_Manuscript_v4_20260824.docx",
    "DiseaseX_Nexus_NC_Manuscript_v4_20260824.pdf",
    "DiseaseX_Nexus_NC_Manuscript_v5_20260824.docx",
    "DiseaseX_Nexus_NC_Manuscript_v5_20260824.pdf",
    "DiseaseX_Nexus_NC_Manuscript_v6_20260824.docx",
    "DiseaseX_Nexus_NC_Manuscript_v6_20260824.pdf",
    "Supplementary_Information_Updated_20260824.docx",
    "Supplementary_Information_Updated_20260824.pdf",
    "Supplementary_Information_Updated_v2_20260824.docx",
    "Supplementary_Information_Updated_v2_20260824.pdf",
    "Supplementary_Information_Updated_v3_20260824.docx",
    "Supplementary_Information_Updated_v3_20260824.pdf",
    "Content_Manifest_Updated_20260824.docx",
    "Content_Manifest_Updated_20260824.pdf",
    "Content_Manifest_Updated_v2_20260824.docx",
    "Content_Manifest_Updated_v2_20260824.pdf",
    "Content_Manifest_Updated_v3_20260824.docx",
    "Content_Manifest_Updated_v3_20260824.pdf",
    "Content_Manifest_Updated_v4_20260824.docx",
    "Content_Manifest_Updated_v4_20260824.pdf",
    "PACKAGE_VALIDATION_20260824.json",
    "PACKAGE_VALIDATION_20260824_v2.json",
    "SHA256SUMS.txt",
    "矢量图.zip",
    "Figure_6_Retrospective_Multi_Pathogen_Evaluation.png",
    "Figure_6_Retrospective_Multi_Pathogen_Evaluation_v2.png",
    "Figure_6_Retrospective_Multi_Pathogen_Evaluation.svg",
    "Main_Figure_6_Retrospective_Multi_Pathogen_Evaluation.png",
    "FIGURE_PROVENANCE_20260824.json",
}

SUPPLEMENTARY_FIGURES = tuple(
    Path("Supplementary") / "SI_Figures" / name
    for name in (
        "Supplementary_Figure_1_HK_Flu_Main_RMSE.png",
        "Supplementary_Figure_2_HK_Flu_Horizon_RMSE.png",
        "Supplementary_Figure_3_HK_Flu_PostCOVID_RMSE.png",
        "Supplementary_Figure_4_Governed_Tracks.png",
    )
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def doc_text(path: Path) -> str:
    doc = Document(path)
    blocks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            blocks.extend(cell.text for cell in row.cells)
    return "\n".join(blocks)


def contains_cjk(text: str) -> bool:
    return any("\u4e00" <= char <= "\u9fff" for char in text)


def should_copy(relative: Path) -> bool:
    if any(part in EXCLUDED_DIR_NAMES for part in relative.parts):
        return False
    if relative.name in EXCLUDED_FILE_NAMES:
        return False
    return True


def build() -> None:
    if DEST.exists():
        shutil.rmtree(DEST)
    DEST.mkdir(parents=True)

    copied = 0
    for source in SRC.rglob("*"):
        relative = source.relative_to(SRC)
        if not source.is_file() or not should_copy(relative):
            continue
        target = DEST / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, target)
        copied += 1

    main_text = doc_text(DEST / MAIN_DOCX)
    si_text = doc_text(DEST / SI_DOCX)
    manifest_text = doc_text(DEST / MANIFEST_DOCX)
    provenance = json.loads((DEST / "FIGURE_PROVENANCE_20260824_v3.json").read_text(encoding="utf-8"))

    required_files = (
        Path(MAIN_DOCX), Path(MAIN_PDF), SI_DOCX, SI_PDF,
        Path(MANIFEST_DOCX), Path(MANIFEST_PDF),
        Path("Figures") / "Figure_6_Evaluation_and_Ablation.png",
        Path("Supplementary") / "SI_Figures" / "Main_Figure_6_Evaluation_and_Ablation.png",
        Path("Supplementary") / "Table4_Figure6_Experiment_20260824" / "Figure_6_Evaluation_and_Ablation.png",
        *SUPPLEMENTARY_FIGURES,
    )
    missing_files = [item.as_posix() for item in required_files if not (DEST / item).is_file()]

    forbidden_caption = "Retrospective multi-pathogen evaluation and ablation"
    checks = {
        "all_required_files_present": not missing_files,
        "main_docx_opens": bool(main_text),
        "supplementary_docx_opens": bool(si_text),
        "manifest_docx_opens": bool(manifest_text),
        "supplementary_text_has_no_cjk": not contains_cjk(si_text),
        "manifest_text_has_no_cjk": not contains_cjk(manifest_text),
        "figure_6_caption_shortened": forbidden_caption not in main_text and forbidden_caption not in si_text,
        "figure_6_bra_no_event_explained": "all 13 test origins contained no positive warning event" in main_text and "all 13 test origins contained no positive warning event" in si_text,
        "figure_1_restored_original": provenance["figure_1"]["status"] == "RESTORED_ORIGINAL",
        "figure_3_restored_original": provenance["figure_3"]["status"] == "RESTORED_ORIGINAL",
        "supplementary_figures_python_generated_english": all(item["status"] == "PYTHON_GENERATED_ENGLISH" for item in provenance["supplementary_figures"]),
        "knowledge_graph_figure_unchanged": provenance["knowledge_graph_figure_changed"] is False,
        "supplementary_v4_embedding_attestation_present": (DEST / "Supplementary" / "SUPPLEMENTARY_ENGLISH_FIGURE_EMBEDDING_V4.json").is_file(),
        "main_v7_embedding_attestation_present": (DEST / "MAIN_ENGLISH_FIGURE_EMBEDDING_V7.json").is_file(),
        "legacy_figure_6_assets_excluded": not any(path.name in EXCLUDED_FILE_NAMES for path in DEST.rglob("*")),
    }
    status = "PASS" if all(checks.values()) else "FAIL"

    hashes: list[tuple[str, str]] = []
    for file_path in sorted(path for path in DEST.rglob("*") if path.is_file()):
        relative = file_path.relative_to(DEST).as_posix()
        if relative in {"SHA256SUMS.txt", "PACKAGE_VALIDATION_20260824_v5.json"}:
            continue
        hashes.append((sha256(file_path), relative))
    (DEST / "SHA256SUMS.txt").write_text(
        "".join(f"{digest}  {relative}\n" for digest, relative in hashes),
        encoding="utf-8",
    )

    validation = {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "status": status,
        "copied_file_count_before_validation_files": copied,
        "missing_files": missing_files,
        "checks": checks,
        "sha256_entry_count": len(hashes),
        "figure_6_bra_dengue_reason": "NOT_EVALUABLE_NO_POSITIVE_WARNING_EVENTS_IN_13_TEST_ORIGINS",
    }
    (DEST / "PACKAGE_VALIDATION_20260824_v5.json").write_text(
        json.dumps(validation, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    if status != "PASS":
        raise RuntimeError(json.dumps(validation, ensure_ascii=False, indent=2))

    if ZIP_PATH.exists():
        ZIP_PATH.unlink()
    with zipfile.ZipFile(ZIP_PATH, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for file_path in sorted(path for path in DEST.rglob("*") if path.is_file()):
            archive.write(file_path, (DEST.name / file_path.relative_to(DEST)).as_posix())
    with zipfile.ZipFile(ZIP_PATH) as archive:
        corrupt = archive.testzip()
        member_count = len(archive.namelist())
    if corrupt:
        raise RuntimeError(f"ZIP CRC failure: {corrupt}")
    print(json.dumps({**validation, "final_directory": str(DEST), "zip_path": str(ZIP_PATH), "zip_sha256": sha256(ZIP_PATH), "zip_member_count": member_count, "zip_crc": "PASS"}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    build()




