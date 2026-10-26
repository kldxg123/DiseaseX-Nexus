from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


INPUT = Path(r"F:\作业\aibola\reports\comparison_experiment_summary_20260821\疫情预测论文对比试验现状与方法总结_20260821.docx")
OUTPUT = Path(r"F:\作业\aibola\reports\comparison_experiment_summary_20260821\疫情预测论文对比试验现状与方法总结_20260821.fixed.docx")


def next_id(nodes, attr):
    values = [int(n.get(qn(attr))) for n in nodes if n.get(qn(attr)) is not None]
    return max(values, default=0) + 1


def add_decimal_numbering(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    abstract_nodes = numbering.findall(qn("w:abstractNum"))
    num_nodes = numbering.findall(qn("w:num"))
    abstract_id = next_id(abstract_nodes, "w:abstractNumId")
    num_id = next_id(num_nodes, "w:numId")

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "decimal")
    lvl.append(fmt)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "%1.")
    lvl.append(text)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    lvl.append(jc)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    ppr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "100")
    spacing.set(qn("w:line"), "264")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    lvl.append(ppr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def assign_numbering(paragraph, num_id: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is not None:
        ppr.remove(numpr)
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    numpr.append(ilvl)
    numpr.append(num)
    ppr.insert(0, numpr)


def main() -> None:
    doc = Document(INPUT)
    groups = [
        [
            "论文 ARIMA 的 WAPE",
            "项目组合模型 WAPE",
            "Qwen受控调优",
            "项目组合模型仍略逊",
            "项目模型的主要优势",
        ],
        [
            "每个 origin 只读取",
            "小波、层级、SARMA",
            "缩放器每折只在",
            "对未来标签、Qwen",
        ],
    ]
    for prefixes in groups:
        num_id = add_decimal_numbering(doc)
        for paragraph in doc.paragraphs:
            if any(paragraph.text.startswith(prefix) for prefix in prefixes):
                assign_numbering(paragraph, num_id)

    # Remove the explicit blank page-break paragraph immediately before section 4.
    for idx, paragraph in enumerate(list(doc.paragraphs)):
        if paragraph.text.startswith("4. 已完成实验的主要结果") and idx > 0:
            previous = doc.paragraphs[idx - 1]
            if previous.text == "" and previous._p.findall(".//" + qn("w:br")):
                previous._element.getparent().remove(previous._element)
            break

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
