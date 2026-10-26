from __future__ import annotations

from pathlib import Path
from typing import Iterable

import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"F:\作业\aibola")
OUT_DIR = ROOT / "reports" / "comparison_experiment_summary_20260821"
ASSET_DIR = ROOT / "tmp_docx_report" / "assets"
OUT_PATH = OUT_DIR / "疫情预测论文对比试验现状与方法总结_20260821.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
LIGHT_GRAY = "F2F4F7"
BLUE_GRAY = "E8EEF5"
PALE_BLUE = "EAF3F8"
GOLD = "7A5A00"
RED = "9B1C1C"
MUTED = "5B6573"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_in: Iterable[float], indent_dxa: int = 120) -> None:
    widths = [int(round(v * 1440)) for v in widths_in]
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for i, cell in enumerate(row.cells):
            width = widths[min(i, len(widths) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, *, size: float | None = None, bold: bool | None = None,
                 color: str | None = None, italic: bool | None = None) -> None:
    run.font.name = "Calibri"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def format_cell_text(cell, *, size: float = 8.5, bold: bool = False, color: str = "000000") -> None:
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        for run in p.runs:
            set_run_font(run, size=size, bold=bold, color=color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float],
              font_size: float = 8.5) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        hdr.cells[i].text = text
        set_cell_shading(hdr.cells[i], BLUE_GRAY)
        format_cell_text(hdr.cells[i], size=font_size, bold=True, color=NAVY)
    for values in rows:
        row = table.add_row()
        for i, text in enumerate(values):
            row.cells[i].text = text
            if len(table.rows) % 2 == 1:
                set_cell_shading(row.cells[i], "FAFBFC")
            format_cell_text(row.cells[i], size=font_size)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.5 if level == 0 else 0.75)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.10
    run = p.add_run(text)
    set_run_font(run, size=10.5)


def add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.10
    run = p.add_run(text)
    set_run_font(run, size=10.5)


def add_paragraph(doc: Document, text: str, *, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=10.5, bold=True, color=NAVY)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=10.5)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5)


def add_callout(doc: Document, label: str, text: str, *, fill: str = PALE_BLUE,
                label_color: str = DARK_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.10
    r1 = p.add_run(label + "  ")
    set_run_font(r1, size=10.5, bold=True, color=label_color)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=8.5, color=MUTED)


def setup_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Title", 24, NAVY, 0, 6),
        ("Subtitle", 12.5, MUTED, 0, 12),
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 11.5, DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        if name != "Subtitle":
            style.font.bold = True

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def build_chart(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    labels = ["论文 ARIMA", "Last-value", "项目组合", "Qwen 调优", "KG 主模型", "滚动均值"]
    values = [22.31, 22.52, 22.89, 23.22, 24.86, 30.97]
    colors = ["#2E74B5", "#6B778D", "#7C3AED", "#9B59B6", "#F39C12", "#AAB2BD"]
    font_path = Path(r"C:\Windows\Fonts\msyh.ttc")
    font = FontProperties(fname=str(font_path)) if font_path.exists() else None
    fig, ax = plt.subplots(figsize=(8.5, 4.2), dpi=180)
    y = list(range(len(labels)))
    ax.barh(y, values, color=colors, height=0.62)
    ax.set_yticks(y, labels=labels, fontproperties=font, fontsize=10)
    ax.invert_yaxis()
    ax.set_xlabel("测试集 WAPE（%，越低越好）", fontproperties=font, fontsize=10)
    ax.set_title("代表性方法的同数据、同切分 WAPE 对比", fontproperties=font, fontsize=13, pad=12)
    ax.set_xlim(0, 35)
    ax.grid(axis="x", color="#D8DEE9", linewidth=0.8, alpha=0.8)
    ax.set_axisbelow(True)
    for i, value in enumerate(values):
        ax.text(value + 0.45, i, f"{value:.2f}%", va="center", fontsize=9)
    for spine in ("top", "right", "left"):
        ax.spines[spine].set_visible(False)
    fig.text(0.5, 0.01, "注：项目“预测吻合度”=100%−WAPE，不是分类准确率。", ha="center",
             fontproperties=font, fontsize=8.5, color="#5B6573")
    plt.tight_layout(rect=[0, 0.05, 1, 1])
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    chart_path = ASSET_DIR / "wape_comparison.png"
    build_chart(chart_path)

    doc = Document()
    setup_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hrun = header.add_run("疫情自动获取系统｜论文基准对比总结")
    set_run_font(hrun, size=8.5, color=MUTED)
    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(8)
    title.add_run("疫情预测论文对比试验\n现状、方法与下一步")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("epidemic-openclaw / 论文知识图谱与受控预测链路")

    meta_rows = [
        ["报告日期", "2026-08-21"],
        ["证据基线", "2026-08-20 R1 closure 本地归档"],
        ["当前状态", "本地实验与浏览器验证完成；vm201 部署因 SSH connection refused 未完成"],
        ["新论文阶段", "2026 PLOS DWT-SARMA-LSTM 已完成预检，等待服务器恢复后执行"],
    ]
    add_table(doc, ["项目", "内容"], meta_rows, [1.35, 5.15], font_size=9)

    add_callout(
        doc,
        "核心结论",
        "现有 2024 RSOS 方法迁移实验中，论文 ARIMA 的测试 WAPE 最低（22.31%）；项目组合模型为 22.89%，差距约 2.62%（相对 WAPE），但项目模型给出更高的预测区间覆盖率。该结果属于历史同数据方法迁移，不是 Fresh 盲测，也不是医学或分类准确率。",
    )

    doc.add_heading("1. 文档范围与结论边界", level=1)
    add_paragraph(doc, "本报告总结已经实际运行的 RSOS 2024 论文方法迁移对比，并说明新选定 PLOS 2026 论文的待执行方案。两者的状态、数据和结论不得混用。")
    add_bullet(doc, "已完成结果：Bleichrodt 等 2024 年 Royal Society Open Science 论文的方法迁移实验。")
    add_bullet(doc, "已完成闭环：冻结数据、同配对 rolling-origin、14 种方法、泄漏哨兵、测试、API/UI staging 浏览器验收。")
    add_bullet(doc, "未完成事项：vm201 线上部署与登录态人工验收；服务器入口目前拒绝连接。")
    add_bullet(doc, "待执行实验：Eligado 等 2026 年 PLOS NTD 的 DWT-SARMA-LSTM 对比，仅完成预注册规划。")

    doc.add_heading("2. 已完成对比实验：论文、数据与任务定义", level=1)
    doc.add_heading("2.1 对比论文", level=2)
    add_paragraph(doc, "论文：Bleichrodt A, Luo R, Kirpich A, Chowell G. Evaluating the forecasting performance of ensemble sub-epidemic frameworks and other time series models for the 2022–2023 mpox epidemic. Royal Society Open Science, 2024, 11(7):240248。DOI：10.1098/rsos.240248。")
    add_paragraph(doc, "该论文有正式 DOI、开放全文、官方代码和原始数据。项目没有在 mpox 原数据上直接宣称完全复现，而是把论文方法迁移到项目的官方新加坡登革热周序列，因此实验身份是 METHOD_TRANSFER_REPLICATION。")

    doc.add_heading("2.2 数据与切分", level=2)
    rows = [
        ["数据来源", "新加坡政府 data.gov.sg 周度登革热病例"],
        ["范围", "2012-W01 至 2019-W48，共412周"],
        ["指标", "cases_new；周频；不插值、不补零"],
        ["训练", "312点（311个标签 + 1个预热点）"],
        ["验证", "50个 origins × H1-H4 = 200个配对"],
        ["测试", "50个 origins；H1-H4 共194个配对（50/49/48/47）"],
        ["比较原则", "所有方法使用完全相同的 origin、horizon 和 target_week"],
        ["评估身份", "RETROSPECTIVE METHOD TRANSFER；测试窗口已消费，非 Fresh"],
    ]
    add_table(doc, ["要素", "冻结定义"], rows, [1.45, 5.05], font_size=9)

    doc.add_heading("3. 实验方法", level=1)
    doc.add_heading("3.1 论文方法组", level=2)
    paper_methods = [
        ["paper_arima", "11周校准窗；AICc 选阶；负预测截断为0。当前同数据实验中表现最好。"],
        ["paper_slr", "局部线性回归，用短校准窗外推未来1-4周。"],
        ["paper_gam", "B样条加GCV的等价实现；在多年季节序列外推时明显发散。"],
        ["sub-epidemic top/second", "GGM/GLM及双波候选，以AICc选择第一名或第二名。"],
        ["sub-epidemic ensemble", "Akaike权重组合或前4名等权组合；100次残差 bootstrap 构造区间。"],
        ["paper_prophet", "Prophet 1.4.0、11周校准窗、周季节项和95%预测区间。"],
    ]
    add_table(doc, ["方法", "实现与作用"], paper_methods, [1.85, 4.65], font_size=9)
    add_callout(doc, "方法忠实度说明", "论文原方法的核心协议得到保留，但 Track 1 发现部分 ARIMA 选阶和 sub-epidemic 候选集与官方产物存在差异；因此不能把全部方法称作逐字节精确复现。", fill="FFF7E6", label_color=GOLD)

    doc.add_heading("3.2 项目方法组", level=2)
    project_methods = [
        ["project_kg_primary", "知识图谱推荐 lagged_count_boosting.v1；使用滞后和滚动计数特征；默认参数50/0.1/10。"],
        ["project_qwen_tuned", "在白名单64组参数内，仅使用验证集选择；最终30棵树、学习率0.2、叶节点最小样本5。"],
        ["project_kg_ensemble", "Qwen调优模型、last-value与seasonal-naive的凸组合；验证集选择权重0.75/0.25/0。"],
        ["统一基线", "last-value、52周seasonal-naive、4周rolling mean。"],
    ]
    add_table(doc, ["方法", "实现与作用"], project_methods, [1.85, 4.65], font_size=9)
    add_paragraph(doc, "特别说明：该次论文对比中的“Qwen 调优”使用与Qwen合同一致的确定性白名单网格，不代表在此轮重新调用大语言模型；其作用是复现受控搜索边界，避免测试集参与选参。")

    doc.add_heading("3.3 指标解释", level=2)
    add_bullet(doc, "WAPE：总绝对误差除以总真实值，越低越好，适合病例计数序列。")
    add_bullet(doc, "预测吻合度：max(0, 1-WAPE)×100%。它只是误差的直观换算，不是分类准确率、临床准确率或公共卫生有效性。")
    add_bullet(doc, "MAE/RMSE：分别反映平均绝对误差和对大误差更敏感的均方根误差。")
    add_bullet(doc, "95% PI覆盖率：真实值落入预测区间的比例；必须与平均区间宽度一起阅读，不能用过宽区间换取表面高覆盖。")

    doc.add_page_break()
    doc.add_heading("4. 已完成实验的主要结果", level=1)
    doc.add_picture(str(chart_path), width=Inches(6.35))
    cap = doc.add_paragraph("图1  代表性方法在同一测试配对集合上的 WAPE 对比")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    for run in cap.runs:
        set_run_font(run, size=8.5, italic=True, color=MUTED)

    key_rows = [
        ["paper_arima", "22.31%", "77.69%", "68.61", "86.32", "80.4%"],
        ["last_value", "22.52%", "77.48%", "69.27", "87.76", "67.0%"],
        ["project_kg_ensemble", "22.89%", "77.11%", "70.41", "88.46", "90.7%"],
        ["project_qwen_tuned", "23.22%", "76.78%", "71.42", "90.08", "94.3%"],
        ["project_kg_primary", "24.86%", "75.14%", "76.48", "97.54", "96.9%"],
        ["rolling_mean_4", "30.97%", "69.03%", "95.26", "115.63", "62.9%"],
        ["seasonal_naive", "81.39%", "18.61%", "250.34", "293.60", "82.0%"],
    ]
    add_table(doc, ["方法", "WAPE", "吻合度", "MAE", "RMSE", "95%覆盖"], key_rows,
              [2.05, 0.75, 0.85, 0.75, 0.85, 1.25], font_size=8.3)

    doc.add_heading("4.1 结果解读", level=2)
    add_number(doc, "论文 ARIMA 的 WAPE 为22.31%，是当前同数据测试中的最低值。")
    add_number(doc, "项目组合模型 WAPE 为22.89%，相对论文 ARIMA 高约2.62%；该差距是真实存在的，不能写成项目显著胜出。")
    add_number(doc, "Qwen受控调优把项目主模型 WAPE 从24.86%降至23.22%，相对改善约6.6%；组合后进一步降至22.89%。")
    add_number(doc, "项目组合模型仍略逊于简单 last-value（22.52%）。原因是测试窗处于较平稳的低发阶段，惯性预测天然占优。")
    add_number(doc, "项目模型的主要优势是不确定性覆盖：KG主模型、Qwen调优和组合模型覆盖率分别约96.9%、94.3%和90.7%，但区间也更宽。")

    add_callout(doc, "学术表述建议", "应写成“项目方法在点预测上接近但未超过论文ARIMA和last-value；受控调优与组合显著改善了KG主模型，且不确定性覆盖更完整”。不要把77%左右吻合度写成分类准确率。")

    doc.add_heading("5. 方法补齐、可信性与验证", level=1)
    add_bullet(doc, "共运行14种方法；验证集200配对、测试集194配对，合计394个时点配对，每个模型行集一致。")
    add_bullet(doc, "future-sentinel 覆盖14种模型：修改 cutoff 后未来数据不改变 cutoff 内预测。")
    add_bullet(doc, "测试集 exactly-once；参数和组合权重只由验证集选择。")
    add_bullet(doc, "旧10种模型逐行字节回归无漂移；双跑 determinism 通过。")
    add_bullet(doc, "123项实验测试通过；ruff、strict mypy、pip check通过。")
    add_bullet(doc, "本地 staging 浏览器44/44断言通过；console error、pageerror和第一方失败请求均为0。")
    add_bullet(doc, "Trusted、Promotion、Active Pointer与正式预测均为零变更。")

    doc.add_heading("6. 当前缺口与部署状态", level=1)
    gap_rows = [
        ["线上部署", "未完成", "vm201 SSH入口持续 connection refused；未修改隧道或SSH配置。"],
        ["登录态人工验收", "待执行", "本地staging已通过，线上登录态需服务器恢复后复验。"],
        ["Track 1忠实度", "有显式差异", "ARIMA选阶、sub-epidemic候选集和官方产物仍有已归因差异。"],
        ["GAM等价性", "未完成", "pygam依赖会破坏旧模型字节一致性，保持OLS B样条+GCV。"],
        ["Fresh盲测", "不适用", "当前测试窗口已被历史运行使用，只能作为回顾性方法迁移基准。"],
    ]
    add_table(doc, ["事项", "状态", "说明"], gap_rows, [1.35, 1.0, 4.15], font_size=8.7)

    doc.add_heading("7. 新论文对比方案（已规划、尚未执行）", level=1)
    add_paragraph(doc, "新选论文：Eligado AA 等，Advancing outbreak detection: Hybridizing machine learning with wavelets for weekly dengue case forecasting，PLOS Neglected Tropical Diseases，2026，DOI 10.1371/journal.pntd.0014444。论文为正式同行评审研究文章，输入仅为周度登革热病例数，公开官方代码。")
    add_callout(doc, "状态边界", "以下内容是预注册方案，不是已经得到的项目结果。作者的 Quezon City 原始病例数据不公开，因此只能做 METHOD_TRANSFER_REPLICATION_ON_FROZEN_PROJECT_DATA，不能称作精确复现。", fill="FFF7E6", label_color=GOLD)

    doc.add_heading("7.1 论文方法", level=2)
    new_methods = [
        ["DWT", "Daubechies 3，level 1，将序列分成低频趋势cA与高频波动cD。"],
        ["SARMA", "cD采用ARMA(0,1)；cA采用SARMA(2,2)×(1,1)[26]。"],
        ["残差LSTM", "对SARMA残差做MinMax缩放，Adam 0.001，batch 4，最多90 epochs，patience 2。"],
        ["输出", "逆DWT重建周度病例预测，并评估95%及95.5%预测区间。"],
    ]
    add_table(doc, ["组成", "预注册实现"], new_methods, [1.35, 5.15], font_size=9)

    doc.add_heading("7.2 双协议比较", level=2)
    add_bullet(doc, "论文对齐协议：前378周训练、后34周固定测试，模拟论文长预测区间。")
    add_bullet(doc, "项目运行协议：312周训练、50个验证 origins、50个测试 origins，H1-H4 rolling-origin。")
    add_bullet(doc, "比较对象：DWT-SARMA、DWT-SARMA-LSTM、项目KG主模型、Qwen调优、项目组合模型及四种统一基线。")
    add_bullet(doc, "LSTM至少运行10个冻结种子并报告全部结果，不得只挑最好的一次。")

    doc.add_heading("7.3 必须修正的潜在泄漏", level=2)
    add_number(doc, "每个 origin 只读取当时已发生的数据；DWT不得对完整测试段一次性拟合。")
    add_number(doc, "小波、层级、SARMA阶数、look-back与LSTM参数只能在训练/验证数据内选择。")
    add_number(doc, "缩放器每折只在训练残差上拟合，early stopping不得读取最终测试标签。")
    add_number(doc, "对未来标签、Qwen上下文、特征和组合权重分别执行future-sentinel。")

    doc.add_heading("8. 后续执行顺序", level=1)
    next_steps = [
        ["1", "恢复连接", "只读核验SSH、项目根目录、Git/Alembic、服务和竞争进程。"],
        ["2", "冻结身份", "重算新加坡数据SHA；固定PLOS PDF与官方代码commit。"],
        ["3", "隔离实现", "在隔离环境实现DWT-SARMA/LSTM和单元测试，不覆盖现有模型。"],
        ["4", "验证选择", "仅用验证集选择参数；冻结10个种子和组合权重。"],
        ["5", "一次测试", "严格同配对运行34周协议及H1-H4协议，test exactly-once。"],
        ["6", "可信验收", "泄漏哨兵、复现、API/DB/UI对账、登录态浏览器和零正式数据变更。"],
        ["7", "交付", "生成逐行预测、指标、图表、回滚说明、manifest和本地ZIP。"],
    ]
    add_table(doc, ["步骤", "阶段", "动作"], next_steps, [0.55, 1.25, 4.70], font_size=8.8)

    doc.add_heading("9. 总体判断", level=1)
    add_paragraph(doc, "现有 RSOS 对比已经证明：知识图谱推荐、受控参数搜索、模型组合、同配对回测和可审计展示可以形成完整的工程闭环；但在当前新加坡登革热测试窗上，简单 ARIMA 与 last-value 仍略优于项目组合模型。项目的价值主要体现在可治理的推荐过程、参数证据、不确定性覆盖和防泄漏，而不是任何数据集上都必然压倒基线。")
    add_paragraph(doc, "新 PLOS 论文更贴近项目目标：同为登革热、仅使用周病例数，并引入DWT-SARMA-LSTM。服务器恢复后，它应作为下一轮主要论文基准；若项目模型仍不占优，应依据验证集结果如实分析数据机制与模型适用性，不能反复查看测试集来追求人为的70%门槛。")

    doc.add_heading("附录A  证据与来源", level=1)
    sources = [
        "本地闭环报告：reports/top_venue_benchmark_r1_closure_20260820/00_FINAL_REPORT.md",
        "本地指标：reports/top_venue_benchmark_r1_closure_20260820/05_METRICS_BY_MODEL_AND_HORIZON.csv",
        "RSOS论文：https://doi.org/10.1098/rsos.240248",
        "PLOS新论文：https://doi.org/10.1371/journal.pntd.0014444",
        "新论文预检：reports/plosntd_dwt_sarma_lstm_benchmark_preflight_20260820/00_SELECTION_AND_PLAN.md",
    ]
    for source in sources:
        add_bullet(doc, source)

    doc.core_properties.title = "疫情预测论文对比试验现状、方法与下一步"
    doc.core_properties.subject = "epidemic-openclaw 论文基准对比总结"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "疫情预测;知识图谱;WAPE;rolling-origin;方法迁移"
    doc.core_properties.comments = "基于本地证据归档生成；未连接远程服务器。"
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
