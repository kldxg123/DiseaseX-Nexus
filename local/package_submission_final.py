from __future__ import annotations

import hashlib
import json
import shutil
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document


ROOT = Path(r"F:\作业\aibola\reports\nc_submission_revision_20260824")
SRC = ROOT / "NC_Submission_Updated_20260824"
DEST = ROOT / "NC_Submission_Final_20260824_Page21_PythonFigures"
ZIP_PATH = ROOT / "NC投稿完善_20260824_正文修订_Python图表_英文界面证据终版.zip"

MAIN_DOCX = "DiseaseX_Nexus_NC_Manuscript_v5_20260824.docx"
SI_DOCX = Path("Supplementary") / "Supplementary_Information_Updated_v2_20260824.docx"
MANIFEST_DOCX = "Content_Manifest_Updated_v2_20260824.docx"

EXCLUDED_DIR_NAMES = {
    "Archive_Previous_20260823",
    "Archive_Previous_20260824_Pre_Page21_CodeFigure_Revision",
    "_qa_main_v5",
    "_qa_si_v2",
    "_qa_manifest_v2",
}

EXCLUDED_FILE_NAMES = {
    "DiseaseX_Nexus_NC_Manuscript_v4_20260824.docx",
    "DiseaseX_Nexus_NC_Manuscript_v4_20260824.pdf",
    "Supplementary_Information_Updated_20260824.docx",
    "Supplementary_Information_Updated_20260824.pdf",
    "Content_Manifest_Updated_20260824.docx",
    "Content_Manifest_Updated_20260824.pdf",
    "Detailed_Experimental_Design_and_Results_20260823.docx",
    "Detailed_Experimental_Design_and_Results_20260823.pdf",
    "PACKAGE_VALIDATION_20260824.json",
    "SHA256SUMS.txt",
    "矢量图.zip",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def doc_text(path: Path) -> str:
    doc = Document(path)
    blocks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            blocks.extend(cell.text for cell in row.cells)
    return "\n".join(blocks)


def contains_cjk(text: str) -> bool:
    return any("\u4e00" <= char <= "\u9fff" for char in text)


def should_copy(relative: Path) -> bool:
    if any(part in EXCLUDED_DIR_NAMES for part in relative.parts):
        return False
    if relative.name in EXCLUDED_FILE_NAMES:
        return False
    return True


def build() -> dict[str, object]:
    if DEST.exists():
        shutil.rmtree(DEST)
    DEST.mkdir(parents=True)

    copied = 0
    for source in SRC.rglob("*"):
        relative = source.relative_to(SRC)
        if not source.is_file() or not should_copy(relative):
            continue
        target = DEST / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, target)
        copied += 1

    main_text = doc_text(DEST / MAIN_DOCX)
    si_text = doc_text(DEST / SI_DOCX)
    manifest_text = doc_text(DEST / MANIFEST_DOCX)

    forbidden_main_tokens = [
        "Data.zip",
        "Results_point.zip",
        "27c3003babd3b4ded5cf5c37a670cacba97028062c7ff3b1e24f82777254237c",
        "f7c289b81ec3d5cf9b075ca8d8fb903b1da20efba29cba312eaea982454b03de",
        "5233f22784cb7d276cc318d400c5935bc143396a91f147015bfea7150b72b0c2",
        "1e213cfe5e938d67cc9d9a6ea8528182841a9e022dec207760585e1ad6c529e4",
    ]
    main_forbidden_hits = [token for token in forbidden_main_tokens if token in main_text]

    figure_provenance = json.loads((DEST / "FIGURE_PROVENANCE_20260824.json").read_text(encoding="utf-8"))
    checks = {
        "main_docx_opens": True,
        "supplementary_docx_opens": True,
        "manifest_docx_opens": True,
        "main_benchmark_has_no_inline_artifact_names_or_hashes": not main_forbidden_hits,
        "supplementary_text_has_no_cjk": not contains_cjk(si_text),
        "manifest_text_has_no_cjk": not contains_cjk(manifest_text),
        "direct_generative_images_used": bool(figure_provenance.get("direct_generative_images_used")),
        "knowledge_graph_figure_changed": bool(figure_provenance.get("knowledge_graph_figure_changed")),
        "python_figure_1_present": (DEST / "Supplementary" / "Python_Generated_Figures_20260824" / "Main_Figure_1_System_Architecture_Python.png").is_file(),
        "python_figure_3_present": (DEST / "Supplementary" / "Python_Generated_Figures_20260824" / "Main_Figure_3_Brazil_Dengue_Python.png").is_file(),
        "python_figure_6_present": (DEST / "Supplementary" / "Table4_Figure6_Experiment_20260824" / "Figure_6_Retrospective_Multi_Pathogen_Evaluation_v2.png").is_file(),
        "authenticated_ui_screenshots_present": all(
            (DEST / "Supplementary" / "UI_Screenshots_20260824" / name).is_file()
            for name in (
                "Supplementary_Figure_5_Authenticated_KG_Recommendation_UI.png",
                "Supplementary_Figure_6_Authenticated_Model_Evaluation_UI.png",
                "Supplementary_Figure_7_Authenticated_Comparison_Experiment_UI.png",
            )
        ),
    }

    required_pass = [
        checks["main_docx_opens"],
        checks["supplementary_docx_opens"],
        checks["manifest_docx_opens"],
        checks["main_benchmark_has_no_inline_artifact_names_or_hashes"],
        checks["supplementary_text_has_no_cjk"],
        checks["manifest_text_has_no_cjk"],
        not checks["direct_generative_images_used"],
        not checks["knowledge_graph_figure_changed"],
        checks["python_figure_1_present"],
        checks["python_figure_3_present"],
        checks["python_figure_6_present"],
        checks["authenticated_ui_screenshots_present"],
    ]

    hashes: list[tuple[str, str]] = []
    for file_path in sorted(path for path in DEST.rglob("*") if path.is_file()):
        relative = file_path.relative_to(DEST).as_posix()
        if relative in {"SHA256SUMS.txt", "PACKAGE_VALIDATION_20260824_v2.json"}:
            continue
        hashes.append((sha256(file_path), relative))

    (DEST / "SHA256SUMS.txt").write_text(
        "".join(f"{digest}  {relative}\n" for digest, relative in hashes),
        encoding="utf-8",
    )

    validation: dict[str, object] = {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "source_directory": str(SRC),
        "final_directory": str(DEST),
        "copied_file_count_before_validation_files": copied,
        "checks": checks,
        "main_forbidden_hits": main_forbidden_hits,
        "sha256_entry_count": len(hashes),
        "status": "PASS" if all(required_pass) else "FAIL",
    }
    (DEST / "PACKAGE_VALIDATION_20260824_v2.json").write_text(
        json.dumps(validation, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    if validation["status"] != "PASS":
        raise RuntimeError(json.dumps(validation, ensure_ascii=False, indent=2))

    if ZIP_PATH.exists():
        ZIP_PATH.unlink()
    with zipfile.ZipFile(ZIP_PATH, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for file_path in sorted(path for path in DEST.rglob("*") if path.is_file()):
            archive.write(file_path, (DEST.name / file_path.relative_to(DEST)).as_posix())

    with zipfile.ZipFile(ZIP_PATH) as archive:
        corrupt = archive.testzip()
        members = archive.namelist()
    if corrupt is not None:
        raise RuntimeError(f"ZIP CRC failure: {corrupt}")

    result = {
        **validation,
        "zip_path": str(ZIP_PATH),
        "zip_sha256": sha256(ZIP_PATH),
        "zip_member_count": len(members),
        "zip_crc": "PASS",
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return result


if __name__ == "__main__":
    build()
