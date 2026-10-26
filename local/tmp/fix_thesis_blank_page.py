from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


DOCX = Path(
    r"F:\作业\aibola\论文材料包_20260814_最终版\01_论文说明文档\疫情自动获取系统_实验架构与预测结果论文材料.docx"
)

replacement = (
    "系统采用 local-first 分层治理：官方数据、论文证据和人工候选先进入可追溯原始层，"
    "规范化后分入 Raw、Trusted、Preview 与 Forecast。知识图谱连接证据、疾病—模型关系和数据画像，"
    "生成受门禁约束的 ExecutionContract；预测由确定性代码执行，Qwen 仅提出白名单内方案。"
)

doc = Document(DOCX)
matched = 0
for paragraph in doc.paragraphs:
    if "local-first" in paragraph.text:
        paragraph.clear()
        run = paragraph.add_run(replacement)
        run.font.name = "Microsoft YaHei"
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10.5)
        paragraph.paragraph_format.space_after = Pt(4)
        matched += 1

if matched != 1:
    candidates = [p.text for p in doc.paragraphs if "local" in p.text.lower()]
    raise RuntimeError(f"expected one matching paragraph, got {matched}; candidates={candidates!r}")

doc.save(DOCX)
print(f"updated={matched}")
