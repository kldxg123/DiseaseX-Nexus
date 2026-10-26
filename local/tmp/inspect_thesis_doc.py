from pathlib import Path

from docx import Document


path = Path(r"F:\作业\aibola\论文材料包_20260813\01_论文说明文档\疫情自动获取系统_实验架构与预测结果论文材料.docx")
doc = Document(path)
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")
for index, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text.strip()
    style = paragraph.style.name if paragraph.style else ""
    if style.startswith("Heading") or any(token in text for token in ("Fresh", "fresh", "轨道", "70%", "结论", "局限")):
        print(f"P{index:03d} [{style}] {text}")
