from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


DOCX = Path(
    r"F:\作业\aibola\论文材料包_20260814_最终版\01_论文说明文档\疫情自动获取系统_实验架构与预测结果论文材料.docx"
)
TMP = DOCX.with_name(DOCX.stem + "_updated.docx")

NAVY = "162033"
LIGHT_GREY = "F4F7FB"
MID_GREY = "D9E2EF"
TEXT = "172235"


def set_run_font(run, size: float = 10.5, bold: bool | None = None, color: str = TEXT) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(cell, color: str = MID_GREY) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        item = tc_borders.find(tag)
        if item is None:
            item = OxmlElement(f"w:{edge}")
            tc_borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), color)


def set_cell_width(cell, width_cm: float) -> None:
    width = int(width_cm * 567)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def configure_table_geometry(table, widths_cm: list[float]) -> None:
    total = int(sum(widths_cm) * 567)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    table.autofit = False
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width_cm in widths_cm:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width_cm * 567)))
        grid.append(grid_col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_cm[index])


doc = Document(DOCX)

for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        if "数据与代码状态：新增第三条 70% 以上历史回放链路；未触发正式发布" in run.text:
            run.text = run.text.replace(
                "数据与代码状态：新增第三条 70% 以上历史回放链路；未触发正式发布",
                "数据与代码状态：三条 70% 以上链路均为非 Fresh；新增 Fresh 候选优化复核；未触发正式发布",
            )

for paragraph in doc.paragraphs:
    text = paragraph.text
    if text.startswith("本材料汇总了“疫情自动获取系统”"):
        run = paragraph.add_run(
            " 补充实验进一步在不读取候选末尾 52 周标签的前提下筛选新 Fresh 链路；所有未注册候选均未通过预注册门禁，因此未发生新的 held-out 消费、预测发布或可视化变更。"
        )
        set_run_font(run)
    elif text.startswith("四条链路均完成了知识图谱推荐"):
        run = paragraph.add_run(
            " 本轮又对未注册流感序列执行两级隔离预筛，但最接近门槛的泰国候选仅达到 65.44% 吻合度，因此没有新增 Fresh 链路。"
        )
        set_run_font(run)
    elif text.startswith("链路 C 的区间覆盖率通过"):
        run = paragraph.add_run(
            " Track C 与后续 Track D 均已被全局注册表标记为 CONSUMED，重新训练不能恢复其 Fresh 身份。"
        )
        set_run_font(run)

reference_anchor = next(
    paragraph for paragraph in doc.paragraphs if paragraph.text.strip() == "参考文献与数据来源"
)
new_elements = []

heading = doc.add_heading("9 Fresh 链路隔离优化复核", level=1)
new_elements.append(heading._p)

p = doc.add_paragraph(
    "本补充实验的目标不是强行得到 70%，而是在不打开真正盲测标签、不影响正式服务的条件下，判断数据库中是否存在值得消费一次性 Fresh 预算的新序列。Track C（智利流感）和 Track D（马来西亚流感）已经处于 CONSUMED 状态，不能通过修改模型、run ID 或 snapshot 恢复 Fresh 身份。"
)
new_elements.append(p._p)

subheading = doc.add_heading("9.1 预筛协议", level=2)
new_elements.append(subheading._p)
for text in (
    "每条未注册候选预留末尾 52 周；SQL 查询只返回此前的调优区间，盲测标签值不进入实验进程。",
    "第一层比较年度季节滞后、滚动均值、幅度校正与组合模型；第二层比较直接 H1–H4 HistGradientBoosting 和 ExtraTrees。",
    "前三个历史 52 周窗口用于选择配置，第四个历史窗口只验证一次；稳定性门禁通过后才允许注册真正 Fresh 窗口。",
):
    bullet = doc.add_paragraph(text, style="List Bullet")
    bullet.paragraph_format.space_after = Pt(3)
    new_elements.append(bullet._p)

subheading = doc.add_heading("9.2 结果与门禁判定", level=2)
new_elements.append(subheading._p)

headers = ["候选", "前三窗中位 WAPE", "前三窗最差 WAPE", "第四窗 WAPE", "第四窗吻合度", "结论"]
rows = [
    ["THA / 流感", "38.83%", "349.42%", "34.56%", "65.44%", "+21.94% 相对基线；未过门禁"],
    ["AUS / 流感", "27.38%", "43.66%", "38.50%", "61.50%", "独立第四窗退化"],
    ["MEX / 流感", "41.52%", "106.10%", "41.93%", "58.07%", "不优于当期基线"],
]
table = doc.add_table(rows=1, cols=len(headers))
table.alignment = 1
widths = [2.7, 2.5, 2.5, 2.1, 2.2, 4.0]
for index, header in enumerate(headers):
    cell = table.rows[0].cells[index]
    shade(cell, NAVY)
    borders(cell, "FFFFFF")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(header)
    set_run_font(run, 8.2, True, "FFFFFF")
header_pr = table.rows[0]._tr.get_or_add_trPr()
header_flag = OxmlElement("w:tblHeader")
header_flag.set(qn("w:val"), "true")
header_pr.append(header_flag)
for row_index, values in enumerate(rows):
    cells = table.add_row().cells
    for index, value in enumerate(values):
        cell = cells[index]
        shade(cell, "FFFFFF" if row_index % 2 == 0 else LIGHT_GREY)
        borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index in (0, 5) else WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(value)
        set_run_font(run, 8.2)
configure_table_geometry(table, widths)
new_elements.append(table._tbl)

p = doc.add_paragraph(
    "最终状态为 FRESH_CHAIN_OPTIMIZATION_NOT_PROMOTED。最接近的泰国候选仍高于 WAPE 30% 上限；澳大利亚在模型选择窗口表现较好，却在独立第四窗明显退化。这说明单变量阳性检出数受到送检量、实验室覆盖、报告制度和季节相位变化的共同影响，增加模型复杂度不足以保证跨年度稳定性。"
)
p.paragraph_format.space_before = Pt(7)
p.paragraph_format.space_after = Pt(5)
new_elements.append(p._p)

subheading = doc.add_heading("9.3 对现有系统与论文结论的影响", level=2)
new_elements.append(subheading._p)
for text in (
    "正式 Git tracked 变化、数据库写入、新 held-out 注册/消费、forecast、Promotion、active pointer 和可视化变化均为 0。",
    "现有 A/B/C/E 四条评估链路及其治理身份保持不变；本轮没有新增第四条超过 70% 的 Fresh 链路。",
    "下一步应联合 specimens_processed 构造检测强度或阳性率特征，并在新官方数据产生前预注册未来窗口。",
):
    bullet = doc.add_paragraph(text, style="List Bullet")
    bullet.paragraph_format.space_after = Pt(3)
    new_elements.append(bullet._p)

for element in new_elements:
    reference_anchor._p.addprevious(element)

doc.save(TMP)
TMP.replace(DOCX)
print(DOCX)
